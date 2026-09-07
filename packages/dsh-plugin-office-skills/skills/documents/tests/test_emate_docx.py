"""Focused e-Mate regressions; stdlib unittest, python-docx and lxml only."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
from docx_common import replace_in_paragraph, iter_all_paragraphs
from docx_create import add_styles
from docx_validate import validate


class EmateDocxTests(unittest.TestCase):
    def test_replacements_never_scan_inserted_text(self):
        for old, new, expected in [('报告', '年度报告', '年度报告年度报告'), ('报告', '报告', '报告报告')]:
            p = Document().add_paragraph('报告报告')
            self.assertEqual(replace_in_paragraph(p, old, new), 2)
            self.assertEqual(p.text, expected)

    def test_cross_run_preserves_unmatched_formatting(self):
        p = Document().add_paragraph()
        p.add_run('前报').bold = True
        p.add_run('告后').italic = True
        self.assertEqual(replace_in_paragraph(p, '报告', '年度报告'), 1)
        self.assertEqual(p.text, '前年度报告后')
        self.assertTrue(p.runs[0].bold)
        self.assertTrue(p.runs[1].italic)
        self.assertEqual(p.runs[1].text, '后')

    def test_linked_headers_and_merged_cells_visited_once(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = '报告'
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = '报告'
        count = sum(replace_in_paragraph(p, '报告', '年度报告') for p in iter_all_paragraphs(doc))
        self.assertEqual(count, 2)
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, '年度报告')

    def test_existing_style_east_asian_font_and_spacing(self):
        doc = Document()
        add_styles(doc, [{'name':'Normal', 'font':'Arial', 'east_asia_font':'Noto Sans CJK SC', 'size_pt':11, 'space_after_pt':6, 'line_spacing':1.35}])
        style = doc.styles['Normal']
        self.assertEqual(style.element.rPr.rFonts.get(qn('w:eastAsia')), 'Noto Sans CJK SC')
        self.assertEqual(style.font.name, 'Arial')
        self.assertEqual(style.paragraph_format.space_after.pt, 6)

    def test_cli_create_read_edit_and_validate(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)
            spec = p / 'spec.json'
            spec.write_text(json.dumps({'blocks':[{'type':'paragraph','text':'报告与中文'}]}, ensure_ascii=False), encoding='utf-8')
            def run(script, *args):
                return json.loads(subprocess.check_output([sys.executable, str(SCRIPTS/script), *map(str,args)], text=True, encoding='utf-8', timeout=10))
            run('docx_create.py',spec,p/'in.docx')
            before = (p/'in.docx').read_bytes()
            run('docx_edit.py','replace',p/'in.docx','--find','报告','--replace','年度报告','-o',p/'out.docx')
            self.assertEqual(run('docx_read.py',p/'out.docx','--text')['body'], ['年度报告与中文'])
            self.assertEqual((p/'in.docx').read_bytes(),before)
            self.assertTrue(run('docx_validate.py',p/'out.docx')['ok'])

    def test_relationships_are_separate_for_document_header_footer(self):
        # Generate distinct image relationships for three actual source parts.
        import base64
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)
            png = p/'a.png'
            png.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl1sAAAAASUVORK5CYII='))
            doc = Document()
            doc.add_paragraph().add_run().add_picture(str(png))
            doc.sections[0].header.paragraphs[0].add_run().add_picture(str(png))
            doc.sections[0].footer.paragraphs[0].add_run().add_picture(str(png))
            doc.save(p/'good.docx')
            self.assertTrue(validate(str(p/'good.docx'))['ok'])
            from lxml import etree
            with zipfile.ZipFile(p/'good.docx') as src, zipfile.ZipFile(p/'bad.docx','w') as dst:
                for name in src.namelist():
                    data = src.read(name)
                    if name == 'word/_rels/header1.xml.rels':
                        root = etree.fromstring(data)
                        for rel in list(root): root.remove(rel)
                        data = etree.tostring(root)
                    dst.writestr(name,data)
            result = validate(str(p/'bad.docx'))
            self.assertFalse(result['ok'])
            errors = [i['detail'] for i in result['issues'] if i['code']=='unresolved-reference']
            self.assertTrue(any('header1.xml' in detail for detail in errors))
            self.assertFalse(any('document.xml' in detail for detail in errors))

if __name__ == '__main__':
    unittest.main()
