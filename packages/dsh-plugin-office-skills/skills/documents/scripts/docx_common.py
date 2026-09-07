#!/usr/bin/env python3
# MIT License. Shared helpers for the docx skill scripts.
"""Shared helpers: paragraph iteration and run-preserving text replacement."""
from __future__ import annotations


def iter_all_paragraphs(doc, include_headers_footers: bool = True):
    """Yield every paragraph in body, tables (recursively), headers, footers."""
    seen = set()
    def unique(container):
        for paragraph in _iter_container(container):
            if paragraph._p not in seen:
                seen.add(paragraph._p)
                yield paragraph
    yield from unique(doc)
    if include_headers_footers:
        for section in doc.sections:
            for part in (
                section.header, section.footer,
                section.first_page_header, section.first_page_footer,
                section.even_page_header, section.even_page_footer,
            ):
                if part is not None:
                    yield from unique(part)


def _iter_container(container):
    for para in container.paragraphs:
        yield para
    for table in container.tables:
        yield from _iter_table(table)


def _iter_table(table):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                yield para
            for nested in cell.tables:
                yield from _iter_table(nested)


def iter_part_roots(doc):
    """Yield the XML root of the body plus every header/footer part."""
    yield doc.element.body
    seen = set()
    for section in doc.sections:
        for part in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            if part is not None and id(part._element) not in seen:
                seen.add(id(part._element))
                yield part._element


def replace_in_paragraph(para, old: str, new: str) -> int:
    """Replace original text matches without rewriting runs or non-text XML.

    A match may span ordinary text nodes/runs. Field results and unsupported
    containers are protected; crossing a drawing, field, tab or break is
    rejected before any mutation. Replacement text inherits its first node's
    run formatting. Existing drawings, field codes and other children survive.
    """
    if not old or old not in para.text:
        return 0
    from docx.oxml.ns import qn

    pieces, nodes, barriers, protected = [], [], set(), []
    offset = 0
    field_depth = 0
    for child in para._p:
        if child.tag == qn("w:pPr"):
            continue
        if child.tag != qn("w:r"):
            text = "".join(t.text or "" for t in child.iter(qn("w:t")))
            barriers.add(offset)
            protected.append((offset, offset + len(text)))
            pieces.append(text)
            offset += len(text)
            barriers.add(offset)
            continue
        for node in child:
            if node.tag == qn("w:rPr"):
                continue
            if node.tag == qn("w:t"):
                text = node.text or ""
                pieces.append(text)
                nodes.append((node, offset, offset + len(text)))
                if field_depth:
                    protected.append((offset, offset + len(text)))
                offset += len(text)
            else:
                barriers.add(offset)
                if node.tag == qn("w:fldChar"):
                    kind = node.get(qn("w:fldCharType"))
                    if kind == "begin":
                        field_depth += 1
                    elif kind == "end":
                        field_depth = max(0, field_depth - 1)
                text = "\t" if node.tag == qn("w:tab") else (
                    "\n" if node.tag in (qn("w:br"), qn("w:cr")) else "")
                pieces.append(text)
                if text:
                    protected.append((offset, offset + len(text)))
                    offset += len(text)
                    barriers.add(offset)
    full = "".join(pieces)
    matches = []
    cursor = 0
    while True:
        start = full.find(old, cursor)
        if start < 0:
            break
        end = start + len(old)
        if any(start < boundary < end for boundary in barriers) or any(
                left < end and right > start for left, right in protected):
            raise ValueError("replacement crosses protected Word structure; document was not saved")
        matches.append((start, end))
        cursor = end
    if matches and any(char in new for char in ("\t", "\r", "\n")):
        raise ValueError("replacement containing tabs or line breaks is unsupported; document was not saved")
    # Validate all matches first, then work backwards on original offsets.
    for start, end in reversed(matches):
        first = True
        for node, left, right in nodes:
            if right <= start or left >= end:
                continue
            cut_start, cut_end = max(start, left) - left, min(end, right) - left
            text = node.text or ""
            node.text = text[:cut_start] + (new if first else "") + text[cut_end:]
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            first = False
    return len(matches)
