"""Core PPTX assembly: create_pptx_with_native_svg."""

from __future__ import annotations

import hashlib
import json
import mimetypes
import os
import posixpath
import random
import re
import shutil
import stat
import subprocess
import tempfile
import uuid
import zipfile
from concurrent.futures import ProcessPoolExecutor, as_completed
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Any
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape, quoteattr

from pptx import Presentation

from pptx_embedded_fonts import (
    FONT_CONTENT_TYPE,
    FONT_REL_TYPE,
    PML_NS as EMBEDDED_FONT_PML_NS,
    REL_NS as EMBEDDED_FONT_REL_NS,
    EmbeddedFontBundle,
    EmbeddedFontError,
    embedded_font_typefaces,
)
from pptx_transitions import (
    AdvanceUpdate,
    DEFAULT_TRANSITION_DURATION,
    EnterUpdate,
    MorphPairExpectation,
    NATIVE_TRANSITIONS,
    apply_slide_motion_xml,
    create_transition_xml,
    normalize_transition_effect_request,
    serialize_source_xml,
    set_directory_use_timings,
    transition_carriers,
    validate_generated_transition_xml,
    validate_pptx_morph_pairs,
    validate_pptx_transition_package,
    validate_seconds,
)
from pptx_animations import (
    ANIMATION_TIMING_OPTION_FIELDS,
    animation_seconds_to_milliseconds,
    create_sequence_timing_xml,
    normalize_animation_effect,
    normalize_animation_effect_request,
    normalize_animation_trigger,
    object_animation_fingerprint,
    pick_animation_effect,
    validate_generated_animation_xml,
    validate_pptx_animation_package,
)
from pptx_opc_validation import (
    canonical_opc_part_path as _canonical_opc_part_path,
    resolve_internal_opc_target as _resolve_internal_opc_target,
    verify_internal_relationships,
)
from pptx_workspace import WorkspaceResourceSpec
from pptx_ooxml.clone import clone_presentation_slides
from pptx_ooxml.package import prune_unreferenced_directory_parts
from language_tags import normalize_language_tag
from hyperlink_contract import (
    HYPERLINK_REL_TYPE,
    trigger_shape_hyperlink_errors,
)

from ..animation_config import (
    MorphPair,
    animation_group_effect_entries,
    resolve_morph_pairs,
    resolve_slide_animation_config,
)
from ..drawingml.context import resolve_text_flow
from ..drawingml.converter import convert_svg_to_slide_shapes
from ..drawingml.theme_colors import (
    ThemeColorError,
    ThemeColorSpec,
    apply_theme_color_spec,
    rewrite_chart_accent_colors,
)
from ..drawingml.theme_fonts import (
    MasterTextStyleSpec,
    ThemeFontSpec,
    apply_master_text_style_spec,
    apply_theme_font_spec,
)
from ..drawingml.utils import EMU_PER_PX
from ..semantic_markers import (
    chrome_token_from_markers,
    page_layout_name_from_svg,
)
from .dimensions import (
    CANVAS_FORMATS,
    resolve_svg_canvas,
)
from .media import (
    PNG_RENDERER,
    get_png_renderer_info, convert_svg_to_png, convert_svg_to_png_cached,
)
from .notes import (
    markdown_to_plain_text,
    create_notes_master_rels_xml,
    create_notes_master_xml,
    create_notes_slide_xml,
    create_notes_slide_rels_xml,
)
from .narration import (
    AUDIO_CONTENT_TYPES,
    AUDIO_REL_TYPE,
    AUDIO_MARKER_PNG_BYTES,
    DEFAULT_NARRATION_START_FLOOR,
    IMAGE_REL_TYPE,
    MEDIA_REL_TYPE,
    apply_recorded_timing,
    inject_narration,
    narration_lead_in_seconds,
    next_shape_id,
    probe_audio_duration,
)
from .slide_xml import (
    create_slide_xml_with_svg, create_slide_rels_xml,
)
from .template_structure import (
    NativeStructureContract,
    OOXML_UINT32_MAX,
    TEMPLATE_PLACEHOLDER_TYPES,
    TemplateElementSpec,
    TemplateSlideSpec,
    TemplateStructureError,
    flat_structure_metadata_errors,
    is_proxy_placeholder,
    match_native_placeholders,
    parse_preserve_slides,
    parse_template_slides,
    template_placeholder_bindings,
)
from .template_validation import validate_pptx_template_package

SLIDE_LAYOUT_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
)
SLIDE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
)
SLIDE_MASTER_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster"
)
NOTES_SLIDE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
)
THEME_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
)
THEME_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.theme+xml"
PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
DML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

for _prefix, _uri in (
    ("p", PML_NS),
    ("a", DML_NS),
    ("r", REL_NS),
    ("p14", P14_NS),
    ("mc", MC_NS),
    ("a14", A14_NS),
    ("m", MATH_NS),
):
    try:
        ET.register_namespace(_prefix, _uri)
    except (ValueError, AttributeError):
        pass


@dataclass(frozen=True)
class PptxStructureContext:
    """Resolved base package structure reused when slide XML is regenerated."""

    slide_layout_targets: dict[int, str]
    slide_master_parts: dict[int, str]

    def slide_layout_target(self, slide_num: int) -> str:
        """Return the slide layout target for a generated slide."""
        try:
            return self.slide_layout_targets[slide_num]
        except KeyError as exc:
            raise RuntimeError(
                f"Missing slide layout relationship for generated slide {slide_num}"
            ) from exc

    def slide_master_part(self, slide_num: int) -> str:
        """Return the slide master package part for a generated slide."""
        try:
            return self.slide_master_parts[slide_num]
        except KeyError as exc:
            raise RuntimeError(
                f"Missing slide master relationship for generated slide {slide_num}"
            ) from exc


@dataclass(frozen=True)
class RoundtripSlidePatch:
    """One authoring slide overlay applied onto its preserved source part."""

    source_ref_ids: frozenset[str]
    edited_ref_ids: frozenset[str]
    deleted_ref_ids: frozenset[str]
    visual_changed: bool
    authoring_visual_changed: bool
    motion_changed: bool
    transition_changed: bool
    transition_replaced: bool
    animation_changed: bool
    notes_changed: bool


@dataclass(frozen=True)
class _NotesMasterReference:
    """Notes master selected for one generated notes slide."""

    package_part: str
    created_theme_part: str | None = None


@dataclass
class _TemplateRuntimeSlide:
    """Parsed slide package state used by explicit Layout structure export."""

    spec: TemplateSlideSpec
    slide_path: Path
    rels_path: Path
    tree: ET.ElementTree
    root: ET.Element
    rels: dict[str, dict[str, str]]
    shapes: dict[str, ET.Element]
    shape_ids_by_svg_id: dict[str, list[str]]


def _relationship_attrs(elem: ET.Element) -> dict[str, str]:
    return {key.rsplit("}", 1)[-1]: value for key, value in elem.attrib.items()}


def _resolve_package_target(source_part: str, target: str) -> str:
    """Resolve a relationship target relative to a package part path."""
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def _relationships_path_for_part(extract_dir: Path, part_name: str) -> Path:
    """Return the package relationship sidecar path for a part name."""
    path = Path(part_name)
    return extract_dir / path.parent / "_rels" / f"{path.name}.rels"


def _find_relationship_target(
    rels_path: Path,
    rel_type: str,
) -> str | None:
    """Find the first relationship target for a relationship type."""
    if not rels_path.exists():
        return None
    root = ET.parse(rels_path).getroot()
    for elem in root:
        attrs = _relationship_attrs(elem)
        if attrs.get("Type") == rel_type:
            return attrs.get("Target")
    return None


def _read_relationships(rels_path: Path) -> dict[str, dict[str, str]]:
    """Return relationship attributes keyed by rId."""
    if not rels_path.exists():
        return {}
    root = ET.parse(rels_path).getroot()
    rels: dict[str, dict[str, str]] = {}
    for elem in root:
        attrs = _relationship_attrs(elem)
        rel_id = attrs.get("Id")
        if rel_id:
            rels[rel_id] = attrs
    return rels


def _find_relationship_id(
    rels_path: Path,
    rel_type: str,
    target: str,
    target_mode: str | None = None,
) -> str | None:
    """Find an existing relationship by type and target."""
    for rel_id, attrs in _read_relationships(rels_path).items():
        if (
            attrs.get("Type") == rel_type
            and attrs.get("Target") == target
            and attrs.get("TargetMode") == target_mode
        ):
            return rel_id
    return None


def _read_slide_layout_targets(extract_dir: Path, slide_count: int) -> PptxStructureContext:
    """Read the actual layout relationship target for every generated slide."""
    slide_layout_targets: dict[int, str] = {}
    slide_master_parts: dict[int, str] = {}
    rels_dir = extract_dir / "ppt" / "slides" / "_rels"
    for slide_num in range(1, slide_count + 1):
        rels_path = rels_dir / f"slide{slide_num}.xml.rels"
        if not rels_path.exists():
            raise RuntimeError(f"Missing slide relationship file: {rels_path}")
        target = _find_relationship_target(rels_path, SLIDE_LAYOUT_REL_TYPE)
        if not target:
            raise RuntimeError(f"Slide {slide_num} has no slide layout relationship")
        slide_layout_targets[slide_num] = target

        slide_part = f"ppt/slides/slide{slide_num}.xml"
        layout_part = _resolve_package_target(slide_part, target)
        layout_rels_path = _relationships_path_for_part(extract_dir, layout_part)
        master_target = _find_relationship_target(layout_rels_path, SLIDE_MASTER_REL_TYPE)
        if not master_target:
            raise RuntimeError(
                f"Slide {slide_num} layout has no slide master relationship"
            )
        slide_master_parts[slide_num] = _resolve_package_target(layout_part, master_target)
    return PptxStructureContext(
        slide_layout_targets=slide_layout_targets,
        slide_master_parts=slide_master_parts,
    )


_SLIDE_BACKGROUND_RE = re.compile(
    r"(?P<prefix><p:cSld\b[^>]*>\s*)"
    r"(?P<bg><p:bg\b.*?</p:bg>)"
    r"(?P<suffix>\s*<p:spTree\b)",
    re.DOTALL,
)


def _extract_slide_background_xml(slide_xml: str) -> str | None:
    """Return the slide-level p:bg XML when it directly precedes spTree."""
    match = _SLIDE_BACKGROUND_RE.search(slide_xml)
    return match.group("bg") if match else None


def _remove_slide_background_xml(slide_xml: str) -> str:
    """Remove a promoted slide-level p:bg from cSld."""
    return _SLIDE_BACKGROUND_RE.sub(r"\g<prefix>\g<suffix>", slide_xml, count=1)


def _put_background_on_part(part_xml: str, background_xml: str) -> str | None:
    """Replace or insert p:bg before a slide/master/layout spTree.

    Returns None when the part carries a p:bg the canonical pattern cannot
    replace; inserting there would leave two p:bg children under p:cSld.
    """
    match = _SLIDE_BACKGROUND_RE.search(part_xml)
    if match:
        return (
            part_xml[:match.start("bg")]
            + background_xml
            + part_xml[match.end("bg"):]
        )
    if "<p:bg" in part_xml:
        return None

    cslide_match = re.search(r"(<p:cSld\b[^>]*>)", part_xml)
    if not cslide_match:
        raise RuntimeError("PPTX slide/master/layout part has no p:cSld element")
    return (
        part_xml[:cslide_match.end()]
        + background_xml
        + part_xml[cslide_match.end():]
    )


def _dominant_variant(
    values_by_slide: dict[int, Any],
) -> tuple[Any | None, list[int]]:
    """Return the most common value and its slides, or None on a tie."""
    slides_by_value: dict[Any, list[int]] = {}
    for slide_num, value in sorted(values_by_slide.items()):
        slides_by_value.setdefault(value, []).append(slide_num)
    if not slides_by_value:
        return None, []
    best_count = max(len(slides) for slides in slides_by_value.values())
    dominant = [
        (value, slides)
        for value, slides in slides_by_value.items()
        if len(slides) == best_count
    ]
    if len(dominant) != 1:
        return None, []
    return dominant[0]


def _is_strict_majority(subset_size: int, total: int) -> bool:
    return subset_size >= 2 and subset_size * 2 > total


def _promote_common_slide_backgrounds_to_masters(
    extract_dir: Path,
    structure: PptxStructureContext,
    slide_count: int,
    *,
    verbose: bool = False,
) -> int:
    """Promote the majority slide background to its shared slide master.

    Every slide in the master group must carry an explicit background —
    a slide without one would start inheriting the promoted master fill.
    Minority slides keep their own slide-level background, which always
    overrides the master fill.
    """
    slides_by_master: dict[str, list[int]] = {}
    for slide_num in range(1, slide_count + 1):
        master_part = structure.slide_master_part(slide_num)
        slides_by_master.setdefault(master_part, []).append(slide_num)

    promoted = 0
    for master_part, slide_nums in slides_by_master.items():
        slide_backgrounds: dict[int, str] = {}
        for slide_num in slide_nums:
            slide_path = extract_dir / "ppt" / "slides" / f"slide{slide_num}.xml"
            slide_xml = slide_path.read_text(encoding="utf-8")
            background_xml = _extract_slide_background_xml(slide_xml)
            if not background_xml:
                slide_backgrounds = {}
                break
            slide_backgrounds[slide_num] = background_xml

        if not slide_backgrounds:
            continue
        background_xml, dominant_slides = _dominant_variant(slide_backgrounds)
        if background_xml is None:
            continue
        if not _is_strict_majority(len(dominant_slides), len(slide_nums)):
            continue

        master_path = extract_dir / master_part
        master_xml = master_path.read_text(encoding="utf-8")
        promoted_master_xml = _put_background_on_part(master_xml, background_xml)
        if promoted_master_xml is None:
            continue
        master_path.write_text(promoted_master_xml, encoding="utf-8")

        for slide_num in dominant_slides:
            slide_path = extract_dir / "ppt" / "slides" / f"slide{slide_num}.xml"
            slide_xml = slide_path.read_text(encoding="utf-8")
            slide_path.write_text(
                _remove_slide_background_xml(slide_xml),
                encoding="utf-8",
            )
            promoted += 1

    if verbose and promoted:
        print(f"  Baseline master background: promoted {promoted} slide background(s)")
    return promoted


_CHROME_TRACE_TOKENS = (
    "logo",
    "footer",
    "header",
    "watermark",
    "chrome",
    "pagenumber",
    "slidenumber",
    "pagenum",
    "slidenum",
)
_TOP_LEVEL_SHAPE_TAGS = {
    f"{{{PML_NS}}}sp",
    f"{{{PML_NS}}}grpSp",
    f"{{{PML_NS}}}pic",
    f"{{{PML_NS}}}cxnSp",
    f"{{{PML_NS}}}graphicFrame",
    f"{{{MC_NS}}}AlternateContent",
}
_FLAT_SYSTEM_PLACEHOLDER_TYPES = frozenset({"dt", "ftr", "sldNum"})
_REL_ATTRS = {
    f"{{{REL_NS}}}embed",
    f"{{{REL_NS}}}link",
    f"{{{REL_NS}}}id",
}


def _chrome_token_from_svg_id(svg_id: str | None) -> str | None:
    """Return the baseline chrome token encoded in a source SVG id."""
    if not svg_id:
        return None
    lower = svg_id.lower()
    compact = re.sub(r"[-_\s]+", "", lower)
    if compact in _CHROME_TRACE_TOKENS:
        return compact
    split_tokens = {token for token in re.split(r"[-_\s]+", lower) if token}
    for token in _CHROME_TRACE_TOKENS:
        if token in split_tokens:
            return token
    return None


def _trace_chrome_shape_ids(
    trace: dict[str, Any] | None,
) -> dict[str, list[str]]:
    """Map chrome token to generated top-level shape ids for one slide."""
    result: dict[str, list[str]] = {}
    if not trace:
        return result
    for event in trace.get("events", []):
        if event.get("decision") != "native":
            continue
        if event.get("animation_override"):
            # Explicitly animated chrome stays slide-local.
            continue
        semantic_role = event.get("data-pptx-role")
        placeholder = event.get("data-pptx-placeholder")
        has_explicit_semantics = (
            semantic_role is not None or placeholder is not None
        )
        token = (
            chrome_token_from_markers(semantic_role, placeholder)
            if has_explicit_semantics
            else _chrome_token_from_svg_id(event.get("id"))
        )
        shape_id = event.get("shape_id")
        if token and shape_id is not None:
            shape_ids = result.setdefault(token, [])
            normalized_shape_id = str(shape_id)
            if normalized_shape_id not in shape_ids:
                shape_ids.append(normalized_shape_id)
    return result


def _trace_native_shape_ids(
    trace: dict[str, Any] | None,
) -> dict[str, list[str]]:
    """Map every traced SVG id to generated top-level shape ids."""
    result: dict[str, list[str]] = {}
    if not trace:
        return result
    for event in trace.get("events", []):
        if event.get("decision") != "native":
            continue
        svg_id = event.get("id")
        shape_id = event.get("shape_id")
        if not svg_id or shape_id is None:
            continue
        shape_ids = result.setdefault(str(svg_id), [])
        normalized = str(shape_id)
        if normalized not in shape_ids:
            shape_ids.append(normalized)
    return result


def _shape_id(elem: ET.Element) -> str | None:
    for cnv in elem.iter(f"{{{PML_NS}}}cNvPr"):
        return cnv.attrib.get("id")
    return None


def _set_shape_name(elem: ET.Element, name: str) -> None:
    """Give one top-level shape a deterministic read-back identity."""
    for cnv in elem.iter(f"{{{PML_NS}}}cNvPr"):
        cnv.set("name", name)
        return
    raise TemplateStructureError(
        f"Cannot name structured shape {name!r}: p:cNvPr is missing"
    )


def _apply_morph_shape_names(
    extract_dir: Path,
    pairs: tuple[MorphPair, ...],
    slide_numbers: dict[str, int],
    shape_ids: dict[tuple[str, str], int],
) -> dict[int, dict[str, str]]:
    """Write forced-Morph names after all structure transformations finish."""
    assignments: dict[int, dict[str, tuple[str, str]]] = {}
    names_by_slide: dict[int, dict[str, str]] = {}
    for pair in pairs:
        for slide_name, group_id in (
            (pair.source_slide, pair.source_group_id),
            (pair.destination_slide, pair.destination_group_id),
        ):
            slide_number = slide_numbers[slide_name]
            shape_id = str(shape_ids[(slide_name, group_id)])
            slide_assignments = assignments.setdefault(slide_number, {})
            previous = slide_assignments.setdefault(
                shape_id,
                (pair.shape_name, group_id),
            )
            if previous[0] != pair.shape_name:
                raise RuntimeError(
                    f'Morph target "{slide_name}/{group_id}" resolves to shape '
                    f'{shape_id} with conflicting names "{previous[0]}" and '
                    f'"{pair.shape_name}"'
                )
            slide_names = names_by_slide.setdefault(slide_number, {})
            previous_group = slide_names.setdefault(pair.shape_name, group_id)
            if previous_group != group_id:
                raise RuntimeError(
                    f'Morph name "{pair.shape_name}" maps to multiple objects '
                    f'on slide "{slide_name}"'
                )

    trace_names: dict[int, dict[str, str]] = {}
    for slide_number, slide_assignments in sorted(assignments.items()):
        slide_path = (
            extract_dir / "ppt" / "slides" / f"slide{slide_number}.xml"
        )
        tree = ET.parse(slide_path)
        root = tree.getroot()
        top_level_shapes = _top_level_shapes_by_id(root)
        desired_names = {
            shape_name
            for shape_name, _group_id in slide_assignments.values()
        }
        for shape_id, shape in top_level_shapes.items():
            if shape_id in slide_assignments:
                continue
            c_nv_pr = next(shape.iter(f"{{{PML_NS}}}cNvPr"), None)
            existing_name = (
                c_nv_pr.get("name") if c_nv_pr is not None else None
            )
            if existing_name in desired_names:
                raise RuntimeError(
                    f'Morph name "{existing_name}" already belongs to an '
                    f'unmapped object on slide {slide_number}'
                )

        for shape_id, (shape_name, group_id) in slide_assignments.items():
            shape = top_level_shapes.get(shape_id)
            if shape is None:
                raise RuntimeError(
                    f'Morph target "{group_id}" no longer resolves to a '
                    f'Slide-local shape on slide {slide_number}'
                )
            _set_shape_name(shape, shape_name)
            trace_names.setdefault(slide_number, {})[group_id] = shape_name
        _write_xml_tree(slide_path, tree)
    return trace_names


def _top_level_shape_name_roster(root: ET.Element) -> tuple[str, ...]:
    """Return the exact visible top-level shape-name sequence for read-back."""
    sp_tree = root.find(f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree")
    if sp_tree is None:
        raise TemplateStructureError("Structured part has no p:cSld/p:spTree")
    names: list[str] = []
    for child in sp_tree:
        if child.tag not in _TOP_LEVEL_SHAPE_TAGS:
            continue
        c_nv_pr = next(child.iter(f"{{{PML_NS}}}cNvPr"), None)
        name = c_nv_pr.get("name") if c_nv_pr is not None else None
        if not name:
            raise TemplateStructureError(
                "Structured part contains a top-level shape without a name"
            )
        names.append(name)
    return tuple(names)


def _top_level_shapes_by_id(root: ET.Element) -> dict[str, ET.Element]:
    sp_tree = root.find(f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree")
    if sp_tree is None:
        return {}
    shapes: dict[str, ET.Element] = {}
    for child in list(sp_tree):
        if child.tag not in _TOP_LEVEL_SHAPE_TAGS:
            continue
        shape_id = _shape_id(child)
        if shape_id:
            shapes[shape_id] = child
    return shapes


def _timing_shape_ids(root: ET.Element) -> set[str]:
    """Return slide-local shape ids referenced by animation timing."""
    return {
        elem.attrib["spid"]
        for elem in root.findall(f".//{{{PML_NS}}}timing//{{{PML_NS}}}spTgt")
        if elem.attrib.get("spid")
    }


def _relationship_ids_in_shape(elem: ET.Element) -> set[str]:
    rel_ids: set[str] = set()
    for node in elem.iter():
        for attr_name, value in node.attrib.items():
            if attr_name in _REL_ATTRS and value:
                rel_ids.add(value)
    return rel_ids


def _shape_relationships_supported(
    elem: ET.Element,
    rels: dict[str, dict[str, str]],
) -> bool:
    """Return whether every shape relation can move to Master/Layout parts."""
    for rel_id in _relationship_ids_in_shape(elem):
        attrs = rels.get(rel_id)
        if not attrs:
            return False
        rel_type = attrs.get("Type")
        target_mode = attrs.get("TargetMode")
        if rel_type == IMAGE_REL_TYPE and not target_mode:
            continue
        if rel_type == HYPERLINK_REL_TYPE and target_mode == "External":
            continue
        if rel_type == SLIDE_REL_TYPE and not target_mode:
            continue
        return False
    return True


def _canonical_shape_xml(
    elem: ET.Element,
    rels: dict[str, dict[str, str]],
) -> bytes:
    """Canonicalize ids and relationship ids for cross-slide equality."""
    clone = ET.fromstring(ET.tostring(elem, encoding="utf-8"))
    for cnv in clone.iter(f"{{{PML_NS}}}cNvPr"):
        cnv.set("id", "ID")
        # Generated names include the slide-local shape id (for example,
        # ``Image 2`` versus ``Image 8``) but do not affect rendering.
        if "name" in cnv.attrib:
            cnv.set("name", "NAME")
    for fld in clone.iter(f"{{{DML_NS}}}fld"):
        # The literal inside a slide-number field is a per-slide render
        # cache that PowerPoint recomputes from the slide position.
        if fld.attrib.get("type") == "slidenum":
            cached = fld.find(f"{{{DML_NS}}}t")
            if cached is not None:
                cached.text = ""
    for node in clone.iter():
        for attr_name, value in list(node.attrib.items()):
            if attr_name not in _REL_ATTRS:
                continue
            attrs = rels.get(value, {})
            node.set(
                attr_name,
                f"{attrs.get('Type', '')}|{attrs.get('Target', '')}|"
                f"{attrs.get('TargetMode', '')}",
            )
    return ET.tostring(clone, encoding="utf-8")


def _ensure_relationship(
    rels_path: Path,
    rel_type: str,
    target: str,
    target_mode: str | None = None,
) -> str:
    existing = _find_relationship_id(
        rels_path,
        rel_type,
        target,
        target_mode,
    )
    if existing:
        return existing
    return _append_relationship(
        rels_path,
        rel_type,
        target,
        target_mode=target_mode,
    )


def _part_name_for_relationships_path(rels_path: Path) -> str:
    """Recover one ``ppt/...`` package part from its relationship sidecar."""
    if rels_path.parent.name != "_rels" or not rels_path.name.endswith(".rels"):
        raise RuntimeError(f"Invalid PPTX relationship path: {rels_path}")
    part_path = rels_path.parent.parent / rels_path.name.removesuffix(".rels")
    parts = part_path.parts
    try:
        ppt_index = len(parts) - 1 - tuple(reversed(parts)).index("ppt")
    except ValueError as exc:
        raise RuntimeError(
            f"Relationship path is not under a ppt package: {rels_path}"
        ) from exc
    return PurePosixPath(*parts[ppt_index:]).as_posix()


def _copy_shape_relationships_to_part(
    elem: ET.Element,
    slide_rels: dict[str, dict[str, str]],
    target_rels_path: Path,
) -> ET.Element:
    """Clone a shape and retarget supported relationship ids to another part."""
    clone = ET.fromstring(ET.tostring(elem, encoding="utf-8"))
    target_part = _part_name_for_relationships_path(target_rels_path)
    for node in clone.iter():
        for attr_name, value in list(node.attrib.items()):
            if attr_name not in _REL_ATTRS:
                continue
            rel = slide_rels.get(value)
            if not rel:
                raise RuntimeError(f"Missing slide relationship for {value}")
            target_mode = rel.get("TargetMode")
            relationship_target = rel["Target"]
            if target_mode != "External":
                resolved_target = _resolve_package_target(
                    "ppt/slides/source.xml",
                    relationship_target,
                )
                relationship_target = posixpath.relpath(
                    resolved_target,
                    posixpath.dirname(target_part),
                )
            new_rid = _ensure_relationship(
                target_rels_path,
                rel["Type"],
                relationship_target,
                target_mode,
            )
            node.set(attr_name, new_rid)
    return clone


def _copy_shape_relationships_to_master(
    elem: ET.Element,
    slide_rels: dict[str, dict[str, str]],
    master_rels_path: Path,
) -> ET.Element:
    """Clone a shape and retarget supported relationship ids to the master."""
    return _copy_shape_relationships_to_part(elem, slide_rels, master_rels_path)


def _next_master_shape_id(master_xml: str) -> int:
    ids = [
        int(match)
        for match in re.findall(r"<p:cNvPr\b[^>]*\bid=\"(\d+)\"", master_xml)
    ]
    return max(ids, default=1) + 1


def _renumber_shape_ids(elem: ET.Element, start_id: int) -> None:
    next_id = start_id
    for cnv in elem.iter(f"{{{PML_NS}}}cNvPr"):
        cnv.set("id", str(next_id))
        next_id += 1


def _append_shape_to_master(master_path: Path, elem: ET.Element) -> None:
    master_xml = master_path.read_text(encoding="utf-8")
    _renumber_shape_ids(elem, _next_master_shape_id(master_xml))
    shape_xml = ET.tostring(elem, encoding="unicode")
    if "</p:spTree>" not in master_xml:
        raise RuntimeError(f"Slide master has no p:spTree: {master_path}")
    master_path.write_text(
        master_xml.replace("</p:spTree>", f"{shape_xml}\n</p:spTree>", 1),
        encoding="utf-8",
    )


def _append_shape_to_part(part_path: Path, elem: ET.Element) -> None:
    """Append a top-level shape to a master/layout spTree with fresh ids."""
    tree = ET.parse(part_path)
    root = tree.getroot()
    sp_tree = root.find(f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree")
    if sp_tree is None:
        raise RuntimeError(f"PPTX part has no p:spTree: {part_path}")
    existing_ids = [
        int(cnv.attrib["id"])
        for cnv in root.iter(f"{{{PML_NS}}}cNvPr")
        if cnv.attrib.get("id", "").isdigit()
    ]
    clone = ET.fromstring(ET.tostring(elem, encoding="utf-8"))
    _renumber_shape_ids(clone, max(existing_ids, default=1) + 1)
    sp_tree.append(clone)
    _write_xml_tree(part_path, tree)


def _write_xml_tree(path: Path, tree: ET.ElementTree) -> None:
    root = tree.getroot()
    if root.tag == f"{{{PACKAGE_REL_NS}}}Relationships":
        ET.register_namespace("", PACKAGE_REL_NS)
    tree.write(path, encoding="utf-8", xml_declaration=True)


def _slide_ref_shape_ids(values: frozenset[str]) -> set[str]:
    """Return Slide-local ids from authoring source references."""
    result: set[str] = set()
    for value in values:
        if value.startswith("slide:") and len(value) > len("slide:"):
            result.add(value.split(":", 1)[1])
    return result


def _roundtrip_synthetic_shape_id(source_order: tuple[int, ...]) -> str:
    """Return the importer's stable identity for a shape with no native id."""
    return "missing-" + "-".join(str(value) for value in source_order)


def _direct_shape_id(element: ET.Element) -> str | None:
    """Return one shape's own cNvPr id without descending into child shapes."""
    if element.tag == f"{{{MC_NS}}}AlternateContent":
        return _shape_id(element)
    for child in element:
        c_nv_pr = child.find(f"{{{PML_NS}}}cNvPr")
        if c_nv_pr is not None:
            return c_nv_pr.get("id")
    return None


def _roundtrip_shape_id_closure(
    element: ET.Element,
    source_order: tuple[int, ...],
) -> set[str]:
    """Return native and synthetic identities in one source shape subtree."""
    identity = _direct_shape_id(element) or _roundtrip_synthetic_shape_id(
        source_order
    )
    identities = _shape_ids_in_element(element) | {identity}
    if element.tag != f"{{{PML_NS}}}grpSp":
        return identities
    child_order = 0
    for child in element:
        if child.tag not in _TOP_LEVEL_SHAPE_TAGS:
            continue
        child_order += 1
        identities.update(
            _roundtrip_shape_id_closure(
                child,
                (*source_order, child_order),
            )
        )
    return identities


def _index_roundtrip_source_shape(
    element: ET.Element,
    source_order: tuple[int, ...],
    shapes: dict[str, ET.Element],
    key_by_element: dict[int, str],
) -> None:
    """Index one native source shape subtree by stable round-trip identity."""
    shape_id = _direct_shape_id(element) or _roundtrip_synthetic_shape_id(
        source_order
    )
    if shape_id in shapes:
        raise TemplateStructureError(
            f"Round-trip source slide repeats shape id {shape_id}"
        )
    shapes[shape_id] = element
    key_by_element[id(element)] = shape_id
    if element.tag != f"{{{PML_NS}}}grpSp":
        return
    child_order = 0
    for child in element:
        if child.tag not in _TOP_LEVEL_SHAPE_TAGS:
            continue
        child_order += 1
        _index_roundtrip_source_shape(
            child,
            (*source_order, child_order),
            shapes,
            key_by_element,
        )


def _roundtrip_source_shape_roster(
    source_tree: ET.Element,
) -> tuple[
    dict[str, ET.Element],
    dict[str, set[str]],
    dict[int, str],
    dict[str, ET.Element],
]:
    """Index source top-level shapes without mutating missing native ids."""
    shapes: dict[str, ET.Element] = {}
    descendant_ids: dict[str, set[str]] = {}
    key_by_element: dict[int, str] = {}
    all_shapes: dict[str, ET.Element] = {}
    source_order = 0
    for child in source_tree:
        if child.tag not in _TOP_LEVEL_SHAPE_TAGS:
            continue
        source_order += 1
        order_path = (source_order,)
        _index_roundtrip_source_shape(
            child,
            order_path,
            all_shapes,
            key_by_element,
        )
        shape_id = key_by_element[id(child)]
        shapes[shape_id] = child
        descendant_ids[shape_id] = _roundtrip_shape_id_closure(
            child,
            order_path,
        )
    return shapes, descendant_ids, key_by_element, all_shapes


def _roundtrip_nested_restore_ids(
    source_shapes: dict[str, ET.Element],
    source_key_by_element: dict[int, str],
    affected_top_ids: set[str],
    unchanged_ids: set[str],
    deleted_ids: set[str],
) -> set[str]:
    """Return outermost unchanged source shapes inside rebuilt top-level groups."""
    restore_ids: set[str] = set()

    def visit(element: ET.Element) -> None:
        shape_id = source_key_by_element[id(element)]
        if shape_id in unchanged_ids:
            restore_ids.add(shape_id)
            return
        if shape_id in deleted_ids or element.tag != f"{{{PML_NS}}}grpSp":
            return
        for child in element:
            if child.tag in _TOP_LEVEL_SHAPE_TAGS:
                visit(child)

    for shape_id in affected_top_ids:
        visit(source_shapes[shape_id])
    return restore_ids


def _restore_roundtrip_nested_source_shapes(
    generated_nodes: list[ET.Element],
    source_shapes: dict[str, ET.Element],
    restore_ids: set[str],
    traced_ids: dict[str, list[str]],
) -> None:
    """Replace converter-safe nested placeholders with exact source shapes."""
    locations: dict[str, tuple[ET.Element, ET.Element]] = {}

    def index_children(parent: ET.Element) -> None:
        for child in parent:
            if child.tag not in _TOP_LEVEL_SHAPE_TAGS:
                continue
            shape_id = _direct_shape_id(child)
            if shape_id is not None:
                previous = locations.setdefault(shape_id, (parent, child))
                if previous[1] is not child:
                    raise TemplateStructureError(
                        "Generated round-trip slide repeats nested shape id "
                        f"{shape_id}"
                    )
            if child.tag == f"{{{PML_NS}}}grpSp":
                index_children(child)

    for node in generated_nodes:
        if node.tag == f"{{{PML_NS}}}grpSp":
            index_children(node)

    for source_id in sorted(restore_ids, key=_roundtrip_shape_sort_key):
        candidate_ids = {source_id}
        candidate_ids.update(traced_ids.get(f"shape-{source_id}", []))
        matches = {
            candidate_id: locations[candidate_id]
            for candidate_id in candidate_ids
            if candidate_id in locations
        }
        if len(matches) != 1:
            raise TemplateStructureError(
                "Unchanged nested round-trip source object did not produce "
                f"exactly one restore placeholder: {source_id}"
            )
        parent, placeholder = next(iter(matches.values()))
        source = source_shapes.get(source_id)
        if source is None:
            raise TemplateStructureError(
                f"Round-trip source shape is missing for nested ref {source_id}"
            )
        clone = ET.fromstring(ET.tostring(source, encoding="utf-8"))
        clone.tail = placeholder.tail
        position = list(parent).index(placeholder)
        parent.remove(placeholder)
        parent.insert(position, clone)


def _roundtrip_shape_sort_key(value: str) -> tuple[int, int | str]:
    """Sort native numeric ids before stable synthetic ids."""
    return (0, int(value)) if value.isdigit() else (1, value)


def _relationship_key(attrs: dict[str, str]) -> tuple[str, str, str | None]:
    """Return the identity used to reuse one source relationship."""
    return (
        attrs.get("Type", ""),
        attrs.get("Target", ""),
        attrs.get("TargetMode"),
    )


def _merge_roundtrip_relationships(
    source_rels: bytes,
    generated_rels_path: Path,
    generated_slide_root: ET.Element,
    required_generated_ids: set[str],
) -> None:
    """Retarget generated references into the preserved source rel roster."""
    try:
        source_root = ET.fromstring(source_rels)
        generated_root = ET.parse(generated_rels_path).getroot()
    except ET.ParseError as exc:
        raise TemplateStructureError(
            "Round-trip slide relationships are not valid XML"
        ) from exc

    source_by_key: dict[tuple[str, str, str | None], str] = {}
    used_ids: set[str] = set()
    for relationship in source_root:
        attrs = _relationship_attrs(relationship)
        rel_id = attrs.get("Id")
        if not rel_id:
            continue
        used_ids.add(rel_id)
        source_by_key.setdefault(_relationship_key(attrs), rel_id)
    numeric_ids = [
        int(match.group(1))
        for rel_id in used_ids
        if (match := re.fullmatch(r"rId(\d+)", rel_id)) is not None
    ]
    next_id = max(numeric_ids, default=0) + 1

    id_map: dict[str, str] = {}
    added_relationship = False
    for relationship in generated_root:
        attrs = _relationship_attrs(relationship)
        generated_id = attrs.get("Id")
        if not generated_id:
            raise TemplateStructureError(
                "Generated round-trip relationship has no Id"
            )
        if generated_id not in required_generated_ids:
            continue
        existing_id = source_by_key.get(_relationship_key(attrs))
        if existing_id is not None:
            id_map[generated_id] = existing_id
            continue

        assigned_id = generated_id
        if assigned_id in used_ids:
            while f"rId{next_id}" in used_ids:
                next_id += 1
            assigned_id = f"rId{next_id}"
            next_id += 1
        clone = ET.fromstring(ET.tostring(relationship, encoding="utf-8"))
        clone.set("Id", assigned_id)
        source_root.append(clone)
        added_relationship = True
        used_ids.add(assigned_id)
        source_by_key[_relationship_key(attrs)] = assigned_id
        id_map[generated_id] = assigned_id

    missing_relationships = required_generated_ids - set(id_map)
    if missing_relationships:
        raise TemplateStructureError(
            "Generated round-trip slide references missing relationship(s): "
            + ", ".join(sorted(missing_relationships))
        )

    for node in generated_slide_root.iter():
        for attr_name, value in list(node.attrib.items()):
            if attr_name in _REL_ATTRS and value in id_map:
                node.set(attr_name, id_map[value])
    if added_relationship:
        _write_xml_tree(generated_rels_path, ET.ElementTree(source_root))
    else:
        generated_rels_path.write_bytes(source_rels)


def _replace_roundtrip_motion(
    source_root: ET.Element,
    generated_root: ET.Element,
    *,
    transition_changed: bool,
    transition_replaced: bool,
    animation_changed: bool,
) -> None:
    """Replace only explicitly changed transition and animation state."""
    if transition_changed and transition_replaced:
        generated_carriers = transition_carriers(generated_root)
        if len(generated_carriers) > 1:
            raise TemplateStructureError(
                "Generated round-trip slide has multiple transition carriers"
            )
        for carrier in transition_carriers(source_root):
            source_root.remove(carrier)
        if generated_carriers:
            clone = ET.fromstring(
                ET.tostring(generated_carriers[0], encoding="utf-8")
            )
            children = list(source_root)
            insert_at = next(
                (
                    index
                    for index, child in enumerate(children)
                    if child.tag in {
                        f"{{{PML_NS}}}timing",
                        f"{{{PML_NS}}}extLst",
                    }
                ),
                len(children),
            )
            source_root.insert(insert_at, clone)
    elif transition_changed:
        _replace_roundtrip_transition_advance(source_root, generated_root)

    if animation_changed:
        timing_tag = f"{{{PML_NS}}}timing"
        existing_timing = source_root.find(timing_tag)
        generated_timing = generated_root.find(timing_tag)
        if existing_timing is not None:
            source_root.remove(existing_timing)
        if generated_timing is None:
            return
        clone = ET.fromstring(ET.tostring(generated_timing, encoding="utf-8"))
        children = list(source_root)
        insert_at = next(
            (
                index
                for index, child in enumerate(children)
                if child.tag == f"{{{PML_NS}}}extLst"
            ),
            len(children),
        )
        source_root.insert(insert_at, clone)


def _replace_roundtrip_transition_advance(
    source_root: ET.Element,
    generated_root: ET.Element,
) -> None:
    """Copy only advance attributes while preserving the source visual effect."""
    source_carriers = transition_carriers(source_root)
    generated_carriers = transition_carriers(generated_root)
    if len(source_carriers) > 1 or len(generated_carriers) > 1:
        raise TemplateStructureError(
            "Round-trip transition advance patch requires at most one carrier"
        )

    transition_tag = f"{{{PML_NS}}}transition"

    def _elements(carrier: ET.Element) -> list[ET.Element]:
        if carrier.tag == transition_tag:
            return [carrier]
        return list(carrier.iter(transition_tag))

    generated_transitions = (
        _elements(generated_carriers[0]) if generated_carriers else []
    )
    advance_values = {
        (transition.get("advClick"), transition.get("advTm"))
        for transition in generated_transitions
    }
    if len(advance_values) > 1:
        raise TemplateStructureError(
            "Generated transition branches disagree on advance attributes"
        )
    advance_click, advance_after = (
        next(iter(advance_values)) if advance_values else (None, None)
    )

    if not source_carriers:
        if generated_carriers:
            clone = ET.fromstring(
                ET.tostring(generated_carriers[0], encoding="utf-8")
            )
            children = list(source_root)
            insert_at = next(
                (
                    index
                    for index, child in enumerate(children)
                    if child.tag in {
                        f"{{{PML_NS}}}timing",
                        f"{{{PML_NS}}}extLst",
                    }
                ),
                len(children),
            )
            source_root.insert(insert_at, clone)
        return

    source_transitions = _elements(source_carriers[0])
    if not source_transitions:
        raise TemplateStructureError(
            "Source logical transition carrier contains no p:transition"
        )
    for transition in source_transitions:
        for attribute, value in (
            ("advClick", advance_click),
            ("advTm", advance_after),
        ):
            if value is None:
                transition.attrib.pop(attribute, None)
            else:
                transition.set(attribute, value)


_RAW_ALTERNATE_CONTENT_RE = re.compile(
    rb"<mc:AlternateContent\b[^>]*>.*?</mc:AlternateContent\s*>",
    re.DOTALL,
)
_RAW_DIRECT_TRANSITION_RE = re.compile(
    rb"<p:transition\b(?:[^>]*/>|[^>]*>.*?</p:transition\s*>)",
    re.DOTALL,
)
_RAW_TIMING_RE = re.compile(
    rb"<p:timing\b(?:[^>]*/>|[^>]*>.*?</p:timing\s*>)",
    re.DOTALL,
)


def _raw_transition_carrier_span(xml_data: bytes) -> tuple[int, int] | None:
    """Locate one raw root transition carrier in PowerPoint slide XML."""
    alternate_spans = [
        match.span()
        for match in _RAW_ALTERNATE_CONTENT_RE.finditer(xml_data)
    ]
    carriers = [
        span
        for span in alternate_spans
        if _RAW_DIRECT_TRANSITION_RE.search(xml_data[span[0]:span[1]])
    ]
    for match in _RAW_DIRECT_TRANSITION_RE.finditer(xml_data):
        if any(start <= match.start() < end for start, end in alternate_spans):
            continue
        carriers.append(match.span())
    if len(carriers) > 1:
        raise TemplateStructureError(
            "Round-trip slide has multiple raw transition carriers"
        )
    return carriers[0] if carriers else None


def _preserve_roundtrip_transition_markup(
    serialized_slide: bytes,
    source_slide: bytes,
) -> bytes:
    """Restore an unchanged source transition carrier byte-for-byte."""
    source_span = _raw_transition_carrier_span(source_slide)
    serialized_span = _raw_transition_carrier_span(serialized_slide)
    if source_span is None:
        if serialized_span is not None:
            raise TemplateStructureError(
                "Round-trip serialization introduced a transition carrier"
            )
        return serialized_slide
    if serialized_span is None:
        raise TemplateStructureError(
            "Round-trip serialization removed the preserved transition carrier"
        )
    return (
        serialized_slide[:serialized_span[0]]
        + source_slide[source_span[0]:source_span[1]]
        + serialized_slide[serialized_span[1]:]
    )


def _raw_timing_span(xml_data: bytes) -> tuple[int, int] | None:
    """Locate one raw logical object-animation timing carrier in slide XML."""
    alternate_spans = [
        match.span()
        for match in _RAW_ALTERNATE_CONTENT_RE.finditer(xml_data)
    ]
    timings = [
        span
        for span in alternate_spans
        if _RAW_TIMING_RE.search(xml_data[span[0]:span[1]])
    ]
    for match in _RAW_TIMING_RE.finditer(xml_data):
        if any(start <= match.start() < end for start, end in alternate_spans):
            continue
        timings.append(match.span())
    if len(timings) > 1:
        raise TemplateStructureError(
            "Round-trip slide has multiple raw object-animation timing carriers"
        )
    return timings[0] if timings else None


def _preserve_roundtrip_timing_markup(
    serialized_slide: bytes,
    source_slide: bytes,
) -> bytes:
    """Restore unchanged source object-animation timing byte-for-byte."""
    source_span = _raw_timing_span(source_slide)
    serialized_span = _raw_timing_span(serialized_slide)
    if source_span is None:
        if serialized_span is not None:
            raise TemplateStructureError(
                "Round-trip serialization introduced object-animation timing"
            )
        return serialized_slide
    if serialized_span is None:
        raise TemplateStructureError(
            "Round-trip serialization removed preserved object-animation timing"
        )
    return (
        serialized_slide[:serialized_span[0]]
        + source_slide[source_span[0]:source_span[1]]
        + serialized_slide[serialized_span[1]:]
    )


def _apply_roundtrip_transition_overlay(
    slide_path: Path,
    *,
    effect: str | None,
    effect_options: dict[str, object],
    duration: float,
    auto_advance: float | None,
    replace_transition: bool,
) -> bool:
    """Patch only transition/advance state while preserving source timing."""
    source_bytes = slide_path.read_bytes()
    source_xml = source_bytes.decode("utf-8")
    source_animation = object_animation_fingerprint(source_xml)
    if not replace_transition:
        enter = EnterUpdate(policy="preserve", duration=duration)
    elif effect is None:
        enter = EnterUpdate(
            policy="none",
            effect=None,
            duration=(duration or DEFAULT_TRANSITION_DURATION),
        )
    else:
        enter = EnterUpdate(
            policy="replace",
            effect=effect,
            duration=duration,
            effect_options=effect_options,
        )
    advance = (
        AdvanceUpdate(mode="click")
        if auto_advance is None
        else AdvanceUpdate(mode="both", after=auto_advance)
    )
    updated_xml, uses_timings = apply_slide_motion_xml(
        source_xml,
        enter=enter,
        advance=advance,
    )
    if object_animation_fingerprint(updated_xml) != source_animation:
        raise RuntimeError(
            "Round-trip transition overlay changed source object animations"
        )
    if replace_transition:
        validate_generated_transition_xml(
            updated_xml,
            effect=effect,
            effect_options=effect_options,
            duration=duration,
            advance_on_click=True,
            advance_after=auto_advance,
        )
    slide_path.write_bytes(
        _preserve_roundtrip_timing_markup(
            updated_xml.encode("utf-8"),
            source_bytes,
        )
    )
    return uses_timings


def _shape_ids_in_element(element: ET.Element) -> set[str]:
    """Return every non-visual DrawingML id owned by one shape subtree."""
    return {
        value
        for node in element.iter(f"{{{PML_NS}}}cNvPr")
        if (value := node.get("id"))
    }


def _rewrite_roundtrip_timing_shape_ids(
    slide_root: ET.Element,
    id_map: dict[str, str],
) -> None:
    """Rewrite every timing reference to a remapped Slide-local shape id."""
    if not id_map:
        return
    timing = slide_root.find(f"{{{PML_NS}}}timing")
    if timing is None:
        return
    for node in timing.iter():
        old_id = node.get("spid")
        if old_id in id_map:
            node.set("spid", id_map[old_id])


def _renumber_roundtrip_generated_shapes(
    generated_shapes: list[ET.Element],
    retained_source_nodes: list[ET.Element],
    slide_root: ET.Element,
    *,
    rewrite_timing: bool,
) -> dict[str, str]:
    """Keep generated additions disjoint from retained source shape ids."""
    used_ids = {
        value
        for element in retained_source_nodes
        for value in _shape_ids_in_element(element)
    }
    all_ids = set(used_ids)
    for element in generated_shapes:
        all_ids.update(_shape_ids_in_element(element))
    numeric_ids = [int(value) for value in all_ids if value.isdigit()]
    next_id = max(numeric_ids, default=1) + 1
    id_map: dict[str, str] = {}
    for element in generated_shapes:
        for node in element.iter(f"{{{PML_NS}}}cNvPr"):
            old_id = node.get("id")
            if not old_id:
                continue
            if old_id in used_ids:
                while str(next_id) in all_ids:
                    next_id += 1
                new_id = str(next_id)
                next_id += 1
                node.set("id", new_id)
                id_map[old_id] = new_id
                all_ids.add(new_id)
                used_ids.add(new_id)
            else:
                used_ids.add(old_id)

    if not id_map:
        return {}
    for element in generated_shapes:
        for connector_tag in ("stCxn", "endCxn"):
            for connector in element.iter(f"{{{PML_NS}}}{connector_tag}"):
                old_id = connector.get("id")
                if old_id in id_map:
                    connector.set("id", id_map[old_id])
    if rewrite_timing:
        _rewrite_roundtrip_timing_shape_ids(slide_root, id_map)
    return id_map


def _apply_roundtrip_slide_overlay(
    source_slide: bytes,
    source_rels: bytes,
    generated_slide_path: Path,
    generated_rels_path: Path,
    patch: RoundtripSlidePatch,
    conversion_trace: dict[str, Any] | None,
) -> dict[int, int]:
    """Overlay authored Slide objects onto the exact source slide package part."""
    try:
        source_root = ET.fromstring(source_slide)
        generated_tree = ET.parse(generated_slide_path)
    except ET.ParseError as exc:
        raise TemplateStructureError(
            "Round-trip source or generated slide is not valid XML"
        ) from exc
    generated_root = generated_tree.getroot()

    source_tree = source_root.find(
        f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree"
    )
    generated_sp_tree = generated_root.find(
        f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree"
    )
    if source_tree is None or generated_sp_tree is None:
        raise TemplateStructureError(
            "Round-trip slide overlay requires source and generated p:spTree"
        )

    (
        source_shapes,
        source_descendant_ids,
        source_key_by_element,
        all_source_shapes,
    ) = _roundtrip_source_shape_roster(source_tree)

    generated_order: list[str] = []
    generated_shapes: dict[str, ET.Element] = {}
    generated_top_id_by_descendant: dict[str, str] = {}
    for child in generated_sp_tree:
        if child.tag not in _TOP_LEVEL_SHAPE_TAGS:
            continue
        shape_id = _shape_id(child)
        if not shape_id:
            raise TemplateStructureError(
                "Generated round-trip slide contains a top-level shape without id"
            )
        if shape_id in generated_shapes:
            raise TemplateStructureError(
                f"Generated round-trip slide repeats top-level shape id {shape_id}"
            )
        generated_order.append(shape_id)
        generated_shapes[shape_id] = child
        for descendant_id in _shape_ids_in_element(child):
            previous = generated_top_id_by_descendant.setdefault(
                descendant_id,
                shape_id,
            )
            if previous != shape_id:
                raise TemplateStructureError(
                    "Generated round-trip slide repeats a shape id across "
                    f"top-level objects: {descendant_id}"
                )

    source_ref_ids = _slide_ref_shape_ids(patch.source_ref_ids)
    edited_ids = _slide_ref_shape_ids(patch.edited_ref_ids)
    deleted_ids = _slide_ref_shape_ids(patch.deleted_ref_ids)
    unchanged_ids = source_ref_ids - edited_ids - deleted_ids
    affected_top_ids = {
        shape_id
        for shape_id, descendant_ids in source_descendant_ids.items()
        if descendant_ids & (edited_ids | deleted_ids)
    }
    nested_restore_ids = _roundtrip_nested_restore_ids(
        source_shapes,
        source_key_by_element,
        affected_top_ids,
        unchanged_ids,
        deleted_ids,
    )
    authored_top_ids = {
        shape_id
        for shape_id in source_shapes
        if shape_id in source_ref_ids or shape_id in affected_top_ids
    }
    deleted_top_ids = authored_top_ids & deleted_ids
    traced_ids = _trace_native_shape_ids(conversion_trace)
    generated_by_source_id: dict[str, ET.Element] = {}
    source_id_by_generated_top_id: dict[str, str] = {}
    for source_id in authored_top_ids - deleted_top_ids:
        candidate_ids = [source_id]
        candidate_ids.extend(traced_ids.get(f"shape-{source_id}", []))
        generated_top_ids = {
            generated_top_id_by_descendant[candidate_id]
            for candidate_id in candidate_ids
            if candidate_id in generated_top_id_by_descendant
        }
        if len(generated_top_ids) > 1:
            raise TemplateStructureError(
                "Generated round-trip source object maps to multiple top-level "
                f"shapes: {source_id}"
            )
        if not generated_top_ids:
            continue
        generated_top_id = next(iter(generated_top_ids))
        previous_source_id = source_id_by_generated_top_id.setdefault(
            generated_top_id,
            source_id,
        )
        if previous_source_id != source_id:
            raise TemplateStructureError(
                "Generated round-trip top-level shape maps to multiple source "
                f"objects: {previous_source_id}, {source_id}"
            )
        generated_by_source_id[source_id] = generated_shapes[generated_top_id]

    authored_order = [
        source_id_by_generated_top_id[generated_top_id]
        for generated_top_id in generated_order
        if generated_top_id in source_id_by_generated_top_id
    ]
    if len(authored_order) != len(set(authored_order)):
        raise TemplateStructureError(
            "Generated round-trip slide has an ambiguous authored shape order"
        )
    missing_authored = (
        authored_top_ids - deleted_top_ids - set(authored_order)
    )
    if missing_authored:
        raise TemplateStructureError(
            "Edited round-trip source object did not produce a DrawingML shape: "
            + ", ".join(sorted(missing_authored, key=_roundtrip_shape_sort_key))
        )

    selected_by_id: dict[str, tuple[ET.Element, bool]] = {}
    for shape_id in authored_order:
        if shape_id in affected_top_ids:
            selected_by_id[shape_id] = (
                generated_by_source_id[shape_id],
                True,
            )
        else:
            selected_by_id[shape_id] = (source_shapes[shape_id], False)
    extra_generated = [
        generated_shapes[shape_id]
        for shape_id in generated_order
        if shape_id not in source_id_by_generated_top_id
    ]

    authored_iterator = iter(authored_order)
    merged_children: list[tuple[ET.Element, bool]] = []
    final_authored_slot = 0
    for child in list(source_tree):
        shape_id = (
            source_key_by_element.get(id(child))
            if child.tag in _TOP_LEVEL_SHAPE_TAGS
            else None
        )
        if shape_id not in authored_top_ids:
            merged_children.append((child, False))
            continue
        next_shape_id = next(authored_iterator, None)
        if next_shape_id is not None:
            merged_children.append(selected_by_id[next_shape_id])
        final_authored_slot = len(merged_children)
    if next(authored_iterator, None) is not None:
        raise TemplateStructureError(
            "Generated authored shape roster exceeds the source overlay slots"
        )
    if not authored_top_ids:
        final_authored_slot = next(
            (
                index
                for index, (element, _generated) in enumerate(merged_children)
                if element.tag in _TOP_LEVEL_SHAPE_TAGS
            ),
            len(merged_children),
        )
    merged_children[final_authored_slot:final_authored_slot] = [
        (element, True)
        for element in extra_generated
    ]

    generated_nodes = [
        element for element, generated in merged_children if generated
    ]
    retained_source_nodes = [
        element for element, generated in merged_children if not generated
    ]
    required_generated_relationship_ids = {
        relationship_id
        for element in generated_nodes
        for relationship_id in _relationship_ids_in_shape(element)
    }
    if patch.transition_changed:
        for carrier in transition_carriers(generated_root):
            required_generated_relationship_ids.update(
                _relationship_ids_in_shape(carrier)
            )
    if patch.animation_changed:
        motion_node = generated_root.find(f"{{{PML_NS}}}timing")
        if motion_node is not None:
            required_generated_relationship_ids.update(
                _relationship_ids_in_shape(motion_node)
            )
    _merge_roundtrip_relationships(
        source_rels,
        generated_rels_path,
        generated_root,
        required_generated_relationship_ids,
    )
    _restore_roundtrip_nested_source_shapes(
        generated_nodes,
        all_source_shapes,
        nested_restore_ids,
        traced_ids,
    )
    if patch.transition_changed or patch.animation_changed:
        _replace_roundtrip_motion(
            source_root,
            generated_root,
            transition_changed=patch.transition_changed,
            transition_replaced=patch.transition_replaced,
            animation_changed=patch.animation_changed,
        )
    restored_shape_id_map = {
        generated_top_id: source_id
        for generated_top_id, source_id in source_id_by_generated_top_id.items()
        if source_id not in affected_top_ids
        and generated_top_id != source_id
    }
    if patch.animation_changed:
        _rewrite_roundtrip_timing_shape_ids(
            source_root,
            restored_shape_id_map,
        )
    renumbered_shape_id_map = _renumber_roundtrip_generated_shapes(
        generated_nodes,
        retained_source_nodes,
        source_root,
        rewrite_timing=patch.animation_changed,
    )

    for child in list(source_tree):
        source_tree.remove(child)
    for child, _generated in merged_children:
        source_tree.append(child)

    visible_shape_ids = {
        value
        for element in source_tree
        for value in _shape_ids_in_element(element)
    }
    if not patch.animation_changed:
        missing_timing_targets = _timing_shape_ids(source_root) - visible_shape_ids
        if missing_timing_targets:
            raise TemplateStructureError(
                "Edited slide removed source animation target(s); update "
                "animations.json explicitly: "
                + ", ".join(sorted(missing_timing_targets, key=int))
            )
    serialized_slide = serialize_source_xml(source_root, source_slide)
    if not patch.transition_changed:
        serialized_slide = _preserve_roundtrip_transition_markup(
            serialized_slide,
            source_slide,
        )
    if not patch.animation_changed:
        serialized_slide = _preserve_roundtrip_timing_markup(
            serialized_slide,
            source_slide,
        )
    generated_slide_path.write_bytes(serialized_slide)
    combined_id_map = {
        **restored_shape_id_map,
        **renumbered_shape_id_map,
    }
    return {
        int(old_id): int(new_id)
        for old_id, new_id in combined_id_map.items()
        if old_id.isdigit() and new_id.isdigit()
    }


def _remove_relationships_by_type(rels_path: Path, rel_type: str) -> int:
    """Remove every relationship of one type and return the removed count."""
    tree = ET.parse(rels_path)
    root = tree.getroot()
    removed = 0
    for relationship in list(root):
        if _relationship_attrs(relationship).get("Type") != rel_type:
            continue
        root.remove(relationship)
        removed += 1
    if removed:
        _write_xml_tree(rels_path, tree)
    return removed


def _apply_slide_notes(
    extract_dir: Path,
    rels_path: Path,
    slide_num: int,
    notes_content: str,
    primary_language: str | None,
    *,
    enable_notes: bool,
) -> _NotesMasterReference | None:
    """Replace one slide's notes relationship and return its notes master."""
    _remove_relationships_by_type(rels_path, NOTES_SLIDE_REL_TYPE)
    if not enable_notes:
        return None
    notes_text = markdown_to_plain_text(notes_content) if notes_content else ''
    if not notes_text:
        return None

    notes_master = _ensure_notes_master(extract_dir, primary_language)
    notes_slides_dir = extract_dir / 'ppt' / 'notesSlides'
    notes_slides_dir.mkdir(exist_ok=True)
    notes_xml_path = notes_slides_dir / f'notesSlide{slide_num}.xml'
    notes_xml_path.write_text(
        create_notes_slide_xml(slide_num, notes_text, primary_language),
        encoding='utf-8',
    )
    notes_rels_dir = notes_slides_dir / '_rels'
    notes_rels_dir.mkdir(exist_ok=True)
    notes_rels_path = notes_rels_dir / f'notesSlide{slide_num}.xml.rels'
    notes_rels_path.write_text(
        create_notes_slide_rels_xml(
            slide_num,
            posixpath.relpath(notes_master.package_part, 'ppt/notesSlides'),
        ),
        encoding='utf-8',
    )
    _ensure_relationship(
        rels_path,
        NOTES_SLIDE_REL_TYPE,
        f'../notesSlides/notesSlide{slide_num}.xml',
    )
    return notes_master


def _roundtrip_resource_payloads(
    resource_root: Path | None,
    resources: tuple[WorkspaceResourceSpec, ...],
) -> dict[str, bytes]:
    """Read changed, validated round-trip resources by package part."""
    changed = [resource for resource in resources if resource.changed]
    if not changed:
        return {}
    if resource_root is None:
        raise TemplateStructureError(
            "Round-trip resource reinjection requires an explicit project root"
        )
    workspace = resource_root.resolve()
    payloads: dict[str, bytes] = {}
    for resource in changed:
        source = (workspace / resource.workspace_path).resolve()
        try:
            source.relative_to(workspace)
        except ValueError as exc:
            raise TemplateStructureError(
                "Round-trip resource resolves outside the project: "
                f"{resource.workspace_path.as_posix()}"
            ) from exc
        if not source.is_file():
            raise TemplateStructureError(
                "Materialized round-trip resource is missing: "
                f"{resource.workspace_path.as_posix()}"
            )
        payload = source.read_bytes()
        current_sha256 = hashlib.sha256(payload).hexdigest()
        if current_sha256 != resource.current_sha256:
            raise TemplateStructureError(
                "Round-trip resource changed after validation: "
                f"{resource.workspace_path.as_posix()}"
            )
        payloads[resource.package_part] = payload
    return payloads


def _reinject_roundtrip_resources(
    extract_dir: Path,
    resource_root: Path | None,
    resources: tuple[WorkspaceResourceSpec, ...],
) -> int:
    """Write changed resource bytes back to their exact source package parts."""
    payloads = _roundtrip_resource_payloads(resource_root, resources)
    written = 0
    for package_part, payload in payloads.items():
        target = extract_dir.joinpath(*PurePosixPath(package_part).parts)
        if not target.is_file():
            raise TemplateStructureError(
                "Round-trip source package part is missing: "
                f"{package_part}"
            )
        target.write_bytes(payload)
        written += 1
    return written


COVER_LAYOUT_NAME = "Cover"


def _next_layout_part_number(extract_dir: Path) -> int:
    layouts_dir = extract_dir / "ppt" / "slideLayouts"
    numbers = [
        int(match.group(1))
        for path in layouts_dir.glob("slideLayout*.xml")
        if (match := re.fullmatch(r"slideLayout(\d+)\.xml", path.name))
    ]
    return max(numbers, default=0) + 1


def _next_slide_layout_id(extract_dir: Path) -> int:
    """Return a package-wide unused id for a new sldLayoutId entry."""
    ids: list[int] = []
    masters_dir = extract_dir / "ppt" / "slideMasters"
    for master_path in sorted(masters_dir.glob("slideMaster*.xml")):
        master_root = ET.parse(master_path).getroot()
        ids.extend(
            int(entry.attrib["id"])
            for entry in master_root.findall(
                f"{{{PML_NS}}}sldLayoutIdLst/{{{PML_NS}}}sldLayoutId"
            )
            if entry.attrib.get("id", "").isdigit()
        )
    presentation_path = extract_dir / "ppt" / "presentation.xml"
    if presentation_path.exists():
        ids.extend(
            int(value)
            for value in re.findall(
                r'\bid="(\d{9,})"', presentation_path.read_text(encoding="utf-8")
            )
        )
    next_id = max([*ids, 2147483648]) + 1
    if next_id > OOXML_UINT32_MAX:
        raise TemplateStructureError(
            "Cannot register another Slide Layout because the OOXML UInt32 "
            "identifier range is exhausted"
        )
    return next_id


def _create_cover_layout(extract_dir: Path, master_part: str, base_layout_part: str) -> str:
    """Clone a layout into a Cover layout that hides master shapes.

    Returns the new layout target relative to slide parts.
    """
    base_layout_path = extract_dir / base_layout_part
    layout_xml = base_layout_path.read_text(encoding="utf-8")

    root_match = re.search(r"<p:sldLayout\b[^>]*>", layout_xml)
    if not root_match:
        raise RuntimeError(f"Slide layout has no p:sldLayout root: {base_layout_part}")
    root_tag = root_match.group(0)
    if "showMasterSp=" in root_tag:
        new_root_tag = re.sub(r'showMasterSp="[^"]*"', 'showMasterSp="0"', root_tag)
    else:
        new_root_tag = root_tag[:-1] + ' showMasterSp="0">'
    layout_xml = layout_xml.replace(root_tag, new_root_tag, 1)
    layout_xml = re.sub(
        r"(<p:cSld\b[^>]*?)\s+name=\"[^\"]*\"",
        rf'\g<1> name="{COVER_LAYOUT_NAME}"',
        layout_xml,
        count=1,
    )

    layout_num = _next_layout_part_number(extract_dir)
    new_layout_part = f"ppt/slideLayouts/slideLayout{layout_num}.xml"
    new_layout_path = extract_dir / new_layout_part
    new_layout_path.write_text(layout_xml, encoding="utf-8")

    base_rels_path = _relationships_path_for_part(extract_dir, base_layout_part)
    new_rels_path = _relationships_path_for_part(extract_dir, new_layout_part)
    new_rels_path.parent.mkdir(exist_ok=True)
    new_rels_path.write_text(
        base_rels_path.read_text(encoding="utf-8"), encoding="utf-8"
    )

    master_path = extract_dir / master_part
    master_rels_path = _relationships_path_for_part(extract_dir, master_part)
    layout_target = posixpath.relpath(
        new_layout_part, posixpath.dirname(master_part)
    )
    rel_id = _append_relationship(master_rels_path, SLIDE_LAYOUT_REL_TYPE, layout_target)

    master_xml = master_path.read_text(encoding="utf-8")
    layout_id = _next_slide_layout_id(extract_dir)
    entry = f'<p:sldLayoutId id="{layout_id}" r:id="{rel_id}"/>'
    if "</p:sldLayoutIdLst>" not in master_xml:
        raise RuntimeError(f"Slide master has no sldLayoutIdLst: {master_part}")
    master_path.write_text(
        master_xml.replace("</p:sldLayoutIdLst>", f"{entry}</p:sldLayoutIdLst>", 1),
        encoding="utf-8",
    )

    content_types_path = extract_dir / "[Content_Types].xml"
    content_types_path.write_text(
        _add_content_type_override(
            content_types_path.read_text(encoding="utf-8"),
            new_layout_part,
            "application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml",
        ),
        encoding="utf-8",
    )
    return posixpath.relpath(new_layout_part, "ppt/slides")


def _create_custom_layout(
    extract_dir: Path,
    master_part: str,
    base_layout_part: str,
    layout_name: str,
    *,
    show_master_shapes: bool = True,
) -> tuple[str, str]:
    """Clone a clean custom layout and register it under its slide master.

    Returns ``(slide_relationship_target, package_part)``.
    """
    base_layout_path = extract_dir / base_layout_part
    tree = ET.parse(base_layout_path)
    root = tree.getroot()
    _reseed_p14_creation_id(root)
    root.set("type", "cust")
    root.set("preserve", "1")
    root.set("showMasterSp", "1" if show_master_shapes else "0")

    c_sld = root.find(f"{{{PML_NS}}}cSld")
    if c_sld is None:
        raise RuntimeError(f"Slide layout has no p:cSld: {base_layout_part}")
    c_sld.set("name", layout_name)
    sp_tree = c_sld.find(f"{{{PML_NS}}}spTree")
    if sp_tree is None:
        raise RuntimeError(f"Slide layout has no p:spTree: {base_layout_part}")
    for child in list(sp_tree):
        if child.tag in _TOP_LEVEL_SHAPE_TAGS:
            sp_tree.remove(child)

    layout_num = _next_layout_part_number(extract_dir)
    new_layout_part = f"ppt/slideLayouts/slideLayout{layout_num}.xml"
    new_layout_path = extract_dir / new_layout_part
    _write_xml_tree(new_layout_path, tree)

    base_rels_path = _relationships_path_for_part(extract_dir, base_layout_part)
    new_rels_path = _relationships_path_for_part(extract_dir, new_layout_part)
    new_rels_path.parent.mkdir(exist_ok=True)
    new_rels_path.write_text(
        base_rels_path.read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    master_target = posixpath.relpath(
        master_part,
        posixpath.dirname(new_layout_part),
    )
    rels_content = new_rels_path.read_text(encoding="utf-8")
    master_rel_ids = [
        rel_id
        for rel_id, attrs in _read_relationships(new_rels_path).items()
        if attrs.get("Type") == SLIDE_MASTER_REL_TYPE
    ]
    if len(master_rel_ids) != 1:
        raise RuntimeError(
            f"Cloned slide layout must have one Master relationship: {new_layout_part}"
        )
    master_rel_id = master_rel_ids[0]
    master_rel_pattern = re.compile(
        rf'(<Relationship\b[^>]*\bId="{re.escape(master_rel_id)}"'
        rf'[^>]*\bTarget=")[^"]*(")'
    )
    rels_content, replaced = master_rel_pattern.subn(
        rf"\g<1>{master_target}\g<2>",
        rels_content,
        count=1,
    )
    if replaced != 1:
        raise RuntimeError(
            f"Could not retarget cloned Layout to Master {master_part}"
        )
    new_rels_path.write_text(rels_content, encoding="utf-8")

    master_path = extract_dir / master_part
    master_rels_path = _relationships_path_for_part(extract_dir, master_part)
    layout_target = posixpath.relpath(new_layout_part, posixpath.dirname(master_part))
    rel_id = _append_relationship(master_rels_path, SLIDE_LAYOUT_REL_TYPE, layout_target)
    master_tree = ET.parse(master_path)
    master_root = master_tree.getroot()
    layout_list = master_root.find(f"{{{PML_NS}}}sldLayoutIdLst")
    if layout_list is None:
        raise RuntimeError(f"Slide master has no sldLayoutIdLst: {master_part}")
    layout_id = _next_slide_layout_id(extract_dir)
    ET.SubElement(
        layout_list,
        f"{{{PML_NS}}}sldLayoutId",
        {"id": str(layout_id), f"{{{REL_NS}}}id": rel_id},
    )
    _write_xml_tree(master_path, master_tree)

    content_types_path = extract_dir / "[Content_Types].xml"
    content_types_path.write_text(
        _add_content_type_override(
            content_types_path.read_text(encoding="utf-8"),
            new_layout_part,
            "application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml",
        ),
        encoding="utf-8",
    )
    return posixpath.relpath(new_layout_part, "ppt/slides"), new_layout_part


def _set_master_picker_name(master_path: Path, master_name: str) -> None:
    """Set the visible PowerPoint Master name on its common slide data."""
    tree = ET.parse(master_path)
    c_sld = tree.getroot().find(f"{{{PML_NS}}}cSld")
    if c_sld is None:
        raise RuntimeError(f"Slide master has no p:cSld: {master_path}")
    c_sld.set("name", master_name)
    _write_xml_tree(master_path, tree)


def _reseed_p14_creation_id(root: ET.Element) -> None:
    """Give a cloned Slide/Master/Layout part a fresh PowerPoint creation id."""
    c_sld = root.find(f"{{{PML_NS}}}cSld")
    if c_sld is None:
        return
    creation_ids = c_sld.findall(
        f"{{{PML_NS}}}extLst/{{{PML_NS}}}ext/"
        f"{{{P14_NS}}}creationId"
    )
    for creation_id in creation_ids:
        value = 0
        while value == 0:
            value = uuid.uuid4().int & OOXML_UINT32_MAX
        creation_id.set("val", str(value))


def _next_master_part_number(extract_dir: Path) -> int:
    numbers = [
        int(match.group(1))
        for path in (extract_dir / "ppt" / "slideMasters").glob("slideMaster*.xml")
        if (match := re.fullmatch(r"slideMaster(\d+)\.xml", path.name))
    ]
    return max(numbers, default=0) + 1


def _next_theme_part_number(extract_dir: Path) -> int:
    numbers = [
        int(match.group(1))
        for path in (extract_dir / "ppt" / "theme").glob("theme*.xml")
        if (match := re.fullmatch(r"theme(\d+)\.xml", path.name))
    ]
    return max(numbers, default=0) + 1


def _clone_master_theme(
    extract_dir: Path,
    source_master_part: str,
    master_part: str,
    master_name: str,
) -> str:
    """Give a cloned Slide Master its own Theme package part."""
    source_master_rels = _relationships_path_for_part(
        extract_dir,
        source_master_part,
    )
    theme_targets = [
        attrs["Target"]
        for attrs in _read_relationships(source_master_rels).values()
        if attrs.get("Type") == THEME_REL_TYPE and attrs.get("Target")
    ]
    if len(theme_targets) != 1:
        raise RuntimeError(
            "Source Slide Master must have one Theme relationship: "
            f"{source_master_part}"
        )
    source_theme_part = _resolve_package_target(
        source_master_part,
        theme_targets[0],
    )
    source_theme_path = extract_dir / source_theme_part
    if not source_theme_path.exists():
        raise RuntimeError(
            f"Slide Master Theme part is missing: {source_theme_part}"
        )

    theme_num = _next_theme_part_number(extract_dir)
    theme_part = f"ppt/theme/theme{theme_num}.xml"
    theme_path = extract_dir / theme_part
    shutil.copyfile(source_theme_path, theme_path)
    theme_tree = ET.parse(theme_path)
    theme_tree.getroot().set("name", f"{master_name} Theme")
    _write_xml_tree(theme_path, theme_tree)

    source_theme_rels = _relationships_path_for_part(
        extract_dir,
        source_theme_part,
    )
    if source_theme_rels.exists():
        theme_rels = _relationships_path_for_part(extract_dir, theme_part)
        theme_rels.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source_theme_rels, theme_rels)

    master_rels = _relationships_path_for_part(extract_dir, master_part)
    theme_relationships = [
        (rel_id, attrs)
        for rel_id, attrs in _read_relationships(master_rels).items()
        if attrs.get("Type") == THEME_REL_TYPE
    ]
    if len(theme_relationships) != 1:
        raise RuntimeError(
            f"Cloned Slide Master must have one Theme relationship: {master_part}"
        )
    theme_rel_id, _attrs = theme_relationships[0]
    theme_target = posixpath.relpath(theme_part, posixpath.dirname(master_part))
    rels_content = master_rels.read_text(encoding="utf-8")
    theme_rel_pattern = re.compile(
        rf'(<Relationship\b[^>]*\bId="{re.escape(theme_rel_id)}"'
        rf'[^>]*\bTarget=")[^"]*(")'
    )
    rels_content, replaced = theme_rel_pattern.subn(
        rf"\g<1>{theme_target}\g<2>",
        rels_content,
        count=1,
    )
    if replaced != 1:
        raise RuntimeError(
            f"Could not retarget cloned Slide Master Theme: {master_part}"
        )
    master_rels.write_text(rels_content, encoding="utf-8")

    content_types_path = extract_dir / "[Content_Types].xml"
    content_types_path.write_text(
        _add_content_type_override(
            content_types_path.read_text(encoding="utf-8"),
            theme_part,
            THEME_CONTENT_TYPE,
        ),
        encoding="utf-8",
    )
    return theme_part


def _clone_structured_master(
    extract_dir: Path,
    source_master_part: str,
    master_name: str,
) -> str:
    """Clone a clean Master part and register it with the Presentation."""
    master_num = _next_master_part_number(extract_dir)
    master_part = f"ppt/slideMasters/slideMaster{master_num}.xml"
    master_path = extract_dir / master_part
    source_master_path = extract_dir / source_master_part
    shutil.copyfile(source_master_path, master_path)

    tree = ET.parse(master_path)
    root = tree.getroot()
    _reseed_p14_creation_id(root)
    c_sld = root.find(f"{{{PML_NS}}}cSld")
    if c_sld is None:
        raise RuntimeError(f"Slide master has no p:cSld: {source_master_part}")
    c_sld.set("name", master_name)
    layout_list = root.find(f"{{{PML_NS}}}sldLayoutIdLst")
    if layout_list is None:
        raise RuntimeError(
            f"Slide master has no p:sldLayoutIdLst: {source_master_part}"
        )
    for entry in list(layout_list):
        layout_list.remove(entry)
    _write_xml_tree(master_path, tree)

    source_rels = _relationships_path_for_part(extract_dir, source_master_part)
    master_rels = _relationships_path_for_part(extract_dir, master_part)
    master_rels.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source_rels, master_rels)
    for rel_id, attrs in tuple(_read_relationships(master_rels).items()):
        if attrs.get("Type") == SLIDE_LAYOUT_REL_TYPE:
            _remove_relationship(master_rels, rel_id)
    _clone_master_theme(
        extract_dir,
        source_master_part,
        master_part,
        master_name,
    )

    presentation_rels = extract_dir / "ppt" / "_rels" / "presentation.xml.rels"
    relationship_target = posixpath.relpath(master_part, "ppt")
    relationship_id = _append_relationship(
        presentation_rels,
        SLIDE_MASTER_REL_TYPE,
        relationship_target,
    )
    presentation_path = extract_dir / "ppt" / "presentation.xml"
    presentation_xml = presentation_path.read_text(encoding="utf-8")
    master_ids = [
        int(value)
        for value in re.findall(r'<p:sldMasterId\b[^>]*\bid="(\d+)"', presentation_xml)
    ]
    master_id = max(master_ids, default=(1 << 31) - 1) + 1
    if master_id > OOXML_UINT32_MAX:
        raise TemplateStructureError("Presentation Master id exceeds OOXML UInt32")
    entry = f'<p:sldMasterId id="{master_id}" r:id="{relationship_id}"/>'
    if "</p:sldMasterIdLst>" not in presentation_xml:
        raise RuntimeError("presentation.xml has no p:sldMasterIdLst")
    presentation_path.write_text(
        presentation_xml.replace(
            "</p:sldMasterIdLst>",
            f"{entry}</p:sldMasterIdLst>",
            1,
        ),
        encoding="utf-8",
    )

    content_types_path = extract_dir / "[Content_Types].xml"
    content_types_path.write_text(
        _add_content_type_override(
            content_types_path.read_text(encoding="utf-8"),
            master_part,
            "application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml",
        ),
        encoding="utf-8",
    )
    return master_part


def _assign_structured_masters(
    extract_dir: Path,
    structure: PptxStructureContext,
    specs: list[TemplateSlideSpec],
) -> dict[str, str]:
    """Create one registered package Master for each explicit SVG Master key."""
    if not specs:
        raise TemplateStructureError("Structured export requires at least one slide")
    source_master = structure.slide_master_part(specs[0].slide_num)
    masters: dict[str, str] = {}
    for spec in specs:
        existing_part = masters.get(spec.master_key)
        if existing_part is None:
            if not masters:
                existing_part = source_master
                _set_master_picker_name(
                    extract_dir / existing_part,
                    spec.master_name,
                )
            else:
                existing_part = _clone_structured_master(
                    extract_dir,
                    source_master,
                    spec.master_name,
                )
            masters[spec.master_key] = existing_part
        structure.slide_master_parts[spec.slide_num] = existing_part
    return masters


def _clear_master_placeholder_shapes(master_path: Path) -> None:
    """Remove base-package placeholders before installing structured content."""
    tree = ET.parse(master_path)
    root = tree.getroot()
    sp_tree = root.find(f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree")
    if sp_tree is None:
        raise RuntimeError(f"Slide master has no p:spTree: {master_path}")
    for child in list(sp_tree):
        if child.find(f".//{{{PML_NS}}}ph") is not None:
            sp_tree.remove(child)
    _write_xml_tree(master_path, tree)


def _set_slide_layout_target(rels_path: Path, target: str) -> None:
    """Point a slide's layout relationship at a different layout part."""
    content = rels_path.read_text(encoding="utf-8")
    rel_id = None
    for existing_id, attrs in _read_relationships(rels_path).items():
        if attrs.get("Type") == SLIDE_LAYOUT_REL_TYPE:
            rel_id = existing_id
            break
    if rel_id is None:
        raise RuntimeError(f"No slide layout relationship in {rels_path}")
    pattern = re.compile(
        rf'(<Relationship\b[^>]*\bId="{re.escape(rel_id)}"[^>]*\bTarget=")[^"]*(")'
    )
    new_content, replaced = pattern.subn(rf"\g<1>{target}\g<2>", content, count=1)
    if not replaced:
        raise RuntimeError(f"Could not retarget layout relationship in {rels_path}")
    rels_path.write_text(new_content, encoding="utf-8")


_BASELINE_LAYOUT_ROLE_TOKENS = (
    ("Cover", frozenset({"cover", "frontcover"}), ("封面",)),
    (
        "Agenda",
        frozenset({"agenda", "contents", "outline", "toc"}),
        ("目录", "议程"),
    ),
    (
        "Section",
        frozenset({"chapter", "divider", "section", "transition"}),
        ("章节", "过渡页"),
    ),
    (
        "Closing",
        frozenset({"closing", "end", "ending", "qa", "thankyou", "thanks"}),
        ("封底", "结束", "结尾", "结语", "致谢", "谢谢"),
    ),
)


def _baseline_layout_role(svg_path: Path) -> str:
    """Use explicit page semantics, then fall back to conservative filename roles."""
    semantic_role = page_layout_name_from_svg(svg_path)
    if semantic_role:
        return semantic_role
    stem = svg_path.stem.casefold()
    tokens = {
        token
        for token in re.split(r"[^0-9a-z]+", stem)
        if token and not token.isdigit()
    }
    for role, english_tokens, cjk_tokens in _BASELINE_LAYOUT_ROLE_TOKENS:
        if tokens.intersection(english_tokens) or any(
            token in stem for token in cjk_tokens
        ):
            return role
    return "Content"


def _layout_identity(layout_path: Path) -> tuple[str, bool]:
    """Return a layout's picker name and master-shape visibility."""
    root = ET.parse(layout_path).getroot()
    c_sld = root.find(f"{{{PML_NS}}}cSld")
    name = c_sld.attrib.get("name", "") if c_sld is not None else ""
    return name, root.attrib.get("showMasterSp", "1") != "0"


def _shared_explicit_slide_background(
    extract_dir: Path,
    slide_nums: list[int],
) -> str | None:
    """Return one exact background only when every family slide carries it."""
    backgrounds: list[str] = []
    for slide_num in slide_nums:
        slide_path = extract_dir / "ppt" / "slides" / f"slide{slide_num}.xml"
        background = _extract_slide_background_xml(
            slide_path.read_text(encoding="utf-8")
        )
        if background is None:
            return None
        backgrounds.append(background)
    if not backgrounds or len(set(backgrounds)) != 1:
        return None
    return backgrounds[0]


def _extract_baseline_layout_families(
    extract_dir: Path,
    structure: PptxStructureContext,
    svg_files: list[Path],
    *,
    verbose: bool = False,
) -> int:
    """Build conservative post-generation layout families for a free SVG deck."""
    families: dict[tuple[str, str, bool], list[int]] = {}
    base_layouts: dict[tuple[str, str, bool], str] = {}
    for slide_num, svg_path in enumerate(svg_files, 1):
        rels_path = (
            extract_dir
            / "ppt"
            / "slides"
            / "_rels"
            / f"slide{slide_num}.xml.rels"
        )
        layout_target = _find_relationship_target(rels_path, SLIDE_LAYOUT_REL_TYPE)
        if not layout_target:
            raise RuntimeError(f"Slide {slide_num} has no slide layout relationship")
        layout_part = _resolve_package_target(
            f"ppt/slides/slide{slide_num}.xml", layout_target
        )
        layout_name, show_master_shapes = _layout_identity(extract_dir / layout_part)
        role = _baseline_layout_role(svg_path)
        if role == "Content" and layout_name == COVER_LAYOUT_NAME:
            role = COVER_LAYOUT_NAME
        key = (
            structure.slide_master_part(slide_num),
            role,
            show_master_shapes,
        )
        families.setdefault(key, []).append(slide_num)
        base_layouts.setdefault(key, layout_part)

    created = 0
    lifted_backgrounds = 0
    role_counts: dict[tuple[str, str], int] = {}
    for key, slide_nums in families.items():
        master_part, role, show_master_shapes = key
        role_key = (master_part, role)
        role_counts[role_key] = role_counts.get(role_key, 0) + 1
        variant = role_counts[role_key]
        layout_name = role if variant == 1 else f"{role} {variant}"
        layout_target, layout_part = _create_custom_layout(
            extract_dir,
            master_part,
            base_layouts[key],
            layout_name,
            show_master_shapes=show_master_shapes,
        )

        background_xml = _shared_explicit_slide_background(
            extract_dir, slide_nums
        )
        if background_xml is not None:
            layout_path = extract_dir / layout_part
            layout_xml = layout_path.read_text(encoding="utf-8")
            updated_layout_xml = _put_background_on_part(layout_xml, background_xml)
            if updated_layout_xml is not None:
                layout_path.write_text(updated_layout_xml, encoding="utf-8")
                for slide_num in slide_nums:
                    slide_path = (
                        extract_dir / "ppt" / "slides" / f"slide{slide_num}.xml"
                    )
                    slide_path.write_text(
                        _remove_slide_background_xml(
                            slide_path.read_text(encoding="utf-8")
                        ),
                        encoding="utf-8",
                    )
                lifted_backgrounds += len(slide_nums)

        for slide_num in slide_nums:
            rels_path = (
                extract_dir
                / "ppt"
                / "slides"
                / "_rels"
                / f"slide{slide_num}.xml.rels"
            )
            _set_slide_layout_target(rels_path, layout_target)
        created += 1

    if verbose and created:
        print(
            "  Baseline layout families: "
            f"created {created} reusable layout(s), "
            f"lifted {lifted_backgrounds} slide background(s)"
        )
    return created


def _promote_common_chrome_shapes_to_layouts(
    extract_dir: Path,
    slide_count: int,
    conversion_traces: list[dict[str, Any]] | None,
    *,
    verbose: bool = False,
) -> int:
    """Promote exact leading chrome shared by every slide in one layout."""
    if not conversion_traces:
        return 0
    trace_by_slide = {
        int(trace.get("slide_num", 0)): trace
        for trace in conversion_traces
        if trace.get("slide_num") is not None
    }
    if len(trace_by_slide) < slide_count:
        return 0

    slides_by_layout: dict[str, list[int]] = {}
    for slide_num in range(1, slide_count + 1):
        rels_path = (
            extract_dir
            / "ppt"
            / "slides"
            / "_rels"
            / f"slide{slide_num}.xml.rels"
        )
        layout_target = _find_relationship_target(rels_path, SLIDE_LAYOUT_REL_TYPE)
        if not layout_target:
            raise RuntimeError(f"Slide {slide_num} has no slide layout relationship")
        layout_part = _resolve_package_target(
            f"ppt/slides/slide{slide_num}.xml",
            layout_target,
        )
        slides_by_layout.setdefault(layout_part, []).append(slide_num)

    promoted = 0
    promoted_roles = 0
    for layout_part, slide_nums in slides_by_layout.items():
        if len(slide_nums) < 2:
            continue
        slide_state: dict[int, dict[str, Any]] = {}
        for slide_num in slide_nums:
            slide_path = extract_dir / "ppt" / "slides" / f"slide{slide_num}.xml"
            rels_path = (
                extract_dir
                / "ppt"
                / "slides"
                / "_rels"
                / f"slide{slide_num}.xml.rels"
            )
            tree = ET.parse(slide_path)
            root = tree.getroot()
            slide_state[slide_num] = {
                "path": slide_path,
                "rels": _read_relationships(rels_path),
                "root": root,
                "shapes": _top_level_shapes_by_id(root),
                "timing_shape_ids": _timing_shape_ids(root),
                "tokens": _trace_chrome_shape_ids(trace_by_slide.get(slide_num)),
                "tree": tree,
            }

        common_tokens = set.intersection(*(
            set(state["tokens"])
            for state in slide_state.values()
        ))
        common_tokens.difference_update(_PAGE_NUMBER_TOKENS)
        candidates: dict[str, dict[int, str]] = {}
        for token in sorted(common_tokens):
            shape_ids_by_slide: dict[int, str] = {}
            canonical_shapes: set[bytes] = set()
            for slide_num in slide_nums:
                state = slide_state[slide_num]
                shape_ids = state["tokens"].get(token, [])
                if len(shape_ids) != 1:
                    break
                shape_id = shape_ids[0]
                shape = state["shapes"].get(shape_id)
                if shape is None or shape_id in state["timing_shape_ids"]:
                    break
                if not _shape_relationships_supported(shape, state["rels"]):
                    break
                shape_ids_by_slide[slide_num] = shape_id
                canonical_shapes.add(_canonical_shape_xml(shape, state["rels"]))
            if len(shape_ids_by_slide) == len(slide_nums) and len(canonical_shapes) == 1:
                candidates[token] = shape_ids_by_slide

        if not candidates:
            continue

        # Layout shapes render behind slide-local shapes. Keep visual z-order by
        # accepting only the identical leading prefix shared by every family page.
        token_by_shape_id = {
            slide_num: {
                shape_ids[slide_num]: token
                for token, shape_ids in candidates.items()
            }
            for slide_num in slide_nums
        }
        leading_orders: list[list[str]] = []
        for slide_num in slide_nums:
            order: list[str] = []
            for shape_id in slide_state[slide_num]["shapes"]:
                token = token_by_shape_id[slide_num].get(shape_id)
                if token is None:
                    break
                order.append(token)
            leading_orders.append(order)
        safe_tokens = list(leading_orders[0])
        for order in leading_orders[1:]:
            common_length = 0
            for expected, actual in zip(safe_tokens, order):
                if expected != actual:
                    break
                common_length += 1
            safe_tokens = safe_tokens[:common_length]
            if not safe_tokens:
                break
        if not safe_tokens:
            continue

        layout_path = extract_dir / layout_part
        layout_rels_path = _relationships_path_for_part(extract_dir, layout_part)
        for token in safe_tokens:
            shape_ids_by_slide = candidates[token]
            first_slide = slide_nums[0]
            first_state = slide_state[first_slide]
            shape = first_state["shapes"][shape_ids_by_slide[first_slide]]
            layout_shape = _copy_shape_relationships_to_part(
                shape,
                first_state["rels"],
                layout_rels_path,
            )
            _append_shape_to_part(layout_path, layout_shape)
            promoted_roles += 1
            for slide_num, shape_id in shape_ids_by_slide.items():
                state = slide_state[slide_num]
                shape_to_remove = state["shapes"].get(shape_id)
                sp_tree = state["root"].find(
                    f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree"
                )
                if sp_tree is not None and shape_to_remove is not None:
                    sp_tree.remove(shape_to_remove)
                    promoted += 1

        for slide_num in slide_nums:
            state = slide_state[slide_num]
            _write_xml_tree(state["path"], state["tree"])

    if verbose and promoted:
        print(
            "  Baseline layout chrome: "
            f"promoted {promoted} slide shape(s) across "
            f"{promoted_roles} shared object(s)"
        )
    return promoted


_TEMPLATE_PLACEHOLDER_PROMPTS = {
    "title": "Click to add title",
    "subtitle": "Click to add subtitle",
    "body": "Click to add text",
    "picture": "Click to add picture",
    "chart": "Click to add chart",
    "table": "Click to add table",
    "object": "Click to add content",
    "media": "Click to add media",
    "date": "Date",
    "footer": "Footer",
}

_PARAGRAPH_BULLET_CHOICE_TAGS = {
    f"{{{DML_NS}}}buNone",
    f"{{{DML_NS}}}buAutoNum",
    f"{{{DML_NS}}}buChar",
    f"{{{DML_NS}}}buBlip",
}
_PARAGRAPH_PROPERTIES_TRAILING_TAGS = {
    f"{{{DML_NS}}}tabLst",
    f"{{{DML_NS}}}defRPr",
    f"{{{DML_NS}}}extLst",
}


def _template_runtime_slides(
    extract_dir: Path,
    specs: list[TemplateSlideSpec],
    conversion_traces: list[dict[str, Any]] | None,
) -> list[_TemplateRuntimeSlide]:
    """Load slide XML state and join it with SVG-to-shape trace ids."""
    if not conversion_traces:
        raise TemplateStructureError(
            "Explicit Layout export requires native conversion traces for every slide"
        )
    trace_by_slide = {
        int(trace.get("slide_num", 0)): trace
        for trace in conversion_traces
        if trace.get("slide_num") is not None
    }
    states: list[_TemplateRuntimeSlide] = []
    for spec in specs:
        trace = trace_by_slide.get(spec.slide_num)
        if trace is None:
            raise TemplateStructureError(
                f"{spec.svg_path.name}: missing native conversion trace"
            )
        slide_path = extract_dir / "ppt" / "slides" / f"slide{spec.slide_num}.xml"
        rels_path = (
            extract_dir
            / "ppt"
            / "slides"
            / "_rels"
            / f"slide{spec.slide_num}.xml.rels"
        )
        tree = ET.parse(slide_path)
        root = tree.getroot()
        states.append(_TemplateRuntimeSlide(
            spec=spec,
            slide_path=slide_path,
            rels_path=rels_path,
            tree=tree,
            root=root,
            rels=_read_relationships(rels_path),
            shapes=_top_level_shapes_by_id(root),
            shape_ids_by_svg_id=_trace_native_shape_ids(trace),
        ))
    return states


def _template_shape_for_item(
    state: _TemplateRuntimeSlide,
    item: TemplateElementSpec,
) -> ET.Element | None:
    """Resolve one metadata item to its generated top-level DrawingML shape."""
    shape_ids = [
        shape_id
        for shape_id in state.shape_ids_by_svg_id.get(item.element_id, [])
        if shape_id in state.shapes
    ]
    if len(shape_ids) == 1:
        return state.shapes[shape_ids[0]]
    if not shape_ids and item.layer and item.order == 0 and item.tag in {"rect", "g"}:
        if _extract_slide_background_xml(
            state.slide_path.read_text(encoding="utf-8")
        ):
            return None
    if not shape_ids:
        text_hint = (
            "; multiline text placeholders require a single-frame text mode "
            "and cannot use --no-merge"
            if item.placeholder and item.placeholder_carrier_tag == "text"
            else ""
        )
        raise TemplateStructureError(
            f"{state.spec.svg_path.name}: metadata element {item.element_id!r} "
            f"did not produce one top-level native shape{text_hint}"
        )
    raise TemplateStructureError(
        f"{state.spec.svg_path.name}: metadata element {item.element_id!r} "
        f"resolved to {len(shape_ids)} top-level shapes; use one direct SVG element"
    )


def _shape_transform(shape: ET.Element) -> ET.Element | None:
    """Return the direct DrawingML transform for one top-level shape."""
    paths = {
        f"{{{PML_NS}}}sp": f"{{{PML_NS}}}spPr/{{{DML_NS}}}xfrm",
        f"{{{PML_NS}}}pic": f"{{{PML_NS}}}spPr/{{{DML_NS}}}xfrm",
        f"{{{PML_NS}}}cxnSp": f"{{{PML_NS}}}spPr/{{{DML_NS}}}xfrm",
        f"{{{PML_NS}}}graphicFrame": f"{{{PML_NS}}}xfrm",
        f"{{{PML_NS}}}grpSp": f"{{{PML_NS}}}grpSpPr/{{{DML_NS}}}xfrm",
    }
    path = paths.get(shape.tag)
    return shape.find(path) if path is not None else None


def _int_attr(elem: ET.Element, name: str, context: str) -> int:
    try:
        return int(elem.attrib[name])
    except (KeyError, ValueError) as exc:
        raise TemplateStructureError(f"{context} has invalid {name!r}") from exc


def _flatten_group_transform(
    group: ET.Element,
    carrier: ET.Element,
    *,
    context: str,
) -> None:
    """Map a single group's child transform into slide coordinates."""
    group_xfrm = _shape_transform(group)
    carrier_xfrm = _shape_transform(carrier)
    if group_xfrm is None or carrier_xfrm is None:
        raise TemplateStructureError(
            f"{context} cannot be unwrapped because its DrawingML transform is missing"
        )
    if any(group_xfrm.get(name) is not None for name in ("rot", "flipH", "flipV")):
        raise TemplateStructureError(
            f"{context} wrapper carries an unsupported group rotation or flip"
        )
    group_off = group_xfrm.find(f"{{{DML_NS}}}off")
    group_ext = group_xfrm.find(f"{{{DML_NS}}}ext")
    child_off = group_xfrm.find(f"{{{DML_NS}}}chOff")
    child_ext = group_xfrm.find(f"{{{DML_NS}}}chExt")
    carrier_off = carrier_xfrm.find(f"{{{DML_NS}}}off")
    carrier_ext = carrier_xfrm.find(f"{{{DML_NS}}}ext")
    if any(value is None for value in (
        group_off,
        group_ext,
        child_off,
        child_ext,
        carrier_off,
        carrier_ext,
    )):
        raise TemplateStructureError(
            f"{context} cannot be unwrapped because its group transform is incomplete"
        )
    child_width = _int_attr(child_ext, "cx", context)
    child_height = _int_attr(child_ext, "cy", context)
    if child_width <= 0 or child_height <= 0:
        raise TemplateStructureError(f"{context} group child extent must be positive")
    scale_x = _int_attr(group_ext, "cx", context) / child_width
    scale_y = _int_attr(group_ext, "cy", context) / child_height
    mapped_x = _int_attr(group_off, "x", context) + round(
        (_int_attr(carrier_off, "x", context) - _int_attr(child_off, "x", context))
        * scale_x
    )
    mapped_y = _int_attr(group_off, "y", context) + round(
        (_int_attr(carrier_off, "y", context) - _int_attr(child_off, "y", context))
        * scale_y
    )
    carrier_off.set("x", str(mapped_x))
    carrier_off.set("y", str(mapped_y))
    carrier_ext.set(
        "cx",
        str(round(_int_attr(carrier_ext, "cx", context) * scale_x)),
    )
    carrier_ext.set(
        "cy",
        str(round(_int_attr(carrier_ext, "cy", context) * scale_y)),
    )


def _unwrap_placeholder_carrier(
    state: _TemplateRuntimeSlide,
    item: TemplateElementSpec,
) -> ET.Element:
    """Remove one SVG-only slot wrapper and return its top-level carrier."""
    wrapper = _template_shape_for_item(state, item)
    if wrapper is None:
        raise TemplateStructureError(
            f"{state.spec.svg_path.name}: placeholder {item.element_id!r} cannot "
            "be a slide background"
        )
    if item.tag != "g" or wrapper.tag != f"{{{PML_NS}}}grpSp":
        return wrapper
    carriers = [
        child for child in wrapper if child.tag in _TOP_LEVEL_SHAPE_TAGS
    ]
    if len(carriers) != 1:
        raise TemplateStructureError(
            f"{state.spec.svg_path.name}: placeholder group {item.element_id!r} "
            f"converted to {len(carriers)} native children; expected one carrier"
        )
    carrier = carriers[0]
    _flatten_group_transform(
        wrapper,
        carrier,
        context=f"{state.spec.svg_path.name} placeholder {item.element_id!r}",
    )
    sp_tree = _slide_sp_tree(state)
    try:
        wrapper_index = list(sp_tree).index(wrapper)
    except ValueError as exc:
        raise TemplateStructureError(
            f"{state.spec.svg_path.name}: placeholder wrapper "
            f"{item.element_id!r} is not top-level"
        ) from exc
    wrapper.remove(carrier)
    sp_tree.remove(wrapper)
    sp_tree.insert(wrapper_index, carrier)
    wrapper_id = _shape_id(wrapper)
    carrier_id = _shape_id(carrier)
    if wrapper_id:
        state.shapes.pop(wrapper_id, None)
    if carrier_id:
        state.shapes[carrier_id] = carrier
    return carrier


def _slide_sp_tree(state: _TemplateRuntimeSlide) -> ET.Element:
    sp_tree = state.root.find(f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree")
    if sp_tree is None:
        raise RuntimeError(f"Slide has no p:spTree: {state.slide_path}")
    return sp_tree


def _append_shape_to_runtime_slide(
    state: _TemplateRuntimeSlide,
    elem: ET.Element,
) -> None:
    """Append a generated helper shape to one Slide with fresh object ids."""
    existing_ids = [
        int(cnv.attrib["id"])
        for cnv in state.root.iter(f"{{{PML_NS}}}cNvPr")
        if cnv.attrib.get("id", "").isdigit()
    ]
    clone = ET.fromstring(ET.tostring(elem, encoding="utf-8"))
    _renumber_shape_ids(clone, max(existing_ids, default=1) + 1)
    _slide_sp_tree(state).append(clone)


def _presentation_slide_size_emu(extract_dir: Path) -> tuple[int, int]:
    """Return the package slide size in EMU."""
    presentation_path = extract_dir / "ppt" / "presentation.xml"
    root = ET.parse(presentation_path).getroot()
    slide_size = root.find(f"{{{PML_NS}}}sldSz")
    if slide_size is None:
        raise TemplateStructureError("presentation.xml has no p:sldSz")
    try:
        width = int(slide_size.attrib["cx"])
        height = int(slide_size.attrib["cy"])
    except (KeyError, ValueError) as exc:
        raise TemplateStructureError("presentation.xml has an invalid p:sldSz") from exc
    if width <= 0 or height <= 0:
        raise TemplateStructureError("presentation.xml p:sldSz must be positive")
    return width, height


def _solid_background_xml_from_shape(
    shape: ET.Element,
    slide_size_emu: tuple[int, int],
) -> str | None:
    """Convert one exact full-slide solid rectangle shape into p:bg XML."""
    if shape.tag != f"{{{PML_NS}}}sp":
        return None
    if shape.find(f"{{{PML_NS}}}txBody") is not None:
        return None
    sp_pr = shape.find(f"{{{PML_NS}}}spPr")
    if sp_pr is None:
        return None
    xfrm = sp_pr.find(f"{{{DML_NS}}}xfrm")
    if xfrm is None or xfrm.attrib:
        return None
    off = xfrm.find(f"{{{DML_NS}}}off")
    ext = xfrm.find(f"{{{DML_NS}}}ext")
    if off is None or ext is None:
        return None
    try:
        bounds = (
            int(off.attrib["x"]),
            int(off.attrib["y"]),
            int(ext.attrib["cx"]),
            int(ext.attrib["cy"]),
        )
    except (KeyError, ValueError):
        return None
    if bounds != (0, 0, *slide_size_emu):
        return None

    geometry = sp_pr.find(f"{{{DML_NS}}}prstGeom")
    if geometry is None or geometry.attrib.get("prst") != "rect":
        return None
    solid_fill = sp_pr.find(f"{{{DML_NS}}}solidFill")
    if solid_fill is None:
        return None
    competing_fills = {
        f"{{{DML_NS}}}noFill",
        f"{{{DML_NS}}}gradFill",
        f"{{{DML_NS}}}blipFill",
        f"{{{DML_NS}}}pattFill",
        f"{{{DML_NS}}}grpFill",
    }
    if any(child.tag in competing_fills for child in sp_pr):
        return None
    line = sp_pr.find(f"{{{DML_NS}}}ln")
    if line is not None and line.find(f"{{{DML_NS}}}noFill") is None:
        return None
    for effect_tag in (f"{{{DML_NS}}}effectLst", f"{{{DML_NS}}}effectDag"):
        effect = sp_pr.find(effect_tag)
        if effect is not None and (effect.attrib or list(effect)):
            return None

    background = ET.Element(f"{{{PML_NS}}}bg")
    background_props = ET.SubElement(background, f"{{{PML_NS}}}bgPr")
    background_props.append(
        ET.fromstring(ET.tostring(solid_fill, encoding="utf-8"))
    )
    ET.SubElement(background_props, f"{{{DML_NS}}}effectLst")
    return ET.tostring(background, encoding="unicode")


def _remove_template_shape(
    state: _TemplateRuntimeSlide,
    shape: ET.Element,
) -> None:
    sp_tree = _slide_sp_tree(state)
    if shape not in list(sp_tree):
        raise TemplateStructureError(
            f"{state.spec.svg_path.name}: structure shape is not slide-local"
        )
    sp_tree.remove(shape)


def _move_template_background(
    states: list[_TemplateRuntimeSlide],
    target_path: Path,
) -> str:
    backgrounds = [
        _extract_slide_background_xml(state.slide_path.read_text(encoding="utf-8"))
        for state in states
    ]
    if not backgrounds or any(background is None for background in backgrounds):
        raise TemplateStructureError(
            "Template background metadata must resolve to an explicit background "
            "on every affected slide"
        )
    canonical_backgrounds = set()
    for background in backgrounds:
        if background is None:
            continue
        wrapper = ET.fromstring(
            f'<root xmlns:p="{PML_NS}" xmlns:a="{DML_NS}">{background}</root>'
        )
        canonical_backgrounds.add(
            ET.tostring(list(wrapper)[0], encoding="utf-8")
        )
    if len(canonical_backgrounds) != 1:
        slide_names = ", ".join(state.spec.svg_path.name for state in states)
        raise TemplateStructureError(
            f"Explicit template background differs across slides: {slide_names}"
        )
    background_xml = backgrounds[0]
    if background_xml is None:
        raise TemplateStructureError("Template background is unexpectedly empty")
    target_xml = target_path.read_text(encoding="utf-8")
    updated = _put_background_on_part(target_xml, background_xml)
    if updated is None:
        raise TemplateStructureError(
            f"Cannot install explicit background on {target_path.name}"
        )
    target_path.write_text(updated, encoding="utf-8")
    for state in states:
        c_sld = state.root.find(f"{{{PML_NS}}}cSld")
        background = (
            c_sld.find(f"{{{PML_NS}}}bg") if c_sld is not None else None
        )
        if c_sld is None or background is None:
            raise TemplateStructureError(
                f"{state.spec.svg_path.name}: explicit background disappeared "
                "during explicit Layout structure assembly"
            )
        c_sld.remove(background)
    return background_xml


def _move_template_solid_background_shapes(
    states: list[_TemplateRuntimeSlide],
    shapes: list[ET.Element],
    target_path: Path,
    slide_size_emu: tuple[int, int],
) -> str | None:
    """Move repeated full-slide solid rects into a master/layout p:bg."""
    backgrounds = [
        _solid_background_xml_from_shape(shape, slide_size_emu)
        for shape in shapes
    ]
    if not any(backgrounds):
        return None
    if any(background is None for background in backgrounds):
        raise TemplateStructureError(
            "A template background resolves to a full-slide solid rect on only "
            "some slides sharing the structure"
        )
    canonical = {background for background in backgrounds if background is not None}
    if len(canonical) != 1:
        slide_names = ", ".join(state.spec.svg_path.name for state in states)
        raise TemplateStructureError(
            f"Explicit template solid background differs across slides: {slide_names}"
        )
    background_xml = backgrounds[0]
    if background_xml is None:
        return None
    target_xml = target_path.read_text(encoding="utf-8")
    updated = _put_background_on_part(target_xml, background_xml)
    if updated is None:
        raise TemplateStructureError(
            f"Cannot install explicit solid background on {target_path.name}"
        )
    target_path.write_text(updated, encoding="utf-8")
    for state, shape in zip(states, shapes):
        _remove_template_shape(state, shape)
    return background_xml


def _set_slide_tree_background(
    state: _TemplateRuntimeSlide,
    background_xml: str,
) -> None:
    """Replace the slide tree's p:bg with explicit background XML."""
    c_sld = state.root.find(f"{{{PML_NS}}}cSld")
    if c_sld is None:
        raise TemplateStructureError(
            f"{state.spec.svg_path.name}: slide has no p:cSld"
        )
    for existing in list(c_sld):
        if existing.tag == f"{{{PML_NS}}}bg":
            c_sld.remove(existing)
    background = ET.fromstring(background_xml)
    sp_tree_tag = f"{{{PML_NS}}}spTree"
    insert_at = next(
        (index for index, child in enumerate(c_sld) if child.tag == sp_tree_tag),
        0,
    )
    c_sld.insert(insert_at, background)


def _apply_template_slide_backgrounds(
    states: list[_TemplateRuntimeSlide],
    slide_size_emu: tuple[int, int],
) -> dict[str, str]:
    """Compile one-page solid backgrounds into slide-level p:bg."""
    applied: dict[str, str] = {}
    for state in states:
        items = [
            item for item in state.spec.elements
            if item.layer == "slide" and item.is_background
        ]
        if not items:
            continue
        item = items[0]
        shape = _template_shape_for_item(state, item)
        if shape is None:
            c_sld = state.root.find(f"{{{PML_NS}}}cSld")
            background = (
                c_sld.find(f"{{{PML_NS}}}bg")
                if c_sld is not None
                else None
            )
            if background is None:
                raise TemplateStructureError(
                    f"{state.spec.svg_path.name}: slide background disappeared "
                    "during explicit Layout structure assembly"
                )
            applied[f"ppt/slides/slide{state.spec.slide_num}.xml"] = ET.tostring(
                background,
                encoding="unicode",
            )
            continue
        background_xml = _solid_background_xml_from_shape(shape, slide_size_emu)
        if background_xml is None:
            raise TemplateStructureError(
                f"{state.spec.svg_path.name}: {item.element_id!r} must remain an "
                "exact full-slide solid rectangle"
            )
        _remove_template_shape(state, shape)
        _set_slide_tree_background(state, background_xml)
        applied[f"ppt/slides/slide{state.spec.slide_num}.xml"] = background_xml
    return applied


def _move_template_static_shape(
    states: list[_TemplateRuntimeSlide],
    item: TemplateElementSpec,
    target_path: Path,
    target_rels_path: Path,
    slide_size_emu: tuple[int, int],
) -> str | None:
    shapes = [_template_shape_for_item(state, item) for state in states]
    if any(shape is None for shape in shapes):
        if not all(shape is None for shape in shapes):
            raise TemplateStructureError(
                f"{item.element_id}: structure item is a background on only some slides"
            )
        return _move_template_background(states, target_path)

    resolved_shapes = [shape for shape in shapes if shape is not None]
    if item.is_background:
        background_xml = _move_template_solid_background_shapes(
            states,
            resolved_shapes,
            target_path,
            slide_size_emu,
        )
        if background_xml is None:
            raise TemplateStructureError(
                f"{item.element_id!r} must compile to one exact p:bg payload"
            )
        return background_xml
    canonical = {
        _canonical_shape_xml(shape, state.rels)
        for state, shape in zip(states, resolved_shapes)
    }
    if len(canonical) != 1:
        slide_names = ", ".join(state.spec.svg_path.name for state in states)
        raise TemplateStructureError(
            f"Explicit structure element {item.element_id!r} differs across slides: "
            f"{slide_names}"
        )
    for state, shape in zip(states, resolved_shapes):
        shape_id = _shape_id(shape)
        if shape_id and shape_id in _timing_shape_ids(state.root):
            raise TemplateStructureError(
                f"{state.spec.svg_path.name}: structure element {item.element_id!r} "
                "is referenced by slide timing"
            )
        if not _shape_relationships_supported(shape, state.rels):
            raise TemplateStructureError(
                f"{state.spec.svg_path.name}: structure element {item.element_id!r} "
                "uses a relationship that cannot move to a template part"
            )

    prototype_state = states[0]
    prototype_shape = resolved_shapes[0]
    target_shape = _copy_shape_relationships_to_part(
        prototype_shape,
        prototype_state.rels,
        target_rels_path,
    )
    _set_shape_name(target_shape, f"{item.element_id} {item.layer.title()}")
    _append_shape_to_part(target_path, target_shape)
    for state, shape in zip(states, resolved_shapes):
        _remove_template_shape(state, shape)
    return None


def _shape_bounds_emu(
    shape: ET.Element,
    override_px: tuple[float, float, float, float] | None,
) -> tuple[int, int, int, int]:
    if override_px is not None:
        x, y, width, height = override_px
        return tuple(
            round(value * EMU_PER_PX)
            for value in (x, y, width, height)
        )

    xfrm = shape.find(f"{{{PML_NS}}}spPr/{{{DML_NS}}}xfrm")
    if xfrm is None:
        xfrm = shape.find(f"{{{PML_NS}}}xfrm")
    if xfrm is None:
        raise TemplateStructureError(
            "Placeholder shape has no directly readable DrawingML transform; "
            "set data-pptx-bounds"
        )
    off = xfrm.find(f"{{{DML_NS}}}off")
    ext = xfrm.find(f"{{{DML_NS}}}ext")
    if off is None or ext is None:
        raise TemplateStructureError("Placeholder transform has no a:off/a:ext")
    try:
        return (
            int(off.attrib["x"]),
            int(off.attrib["y"]),
            int(ext.attrib["cx"]),
            int(ext.attrib["cy"]),
        )
    except (KeyError, ValueError) as exc:
        raise TemplateStructureError("Placeholder transform is invalid") from exc


def _replace_shape_xfrm(
    sp_pr: ET.Element,
    bounds: tuple[int, int, int, int],
) -> None:
    for existing in list(sp_pr):
        if existing.tag == f"{{{DML_NS}}}xfrm":
            sp_pr.remove(existing)
    x, y, width, height = bounds
    xfrm = ET.Element(f"{{{DML_NS}}}xfrm")
    ET.SubElement(xfrm, f"{{{DML_NS}}}off", {"x": str(x), "y": str(y)})
    ET.SubElement(
        xfrm,
        f"{{{DML_NS}}}ext",
        {"cx": str(width), "cy": str(height)},
    )
    sp_pr.insert(0, xfrm)


def _placeholder_vertical_anchor(
    source_bounds: tuple[int, int, int, int],
    target_bounds: tuple[int, int, int, int],
) -> str:
    """Preserve an intentionally centered carrier inside its full slot frame."""
    _, source_y, _, source_height = source_bounds
    _, target_y, _, target_height = target_bounds
    source_center = source_y + source_height / 2
    target_center = target_y + target_height / 2
    return (
        "ctr"
        if abs(source_center - target_center) <= target_height * 0.2
        else "t"
    )


def _normalize_placeholder_body_properties(
    body_pr: ET.Element,
    source_bounds: tuple[int, int, int, int],
    target_bounds: tuple[int, int, int, int],
) -> None:
    """Make a full-frame placeholder wrap text while preserving vertical intent."""
    body_pr.set("wrap", "square")
    body_pr.set("anchor", _placeholder_vertical_anchor(source_bounds, target_bounds))
    body_pr.set("anchorCtr", "0")
    autofit_tags = {
        f"{{{DML_NS}}}noAutofit",
        f"{{{DML_NS}}}normAutofit",
        f"{{{DML_NS}}}spAutoFit",
    }
    for child in list(body_pr):
        if child.tag in autofit_tags:
            body_pr.remove(child)
    body_pr.append(ET.Element(f"{{{DML_NS}}}noAutofit"))


def _apply_layout_frame_to_placeholder_carrier(
    shape: ET.Element,
    item: TemplateElementSpec,
) -> None:
    """Use the reusable Layout bounds on one template-review Slide carrier."""
    if item.placeholder_bounds is None:
        raise TemplateStructureError(
            f"Placeholder {item.element_id!r} has no reusable Layout bounds"
        )
    source_bounds = _shape_bounds_emu(shape, None)
    target_bounds = _shape_bounds_emu(shape, item.placeholder_bounds)
    if shape.tag in {f"{{{PML_NS}}}sp", f"{{{PML_NS}}}pic"}:
        sp_pr = shape.find(f"{{{PML_NS}}}spPr")
        if sp_pr is None:
            raise TemplateStructureError(
                f"Placeholder {item.element_id!r} has no p:spPr"
            )
        _replace_shape_xfrm(sp_pr, target_bounds)
    elif shape.tag == f"{{{PML_NS}}}graphicFrame":
        xfrm = shape.find(f"{{{PML_NS}}}xfrm")
        if xfrm is None:
            xfrm = ET.Element(f"{{{PML_NS}}}xfrm")
            shape.insert(1, xfrm)
        for child in list(xfrm):
            if child.tag in {f"{{{DML_NS}}}off", f"{{{DML_NS}}}ext"}:
                xfrm.remove(child)
        x, y, width, height = target_bounds
        ET.SubElement(xfrm, f"{{{DML_NS}}}off", {"x": str(x), "y": str(y)})
        ET.SubElement(
            xfrm,
            f"{{{DML_NS}}}ext",
            {"cx": str(width), "cy": str(height)},
        )
    else:
        raise TemplateStructureError(
            f"Placeholder {item.element_id!r} cannot use Layout bounds on "
            f"DrawingML element {shape.tag.rsplit('}', 1)[-1]!r}"
        )

    tx_body = shape.find(f"{{{PML_NS}}}txBody")
    if tx_body is None:
        return
    body_pr = tx_body.find(f"{{{DML_NS}}}bodyPr")
    if body_pr is None:
        body_pr = ET.Element(f"{{{DML_NS}}}bodyPr")
        tx_body.insert(0, body_pr)
    _normalize_placeholder_body_properties(
        body_pr,
        source_bounds,
        target_bounds,
    )


def _layout_level_one_paragraph_properties(
    list_style: ET.Element,
) -> ET.Element:
    """Return the Layout list style's level-one paragraph properties."""
    level_tag = f"{{{DML_NS}}}lvl1pPr"
    level_props = list_style.find(level_tag)
    if level_props is None:
        level_props = ET.Element(level_tag)
        trailing_tags = {
            f"{{{DML_NS}}}lvl{level}pPr" for level in range(2, 10)
        }
        trailing_tags.add(f"{{{DML_NS}}}extLst")
        insert_at = next(
            (
                index
                for index, child in enumerate(list_style)
                if child.tag in trailing_tags
            ),
            len(list_style),
        )
        list_style.insert(insert_at, level_props)
    return level_props


def _set_layout_level_one_default_size(
    list_style: ET.Element,
    source_run_pr: ET.Element | None,
) -> None:
    """Persist the prototype run size as the Layout's level-one text default."""
    if source_run_pr is None or source_run_pr.get("sz") is None:
        return
    level_props = _layout_level_one_paragraph_properties(list_style)
    default_props = level_props.find(f"{{{DML_NS}}}defRPr")
    if default_props is None:
        default_props = ET.Element(f"{{{DML_NS}}}defRPr")
        ext_tag = f"{{{DML_NS}}}extLst"
        insert_at = next(
            (
                index
                for index, child in enumerate(level_props)
                if child.tag == ext_tag
            ),
            len(level_props),
        )
        level_props.insert(insert_at, default_props)
    default_props.set("sz", source_run_pr.get("sz", ""))


def _set_no_bullet_paragraph_properties(
    paragraph_props: ET.Element,
    *,
    replace_existing: bool = False,
) -> None:
    """Disable inherited bullets and hanging indent for a prose paragraph."""
    if replace_existing:
        for child in list(paragraph_props):
            if child.tag in _PARAGRAPH_BULLET_CHOICE_TAGS:
                paragraph_props.remove(child)

    bullet_choice = next(
        (
            child
            for child in paragraph_props
            if child.tag in _PARAGRAPH_BULLET_CHOICE_TAGS
        ),
        None,
    )
    if bullet_choice is not None and bullet_choice.tag != f"{{{DML_NS}}}buNone":
        return
    if bullet_choice is None:
        insert_at = next(
            (
                index
                for index, child in enumerate(paragraph_props)
                if child.tag in _PARAGRAPH_PROPERTIES_TRAILING_TAGS
            ),
            len(paragraph_props),
        )
        paragraph_props.insert(insert_at, ET.Element(f"{{{DML_NS}}}buNone"))

    paragraph_props.set("marL", "0")
    paragraph_props.set("indent", "0")


def _placeholder_text_body(
    source_shape: ET.Element,
    item: TemplateElementSpec,
) -> ET.Element:
    tx_body = ET.Element(f"{{{PML_NS}}}txBody")
    source_tx_body = source_shape.find(f"{{{PML_NS}}}txBody")
    source_body_pr = (
        source_tx_body.find(f"{{{DML_NS}}}bodyPr")
        if source_tx_body is not None
        else None
    )
    source_lst_style = (
        source_tx_body.find(f"{{{DML_NS}}}lstStyle")
        if source_tx_body is not None
        else None
    )
    source_run_pr = (
        source_tx_body.find(f".//{{{DML_NS}}}rPr")
        if source_tx_body is not None
        else None
    )
    body_pr = (
        ET.fromstring(ET.tostring(source_body_pr, encoding="utf-8"))
        if source_body_pr is not None
        else ET.Element(f"{{{DML_NS}}}bodyPr")
    )
    target_bounds = _shape_bounds_emu(source_shape, item.placeholder_bounds)
    try:
        source_bounds = _shape_bounds_emu(source_shape, None)
    except TemplateStructureError:
        # Composite proxy content may compile to p:grpSp, whose transform is
        # intentionally not reused for the Layout's synthetic p:sp carrier.
        # The explicit design-zone bounds remain the authoritative frame.
        source_bounds = target_bounds
    _normalize_placeholder_body_properties(
        body_pr,
        source_bounds,
        target_bounds,
    )
    tx_body.append(body_pr)
    list_style = (
        ET.fromstring(ET.tostring(source_lst_style, encoding="utf-8"))
        if source_lst_style is not None
        else ET.Element(f"{{{DML_NS}}}lstStyle")
    )
    _set_layout_level_one_default_size(list_style, source_run_pr)
    if item.placeholder in {"body", "subtitle"}:
        _set_no_bullet_paragraph_properties(
            _layout_level_one_paragraph_properties(list_style),
            replace_existing=item.placeholder == "body",
        )
    tx_body.append(list_style)

    paragraph = ET.SubElement(tx_body, f"{{{DML_NS}}}p")
    source_paragraph_props = (
        source_tx_body.find(f"{{{DML_NS}}}p/{{{DML_NS}}}pPr")
        if source_tx_body is not None
        else None
    )
    paragraph_props = (
        ET.fromstring(ET.tostring(source_paragraph_props, encoding="utf-8"))
        if source_paragraph_props is not None
        else None
    )
    if item.placeholder in {"body", "subtitle"}:
        if paragraph_props is None:
            paragraph_props = ET.Element(f"{{{DML_NS}}}pPr")
        _set_no_bullet_paragraph_properties(paragraph_props)
    if paragraph_props is not None:
        paragraph.append(paragraph_props)
    if item.placeholder in {"slide-number", "date"}:
        field_type = (
            "slidenum"
            if item.placeholder == "slide-number"
            else "datetimeFigureOut"
        )
        field = ET.SubElement(
            paragraph,
            f"{{{DML_NS}}}fld",
            {"id": f"{{{str(uuid.uuid4()).upper()}}}", "type": field_type},
        )
        if source_run_pr is not None:
            field.append(ET.fromstring(ET.tostring(source_run_pr, encoding="utf-8")))
        source_text = ""
        if source_tx_body is not None:
            source_text = "".join(
                text.text or ""
                for text in source_tx_body.findall(f".//{{{DML_NS}}}t")
            )
        field_text = (
            "‹#›"
            if item.placeholder == "slide-number"
            else source_text or "Date"
        )
        ET.SubElement(field, f"{{{DML_NS}}}t").text = field_text
    else:
        run = ET.SubElement(paragraph, f"{{{DML_NS}}}r")
        if source_run_pr is not None:
            run.append(ET.fromstring(ET.tostring(source_run_pr, encoding="utf-8")))
        ET.SubElement(run, f"{{{DML_NS}}}t").text = _TEMPLATE_PLACEHOLDER_PROMPTS.get(
            item.placeholder or "",
            "Click to add content",
        )
    ET.SubElement(paragraph, f"{{{DML_NS}}}endParaRPr", {"lang": "en-US"})
    return tx_body


def _set_placeholder_no_inherited_bullets(
    shape: ET.Element,
    item: TemplateElementSpec,
) -> None:
    """Keep prose bullet-free while preserving explicit subtitle bullets."""
    if item.placeholder not in {"body", "subtitle"}:
        return
    tx_body = shape.find(f"{{{PML_NS}}}txBody")
    if tx_body is None:
        return
    for paragraph in tx_body.findall(f"{{{DML_NS}}}p"):
        paragraph_props = paragraph.find(f"{{{DML_NS}}}pPr")
        if paragraph_props is None:
            paragraph_props = ET.Element(f"{{{DML_NS}}}pPr")
            paragraph.insert(0, paragraph_props)
        _set_no_bullet_paragraph_properties(
            paragraph_props,
            replace_existing=item.placeholder == "body",
        )


def _set_placeholder_theme_font_role(
    shape: ET.Element,
    item: TemplateElementSpec,
    theme_font_spec: ThemeFontSpec | None,
) -> None:
    """Force semantic text placeholders onto the correct theme font role."""
    if theme_font_spec is None:
        return
    if item.placeholder == "title":
        prefix = "+mj"
    elif item.placeholder in TEMPLATE_PLACEHOLDER_TYPES:
        prefix = "+mn"
    else:
        return
    for props_tag in ("rPr", "defRPr", "endParaRPr"):
        for props in shape.iter(f"{{{DML_NS}}}{props_tag}"):
            for font_tag, suffix in (("latin", "lt"), ("ea", "ea"), ("cs", "cs")):
                font = props.find(f"{{{DML_NS}}}{font_tag}")
                if font is not None:
                    font.set("typeface", f"{prefix}-{suffix}")


def _layout_placeholder_shape(
    source_shape: ET.Element,
    item: TemplateElementSpec,
    placeholder_idx: int | None,
    theme_font_spec: ThemeFontSpec | None = None,
) -> ET.Element:
    """Build one reusable p:sp placeholder from a prototype slide object."""
    placeholder_type = TEMPLATE_PLACEHOLDER_TYPES.get(item.placeholder or "")
    if placeholder_type is None:
        raise TemplateStructureError(
            f"Unsupported placeholder type: {item.placeholder!r}"
        )
    bounds = _shape_bounds_emu(source_shape, item.placeholder_bounds)
    shape = ET.Element(f"{{{PML_NS}}}sp")
    nv_sp_pr = ET.SubElement(shape, f"{{{PML_NS}}}nvSpPr")
    ET.SubElement(
        nv_sp_pr,
        f"{{{PML_NS}}}cNvPr",
        {"id": "2", "name": f"{item.element_id} Placeholder"},
    )
    c_nv_sp_pr = ET.SubElement(nv_sp_pr, f"{{{PML_NS}}}cNvSpPr")
    ET.SubElement(c_nv_sp_pr, f"{{{DML_NS}}}spLocks", {"noGrp": "1"})
    nv_pr = ET.SubElement(nv_sp_pr, f"{{{PML_NS}}}nvPr")
    placeholder_attrs = {"type": placeholder_type}
    if placeholder_idx is not None:
        placeholder_attrs["idx"] = str(placeholder_idx)
    elif item.placeholder != "title":
        raise TemplateStructureError(
            f"Placeholder {item.element_id!r} requires an idx"
        )
    ET.SubElement(nv_pr, f"{{{PML_NS}}}ph", placeholder_attrs)

    source_sp_pr = (
        source_shape.find(f"{{{PML_NS}}}spPr")
        if source_shape.tag == f"{{{PML_NS}}}sp"
        else None
    )
    if source_sp_pr is not None:
        sp_pr = ET.fromstring(ET.tostring(source_sp_pr, encoding="utf-8"))
    else:
        sp_pr = ET.Element(f"{{{PML_NS}}}spPr")
        geometry = ET.SubElement(sp_pr, f"{{{DML_NS}}}prstGeom", {"prst": "rect"})
        ET.SubElement(geometry, f"{{{DML_NS}}}avLst")
        ET.SubElement(sp_pr, f"{{{DML_NS}}}noFill")
        line = ET.SubElement(sp_pr, f"{{{DML_NS}}}ln")
        ET.SubElement(line, f"{{{DML_NS}}}noFill")
    _replace_shape_xfrm(sp_pr, bounds)
    shape.append(sp_pr)
    shape.append(_placeholder_text_body(source_shape, item))
    _set_placeholder_theme_font_role(shape, item, theme_font_spec)
    return shape


def _placeholder_binding_proxy(
    layout_placeholder: ET.Element,
    item: TemplateElementSpec,
) -> ET.Element:
    """Bind a Layout slot invisibly while leaving its visible content ordinary.

    An unbound object placeholder can leak its inherited empty frame into a
    finished Slide in non-PowerPoint renderers. A hidden matching proxy suppresses
    that inheritance. The zero-width transparent run avoids a LibreOffice empty-
    placeholder black fill without adding visible content.
    """
    proxy = ET.fromstring(ET.tostring(layout_placeholder, encoding="utf-8"))
    c_nv_pr = next(proxy.iter(f"{{{PML_NS}}}cNvPr"), None)
    if c_nv_pr is None:
        raise TemplateStructureError(
            f"Cannot create placeholder binding for {item.element_id!r}: "
            "p:cNvPr is missing"
        )
    placeholder = proxy.find(f".//{{{PML_NS}}}ph")
    if placeholder is None:
        raise TemplateStructureError(
            f"Cannot create placeholder binding for {item.element_id!r}: "
            "p:ph is missing"
        )
    c_nv_pr.set(
        "name",
        "Placeholder Binding "
        f"{placeholder.get('type', 'body')} {placeholder.get('idx', '0')}",
    )
    c_nv_pr.set("hidden", "1")
    tx_body = proxy.find(f"{{{PML_NS}}}txBody")
    if tx_body is None:
        raise TemplateStructureError(
            f"Cannot create placeholder binding for {item.element_id!r}: "
            "p:txBody is missing"
        )
    for child in list(tx_body):
        if child.tag == f"{{{DML_NS}}}p":
            tx_body.remove(child)
    paragraph = ET.SubElement(tx_body, f"{{{DML_NS}}}p")
    run = ET.SubElement(paragraph, f"{{{DML_NS}}}r")
    run_props = ET.SubElement(
        run,
        f"{{{DML_NS}}}rPr",
        {"lang": "en-US", "sz": "100"},
    )
    solid_fill = ET.SubElement(run_props, f"{{{DML_NS}}}solidFill")
    color = ET.SubElement(solid_fill, f"{{{DML_NS}}}srgbClr", {"val": "FFFFFF"})
    ET.SubElement(color, f"{{{DML_NS}}}alpha", {"val": "0"})
    ET.SubElement(run, f"{{{DML_NS}}}t").text = "\u200b"
    ET.SubElement(
        paragraph,
        f"{{{DML_NS}}}endParaRPr",
        {"lang": "en-US"},
    )
    return proxy


def _patch_slide_placeholder(
    shape: ET.Element,
    item: TemplateElementSpec,
    placeholder_idx: int | None,
    placeholder_type: str | None = None,
    theme_font_spec: ThemeFontSpec | None = None,
) -> None:
    resolved_type = placeholder_type or TEMPLATE_PLACEHOLDER_TYPES.get(
        item.placeholder or ""
    )
    if resolved_type is None:
        raise TemplateStructureError(
            f"Unsupported placeholder type: {item.placeholder!r}"
        )
    nv_paths = {
        f"{{{PML_NS}}}sp": f"{{{PML_NS}}}nvSpPr/{{{PML_NS}}}nvPr",
        f"{{{PML_NS}}}pic": f"{{{PML_NS}}}nvPicPr/{{{PML_NS}}}nvPr",
        f"{{{PML_NS}}}graphicFrame": (
            f"{{{PML_NS}}}nvGraphicFramePr/{{{PML_NS}}}nvPr"
        ),
    }
    nv_path = nv_paths.get(shape.tag)
    if nv_path is None:
        raise TemplateStructureError(
            f"Placeholder {item.element_id!r} converted to unsupported "
            f"DrawingML element {shape.tag.rsplit('}', 1)[-1]!r}; text/picture/"
            "native chart/table placeholders must remain one top-level object"
        )
    nv_pr = shape.find(nv_path)
    if nv_pr is None:
        raise TemplateStructureError(
            f"Placeholder {item.element_id!r} has no non-visual properties"
        )
    _set_placeholder_no_inherited_bullets(shape, item)
    _set_placeholder_theme_font_role(shape, item, theme_font_spec)
    for existing in list(nv_pr):
        if existing.tag == f"{{{PML_NS}}}ph":
            nv_pr.remove(existing)
    placeholder_attrs: dict[str, str] = {}
    if (
        placeholder_type is not None
        or resolved_type != "obj"
        or placeholder_idx is None
    ):
        placeholder_attrs["type"] = resolved_type
    if placeholder_idx is not None:
        placeholder_attrs["idx"] = str(placeholder_idx)
    ph = ET.Element(f"{{{PML_NS}}}ph", placeholder_attrs)
    ext_tag = f"{{{PML_NS}}}extLst"
    insert_at = next(
        (idx for idx, child in enumerate(nv_pr) if child.tag == ext_tag),
        len(nv_pr),
    )
    nv_pr.insert(insert_at, ph)


def _set_template_layout_header_footer(
    layout_path: Path,
    placeholders: tuple[TemplateElementSpec, ...],
) -> None:
    """Enable declared footer fields for slides newly created from the layout."""
    kinds = {item.placeholder for item in placeholders}
    if not kinds.intersection({"date", "footer", "slide-number"}):
        return
    tree = ET.parse(layout_path)
    root = tree.getroot()
    hf = root.find(f"{{{PML_NS}}}hf")
    if hf is None:
        hf = ET.Element(f"{{{PML_NS}}}hf")
        trailing_tags = {
            f"{{{PML_NS}}}timing",
            f"{{{PML_NS}}}transition",
            f"{{{PML_NS}}}extLst",
        }
        insert_at = next(
            (idx for idx, child in enumerate(root) if child.tag in trailing_tags),
            len(root),
        )
        root.insert(insert_at, hf)
    hf.set("hdr", "0")
    hf.set("dt", "1" if "date" in kinds else "0")
    hf.set("ftr", "1" if "footer" in kinds else "0")
    hf.set("sldNum", "1" if "slide-number" in kinds else "0")
    _write_xml_tree(layout_path, tree)


def _apply_explicit_layout_structure(
    extract_dir: Path,
    structure: PptxStructureContext,
    specs: list[TemplateSlideSpec],
    conversion_traces: list[dict[str, Any]] | None,
    theme_font_spec: ThemeFontSpec | None,
    *,
    use_layout_placeholder_frames: bool = False,
    verbose: bool = False,
) -> tuple[
    dict[str, str | None],
    dict[str, tuple[str, ...]],
    dict[str, str],
    dict[str, str],
]:
    """Materialize explicit SVG master/layout/placeholder metadata into OOXML."""
    master_parts_by_key = _assign_structured_masters(
        extract_dir,
        structure,
        specs,
    )
    states = _template_runtime_slides(extract_dir, specs, conversion_traces)
    states_by_slide = {state.spec.slide_num: state for state in states}
    slide_size_emu = _presentation_slide_size_emu(extract_dir)

    expected_backgrounds: dict[str, str | None] = {}
    expected_shape_rosters: dict[str, tuple[str, ...]] = {}
    states_by_master: dict[str, list[_TemplateRuntimeSlide]] = {}
    for state in states:
        master_part = master_parts_by_key[state.spec.master_key]
        states_by_master.setdefault(master_part, []).append(state)
    master_shape_count = 0
    for master_part, master_states in states_by_master.items():
        master_path = extract_dir / master_part
        master_rels_path = _relationships_path_for_part(extract_dir, master_part)
        expected_backgrounds[master_part] = _extract_slide_background_xml(
            master_path.read_text(encoding="utf-8")
        )
        _clear_master_placeholder_shapes(master_path)
        master_items = master_states[0].spec.master_elements
        for item in master_items:
            background_xml = _move_template_static_shape(
                master_states,
                item,
                master_path,
                master_rels_path,
                slide_size_emu,
            )
            if background_xml is not None:
                expected_backgrounds[master_part] = background_xml
            master_shape_count += 1

    specs_by_layout: dict[str, list[TemplateSlideSpec]] = {}
    for spec in specs:
        specs_by_layout.setdefault(spec.layout_key, []).append(spec)
    placeholder_count = 0
    layout_shape_count = 0
    created_layout_parts: set[str] = set()
    layout_parts_by_key: dict[str, str] = {}
    for layout_key, layout_specs in specs_by_layout.items():
        layout_states = [states_by_slide[spec.slide_num] for spec in layout_specs]
        master_parts = {
            structure.slide_master_part(spec.slide_num) for spec in layout_specs
        }
        if len(master_parts) != 1:
            raise TemplateStructureError(
                f"Layout {layout_key!r} spans multiple slide masters; use distinct "
                "layout keys per master"
            )
        master_part = next(iter(master_parts))
        prototype = layout_specs[0]
        base_target = structure.slide_layout_target(prototype.slide_num)
        base_layout_part = _resolve_package_target(
            f"ppt/slides/slide{prototype.slide_num}.xml",
            base_target,
        )
        layout_target, layout_part = _create_custom_layout(
            extract_dir,
            master_part,
            base_layout_part,
            prototype.layout_name,
            show_master_shapes=prototype.layout_show_master_shapes,
        )
        layout_path = extract_dir / layout_part
        layout_rels_path = _relationships_path_for_part(extract_dir, layout_part)
        created_layout_parts.add(layout_part)
        layout_parts_by_key[layout_key] = layout_part

        placeholder_bindings = {
            binding.element.element_id: binding
            for binding in template_placeholder_bindings(prototype)
        }
        for item in prototype.elements:
            if item.layer == "layout":
                background_xml = _move_template_static_shape(
                    layout_states,
                    item,
                    layout_path,
                    layout_rels_path,
                    slide_size_emu,
                )
                if background_xml is not None:
                    expected_backgrounds[layout_part] = background_xml
                layout_shape_count += 1
                continue
            if not item.placeholder:
                continue
            proxy_binding = is_proxy_placeholder(item)
            if proxy_binding:
                placeholder_shapes = [
                    _template_shape_for_item(state, item)
                    for state in layout_states
                ]
                if any(shape is None for shape in placeholder_shapes):
                    raise TemplateStructureError(
                        f"Placeholder {item.element_id!r} cannot be a slide background"
                    )
                resolved_shapes = [
                    shape for shape in placeholder_shapes if shape is not None
                ]
            else:
                resolved_shapes = [
                    _unwrap_placeholder_carrier(state, item)
                    for state in layout_states
                ]
            prototype_shape = resolved_shapes[0]
            binding = placeholder_bindings[item.element_id]
            assigned_idx = binding.assigned_idx
            layout_placeholder = _layout_placeholder_shape(
                prototype_shape,
                item,
                assigned_idx,
                theme_font_spec,
            )
            _append_shape_to_part(layout_path, layout_placeholder)
            for state, shape in zip(layout_states, resolved_shapes):
                if not proxy_binding:
                    _patch_slide_placeholder(
                        shape,
                        item,
                        assigned_idx,
                        theme_font_spec=theme_font_spec,
                    )
                    if use_layout_placeholder_frames:
                        _apply_layout_frame_to_placeholder_carrier(shape, item)
                    _set_shape_name(
                        shape,
                        f"{item.element_id} Placeholder Carrier",
                    )
                else:
                    _set_shape_name(
                        shape,
                        f"{item.element_id} Proxy Content",
                    )
                    _append_shape_to_runtime_slide(
                        state,
                        _placeholder_binding_proxy(layout_placeholder, item),
                    )
            placeholder_count += 1

        _set_template_layout_header_footer(layout_path, prototype.placeholders)
        for state in layout_states:
            _set_slide_layout_target(state.rels_path, layout_target)

    slide_backgrounds = _apply_template_slide_backgrounds(
        states,
        slide_size_emu,
    )
    expected_backgrounds.update(slide_backgrounds)
    for state in states:
        state.root.set(
            "showMasterSp",
            "1" if state.spec.slide_show_inherited_shapes else "0",
        )
        _write_xml_tree(state.slide_path, state.tree)
        expected_backgrounds.setdefault(
            f"ppt/slides/slide{state.spec.slide_num}.xml",
            None,
        )
        expected_shape_rosters[
            f"ppt/slides/slide{state.spec.slide_num}.xml"
        ] = _top_level_shape_name_roster(state.root)
    for part in states_by_master:
        expected_backgrounds.setdefault(
            part,
            _extract_slide_background_xml(
                (extract_dir / part).read_text(encoding="utf-8")
            ),
        )
    for part in created_layout_parts:
        expected_backgrounds.setdefault(part, None)
    for part in (*states_by_master, *created_layout_parts):
        expected_shape_rosters[part] = _top_level_shape_name_roster(
            ET.parse(extract_dir / part).getroot()
        )

    if verbose:
        print(
            "  Explicit Layout structure: "
            f"{len(states_by_master)} master(s), "
            f"{len(specs_by_layout)} layout(s), "
            f"{master_shape_count} master element(s), "
            f"{layout_shape_count} layout element(s), "
            f"{len(slide_backgrounds)} slide background(s), "
            f"{placeholder_count} placeholder definition(s)"
        )
    return (
        expected_backgrounds,
        expected_shape_rosters,
        layout_parts_by_key,
        master_parts_by_key,
    )


def _apply_preserved_structure(
    extract_dir: Path,
    specs: list[TemplateSlideSpec],
    contract: NativeStructureContract,
    conversion_traces: list[dict[str, Any]] | None,
    *,
    verbose: bool = False,
) -> None:
    """Drop preview-only inherited layers and bind content to source placeholders."""
    states = _template_runtime_slides(extract_dir, specs, conversion_traces)
    removed_preview_shapes = 0
    removed_preview_backgrounds = 0
    placeholder_count = 0
    for state in states:
        removed_background = False
        for item in state.spec.elements:
            if item.layer not in {"master", "layout"}:
                continue
            shape = _template_shape_for_item(state, item)
            if shape is not None:
                _remove_template_shape(state, shape)
                removed_preview_shapes += 1
                continue
            if removed_background:
                raise TemplateStructureError(
                    f"{state.spec.svg_path.name}: multiple inherited preview "
                    "backgrounds resolved to one slide background"
                )
            common_slide = state.root.find(f"{{{PML_NS}}}cSld")
            background = (
                common_slide.find(f"{{{PML_NS}}}bg")
                if common_slide is not None
                else None
            )
            if common_slide is None or background is None:
                raise TemplateStructureError(
                    f"{state.spec.svg_path.name}: inherited preview background "
                    "did not produce a removable slide background"
                )
            common_slide.remove(background)
            removed_background = True
            removed_preview_backgrounds += 1

        layout = contract.layout(state.spec.layout_key)
        for item, source_placeholder in match_native_placeholders(state.spec, layout):
            shape = _template_shape_for_item(state, item)
            if shape is None:
                raise TemplateStructureError(
                    f"{state.spec.svg_path.name}: placeholder {item.element_id!r} "
                    "cannot resolve to a slide background"
                )
            _patch_slide_placeholder(
                shape,
                item,
                source_placeholder.idx,
                source_placeholder.placeholder_type,
            )
            placeholder_count += 1
        _write_xml_tree(state.slide_path, state.tree)

    if verbose:
        print(
            "  Preserved structure: "
            f"{len({spec.layout_key for spec in specs})} source layout(s), "
            f"{removed_preview_shapes} preview shape(s) removed, "
            f"{removed_preview_backgrounds} preview background(s) removed, "
            f"{placeholder_count} source placeholder binding(s)"
        )


def _promote_common_chrome_shapes_to_masters(
    extract_dir: Path,
    structure: PptxStructureContext,
    slide_count: int,
    conversion_traces: list[dict[str, Any]] | None,
    *,
    verbose: bool = False,
) -> int:
    """Promote explicit repeated chrome SVG ids to their shared master."""
    if not conversion_traces:
        return 0
    trace_by_slide = {
        int(trace.get("slide_num", 0)): trace
        for trace in conversion_traces
        if trace.get("slide_num") is not None
    }
    if len(trace_by_slide) < slide_count:
        return 0

    slides_by_master: dict[str, list[int]] = {}
    for slide_num in range(1, slide_count + 1):
        master_part = structure.slide_master_part(slide_num)
        slides_by_master.setdefault(master_part, []).append(slide_num)

    promoted = 0
    promoted_roles = 0
    for master_part, slide_nums in slides_by_master.items():
        if len(slide_nums) < 2:
            continue
        slide_state: dict[int, dict[str, Any]] = {}
        for slide_num in slide_nums:
            slide_path = extract_dir / "ppt" / "slides" / f"slide{slide_num}.xml"
            rels_path = extract_dir / "ppt" / "slides" / "_rels" / f"slide{slide_num}.xml.rels"
            tree = ET.parse(slide_path)
            root = tree.getroot()
            slide_state[slide_num] = {
                "path": slide_path,
                "rels": _read_relationships(rels_path),
                "root": root,
                "shapes": _top_level_shapes_by_id(root),
                "timing_shape_ids": _timing_shape_ids(root),
                "tokens": _trace_chrome_shape_ids(trace_by_slide.get(slide_num)),
                "tree": tree,
            }

        # Per token, find the strict-majority identical variant. Slides
        # outside every dominant set become cover-layout minority slides.
        candidate_sets: dict[str, dict[int, str]] = {}
        all_tokens = sorted({
            token
            for state in slide_state.values()
            for token in state["tokens"]
        })
        for token in all_tokens:
            carriers: dict[int, str] = {}
            canonical_by_slide: dict[int, bytes] = {}
            for slide_num in slide_nums:
                state = slide_state[slide_num]
                shape_ids = state["tokens"].get(token, [])
                if len(shape_ids) != 1:
                    continue
                shape_id = shape_ids[0]
                shape = state["shapes"].get(shape_id)
                if shape is None:
                    continue
                if shape_id in state["timing_shape_ids"]:
                    continue
                if not _shape_relationships_supported(shape, state["rels"]):
                    continue
                carriers[slide_num] = shape_id
                canonical_by_slide[slide_num] = _canonical_shape_xml(
                    shape,
                    state["rels"],
                )
            dominant_xml, dominant_slides = _dominant_variant(canonical_by_slide)
            if dominant_xml is None:
                continue
            if not _is_strict_majority(len(dominant_slides), len(slide_nums)):
                continue
            candidate_sets[token] = {
                slide_num: carriers[slide_num] for slide_num in dominant_slides
            }

        if not candidate_sets:
            continue
        content_slides = sorted(set.intersection(
            *(set(slides) for slides in candidate_sets.values())
        ))
        if not _is_strict_majority(len(content_slides), len(slide_nums)):
            continue

        promotions: list[tuple[str, dict[int, str]]] = []
        claimed_shape_ids: dict[int, set[str]] = {
            slide_num: set() for slide_num in content_slides
        }
        for token in sorted(candidate_sets):
            shape_ids_by_slide = {
                slide_num: candidate_sets[token][slide_num]
                for slide_num in content_slides
            }
            # A flattened nested chrome group can emit several semantic trace
            # ids for the same generated DrawingML shape. Claim it once.
            if any(
                shape_ids_by_slide[slide_num] in claimed_shape_ids[slide_num]
                for slide_num in content_slides
            ):
                continue
            for slide_num, shape_id in shape_ids_by_slide.items():
                claimed_shape_ids[slide_num].add(shape_id)
            promotions.append((token, shape_ids_by_slide))

        if not promotions:
            continue

        # Master shapes always render behind slide-local shapes. Preserve the
        # original z-order by promoting only a common leading chrome prefix;
        # overlay headers/footers remain slide-local.
        token_by_shape_id = {
            slide_num: {
                shape_ids[slide_num]: token
                for token, shape_ids in promotions
            }
            for slide_num in content_slides
        }
        leading_token_orders: list[list[str]] = []
        for slide_num in content_slides:
            order: list[str] = []
            for shape_id in slide_state[slide_num]["shapes"]:
                token = token_by_shape_id[slide_num].get(shape_id)
                if token is None:
                    break
                order.append(token)
            leading_token_orders.append(order)

        safe_tokens = list(leading_token_orders[0])
        for order in leading_token_orders[1:]:
            common_length = 0
            for expected, actual in zip(safe_tokens, order):
                if expected != actual:
                    break
                common_length += 1
            safe_tokens = safe_tokens[:common_length]
            if not safe_tokens:
                break
        promotion_by_token = {token: shape_ids for token, shape_ids in promotions}
        promotions = [
            (token, promotion_by_token[token])
            for token in safe_tokens
        ]

        if not promotions:
            continue

        master_path = extract_dir / master_part
        master_rels_path = _relationships_path_for_part(extract_dir, master_part)
        for _token, shape_ids_by_slide in promotions:
            first_slide = content_slides[0]
            first_state = slide_state[first_slide]
            shape = first_state["shapes"][shape_ids_by_slide[first_slide]]
            master_shape = _copy_shape_relationships_to_master(
                shape,
                first_state["rels"],
                master_rels_path,
            )
            _append_shape_to_master(master_path, master_shape)
            promoted_roles += 1

            for slide_num, shape_id in shape_ids_by_slide.items():
                state = slide_state[slide_num]
                shape_to_remove = state["shapes"].get(shape_id)
                sp_tree = state["root"].find(f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree")
                if sp_tree is not None and shape_to_remove is not None:
                    sp_tree.remove(shape_to_remove)
                    promoted += 1

        for slide_num in content_slides:
            state = slide_state[slide_num]
            _write_xml_tree(state["path"], state["tree"])

        # Minority slides (covers, section pages) keep every shape
        # slide-local and move to a Cover layout that hides the newly
        # promoted master chrome, so their rendering never changes.
        minority_slides = [
            slide_num for slide_num in slide_nums
            if slide_num not in set(content_slides)
        ]
        if minority_slides:
            first_minority_rels = (
                extract_dir / "ppt" / "slides" / "_rels"
                / f"slide{minority_slides[0]}.xml.rels"
            )
            base_target = _find_relationship_target(
                first_minority_rels, SLIDE_LAYOUT_REL_TYPE
            )
            if not base_target:
                raise RuntimeError(
                    f"Slide {minority_slides[0]} has no slide layout relationship"
                )
            base_layout_part = _resolve_package_target(
                f"ppt/slides/slide{minority_slides[0]}.xml", base_target
            )
            cover_target = _create_cover_layout(
                extract_dir, master_part, base_layout_part
            )
            for slide_num in minority_slides:
                rels_path = (
                    extract_dir / "ppt" / "slides" / "_rels"
                    / f"slide{slide_num}.xml.rels"
                )
                _set_slide_layout_target(rels_path, cover_target)
            if verbose:
                print(
                    "  Baseline cover layout: "
                    f"{len(minority_slides)} slide(s) keep slide-local chrome"
                )

    if verbose and promoted:
        print(
            "  Baseline master chrome: "
            f"promoted {promoted} slide shape(s) across {promoted_roles} shared object(s)"
        )
    return promoted


_PAGE_NUMBER_TOKENS = {"pagenumber", "pagenum", "slidenumber"}


def _first_slide_number(extract_dir: Path) -> int:
    """Read firstSlideNum from presentation.xml (defaults to 1)."""
    presentation_path = extract_dir / "ppt" / "presentation.xml"
    try:
        root = ET.parse(presentation_path).getroot()
    except (OSError, ET.ParseError):
        return 1
    raw = root.attrib.get("firstSlideNum")
    if raw is None:
        return 1
    try:
        return int(raw)
    except ValueError:
        return 1


def _shape_with_id(root: ET.Element, shape_id: str) -> ET.Element | None:
    """Find a p:sp anywhere in the slide tree by its cNvPr id."""
    for shape in root.iter(f"{{{PML_NS}}}sp"):
        cnv = shape.find(f"{{{PML_NS}}}nvSpPr/{{{PML_NS}}}cNvPr")
        if cnv is not None and cnv.attrib.get("id") == shape_id:
            return shape
    return None


def _replace_literal_run_with_slidenum_field(
    shape: ET.Element,
    expected_text: str,
    field_guid: str,
) -> bool:
    """Swap a single literal page-number run for an a:fld slidenum field."""
    tx_body = shape.find(f"{{{PML_NS}}}txBody")
    if tx_body is None:
        return False
    a_t = f"{{{DML_NS}}}t"
    total_text = "".join(t.text or "" for t in tx_body.iter(a_t))
    if total_text.strip() != expected_text:
        return False
    text_runs = [
        (paragraph, run)
        for paragraph in tx_body.iter(f"{{{DML_NS}}}p")
        for run in paragraph.findall(f"{{{DML_NS}}}r")
        if (run.findtext(a_t) or "").strip()
    ]
    if len(text_runs) != 1:
        return False
    paragraph, run = text_runs[0]
    if (run.findtext(a_t) or "").strip() != expected_text:
        return False

    fld = ET.Element(f"{{{DML_NS}}}fld", {"id": field_guid, "type": "slidenum"})
    r_pr = run.find(f"{{{DML_NS}}}rPr")
    if r_pr is not None:
        fld.append(ET.fromstring(ET.tostring(r_pr, encoding="utf-8")))
    fld_text = ET.SubElement(fld, a_t)
    fld_text.text = expected_text
    index = list(paragraph).index(run)
    paragraph.remove(run)
    paragraph.insert(index, fld)
    return True


def _convert_page_number_texts_to_fields(
    extract_dir: Path,
    slide_count: int,
    conversion_traces: list[dict[str, Any]] | None,
    *,
    context: str = "Baseline",
    verbose: bool = False,
) -> int:
    """Replace literal page-number chrome text with auto-updating fields.

    Only converts when the traced pageNumber/slideNumber shape's whole text
    equals the slide's expected display number (honoring firstSlideNum), so
    schemes like content-only numbering keep their literal text untouched.
    """
    if not conversion_traces:
        return 0
    trace_by_slide = {
        int(trace.get("slide_num", 0)): trace
        for trace in conversion_traces
        if trace.get("slide_num") is not None
    }
    first_slide_number = _first_slide_number(extract_dir)
    field_guid = f"{{{str(uuid.uuid4()).upper()}}}"

    converted = 0
    for slide_num in range(1, slide_count + 1):
        tokens = _trace_chrome_shape_ids(trace_by_slide.get(slide_num))
        shape_ids = sorted({
            shape_id
            for token, ids in tokens.items()
            if token in _PAGE_NUMBER_TOKENS
            for shape_id in ids
        })
        if len(shape_ids) != 1:
            continue
        slide_path = extract_dir / "ppt" / "slides" / f"slide{slide_num}.xml"
        tree = ET.parse(slide_path)
        shape = _shape_with_id(tree.getroot(), shape_ids[0])
        if shape is None:
            continue
        expected_text = str(first_slide_number + slide_num - 1)
        if _replace_literal_run_with_slidenum_field(shape, expected_text, field_guid):
            _write_xml_tree(slide_path, tree)
            converted += 1

    if verbose and converted:
        print(
            f"  {context} slide-number fields: "
            f"converted {converted} page number(s)"
        )
    return converted


def _remove_relationship(rels_path: Path, rel_id: str) -> None:
    """Remove one relationship entry by rId."""
    rels_content = rels_path.read_text(encoding="utf-8")
    pattern = re.compile(
        rf'[ \t]*<Relationship\b[^>]*\bId="{re.escape(rel_id)}"[^>]*/>[ \t]*\n?'
    )
    new_content, removed = pattern.subn("", rels_content, count=1)
    if not removed:
        raise RuntimeError(f"Relationship {rel_id} not found in {rels_path}")
    rels_path.write_text(new_content, encoding="utf-8")


def _remove_content_type_override(content_types_path: Path, part_name: str) -> None:
    """Remove the Override content-type entry for a deleted package part."""
    normalized = "/" + part_name.lstrip("/")
    content = content_types_path.read_text(encoding="utf-8")
    pattern = re.compile(
        rf'[ \t]*<Override\b[^>]*\bPartName="{re.escape(normalized)}"[^>]*/>[ \t]*\n?'
    )
    new_content, removed = pattern.subn("", content, count=1)
    if removed:
        content_types_path.write_text(new_content, encoding="utf-8")


def _remove_trailing_layout_definition_slides(
    extract_dir: Path,
    public_slide_count: int,
    total_slide_count: int,
) -> int:
    """Remove internal carrier slides after their Layouts are registered."""
    if total_slide_count <= public_slide_count:
        return 0
    presentation_part = "ppt/presentation.xml"
    presentation_path = extract_dir / presentation_part
    presentation_rels = _relationships_path_for_part(
        extract_dir,
        presentation_part,
    )
    tree = ET.parse(presentation_path)
    root = tree.getroot()
    slide_list = root.find(f"{{{PML_NS}}}sldIdLst")
    if slide_list is None:
        raise RuntimeError("presentation.xml has no p:sldIdLst")
    entries_by_rel_id = {
        entry.get(f"{{{REL_NS}}}id", ""): entry
        for entry in slide_list.findall(f"{{{PML_NS}}}sldId")
    }
    relationships = _read_relationships(presentation_rels)
    content_types_path = extract_dir / "[Content_Types].xml"
    removed = 0
    for slide_num in range(public_slide_count + 1, total_slide_count + 1):
        slide_part = f"ppt/slides/slide{slide_num}.xml"
        rel_ids = [
            rel_id
            for rel_id, attrs in relationships.items()
            if attrs.get("Type") == SLIDE_REL_TYPE
            and _resolve_package_target(
                presentation_part,
                attrs.get("Target", ""),
            ) == slide_part
        ]
        if len(rel_ids) != 1:
            raise RuntimeError(
                f"Internal Layout carrier {slide_part} must have exactly one "
                "Presentation relationship"
            )
        rel_id = rel_ids[0]
        entry = entries_by_rel_id.get(rel_id)
        if entry is None:
            raise RuntimeError(
                f"presentation.xml has no p:sldId entry for {slide_part}"
            )
        slide_list.remove(entry)
        _remove_relationship(presentation_rels, rel_id)
        slide_path = extract_dir / slide_part
        slide_rels = _relationships_path_for_part(extract_dir, slide_part)
        if not slide_path.is_file() or not slide_rels.is_file():
            raise RuntimeError(
                f"Internal Layout carrier package parts are incomplete: {slide_part}"
            )
        slide_path.unlink()
        slide_rels.unlink()
        _remove_content_type_override(content_types_path, slide_part)
        removed += 1
    _write_xml_tree(presentation_path, tree)
    return removed


def _prune_unreferenced_definition_payload_parts(
    extract_dir: Path,
    *,
    preserved_parts: frozenset[str] = frozenset(),
) -> int:
    """Remove generated native/media payload left by deleted carrier slides.

    Unselected complete Slide prototypes are first converted as carrier slides
    so their reusable Layout can be registered. Removing those carriers may
    leave chart, workbook, or media parts with no remaining relationship. Run
    an iterative incoming-reference sweep so chart-owned workbooks/styles are
    removed after their orphan chart part and relationship sidecar disappear.
    """
    candidate_prefixes = (
        "ppt/charts/",
        "ppt/embeddings/",
        "ppt/media/",
    )
    content_types_path = extract_dir / "[Content_Types].xml"
    removed = 0
    while True:
        referenced_parts: set[str] = set()
        for rels_path in extract_dir.rglob("*.rels"):
            rels_rel = rels_path.relative_to(extract_dir).as_posix()
            try:
                root = ET.parse(rels_path).getroot()
            except ET.ParseError as exc:
                raise RuntimeError(
                    f"Invalid relationships XML while pruning {rels_rel}: {exc}"
                ) from exc
            for elem in root:
                attrs = _relationship_attrs(elem)
                if attrs.get("TargetMode", "").lower() == "external":
                    continue
                target = attrs.get("Target")
                if not target:
                    continue
                resolved = _resolve_internal_opc_target(rels_rel, target)
                if resolved is not None:
                    referenced_parts.add(resolved)

        orphan_paths: list[tuple[Path, str]] = []
        for path in extract_dir.rglob("*"):
            if not path.is_file():
                continue
            part_name = path.relative_to(extract_dir).as_posix()
            if "/_rels/" in part_name:
                continue
            if not part_name.startswith(candidate_prefixes):
                continue
            if part_name in preserved_parts:
                continue
            canonical = _canonical_opc_part_path(part_name)
            if canonical is not None and canonical not in referenced_parts:
                orphan_paths.append((path, part_name))
        if not orphan_paths:
            break

        for path, part_name in orphan_paths:
            rels_path = _relationships_path_for_part(extract_dir, part_name)
            if rels_path.is_file():
                rels_path.unlink()
            path.unlink()
            _remove_content_type_override(content_types_path, part_name)
            removed += 1
    return removed


def _prune_unused_slide_layouts(
    extract_dir: Path,
    structure: PptxStructureContext,
    slide_count: int,
    *,
    verbose: bool = False,
) -> int:
    """Remove base-template slide layouts no generated slide references.

    The python-pptx base package ships the full Office layout set; unused
    entries only pollute the PowerPoint new-slide picker. Layouts referenced
    by any generated slide are always kept, and a master keeps its layout
    list untouched unless at least one referenced layout remains in it.
    """
    # Read layout references live: earlier baseline passes may have rebound
    # minority slides to a Cover layout that the initial structure context
    # does not know about.
    referenced_layouts: set[str] = set()
    for slide_num in range(1, slide_count + 1):
        rels_path = (
            extract_dir / "ppt" / "slides" / "_rels" / f"slide{slide_num}.xml.rels"
        )
        target = _find_relationship_target(rels_path, SLIDE_LAYOUT_REL_TYPE)
        if not target:
            raise RuntimeError(f"Slide {slide_num} has no slide layout relationship")
        referenced_layouts.add(
            _resolve_package_target(f"ppt/slides/slide{slide_num}.xml", target)
        )

    pruned = 0
    content_types_path = extract_dir / "[Content_Types].xml"
    for master_part in sorted(set(structure.slide_master_parts.values())):
        master_path = extract_dir / master_part
        master_rels_path = _relationships_path_for_part(extract_dir, master_part)
        layout_rels = {
            rel_id: _resolve_package_target(master_part, attrs.get("Target", ""))
            for rel_id, attrs in _read_relationships(master_rels_path).items()
            if attrs.get("Type") == SLIDE_LAYOUT_REL_TYPE
        }
        if not any(part in referenced_layouts for part in layout_rels.values()):
            continue

        master_xml = master_path.read_text(encoding="utf-8")
        for rel_id, layout_part in sorted(layout_rels.items()):
            if layout_part in referenced_layouts:
                continue
            entry_re = re.compile(
                rf'[ \t]*<p:sldLayoutId\b[^>]*\br:id="{re.escape(rel_id)}"[^>]*/>[ \t]*\n?'
            )
            master_xml, removed = entry_re.subn("", master_xml, count=1)
            if not removed:
                raise RuntimeError(
                    f"Slide master {master_part} has no sldLayoutId entry for {rel_id}"
                )
            _remove_relationship(master_rels_path, rel_id)
            (extract_dir / layout_part).unlink()
            layout_rels_path = _relationships_path_for_part(extract_dir, layout_part)
            if layout_rels_path.exists():
                layout_rels_path.unlink()
            _remove_content_type_override(content_types_path, layout_part)
            pruned += 1
        master_path.write_text(master_xml, encoding="utf-8")

    if verbose and pruned:
        print(f"  Layout prune: removed {pruned} unused base layout(s)")
    return pruned


def _flat_structure_name(value: str | None) -> str:
    """Return one compact package identity for a free-design deck."""
    normalized = " ".join((value or "").split()).strip()
    return (normalized or "Free Design")[:120]


def _flat_placeholder_type(shape: ET.Element) -> str | None:
    """Return one system placeholder type carried by a top-level shape."""
    if shape.tag != f"{{{PML_NS}}}sp":
        return None
    placeholder = shape.find(
        f"{{{PML_NS}}}nvSpPr/{{{PML_NS}}}nvPr/{{{PML_NS}}}ph"
    )
    return placeholder.get("type") if placeholder is not None else None


def _clean_flat_structure_part(
    part_path: Path,
    name: str,
    *,
    is_layout: bool,
) -> tuple[int, tuple[str, ...]]:
    """Keep only standard footer hooks in one project-owned flat shell."""
    try:
        tree = ET.parse(part_path)
    except (OSError, ET.ParseError) as exc:
        raise RuntimeError(
            f"Cannot parse flat structure part {part_path}: {exc}"
        ) from exc
    root = tree.getroot()
    common_slide = root.find(f"{{{PML_NS}}}cSld")
    if common_slide is None:
        raise RuntimeError(f"Flat structure part has no p:cSld: {part_path}")
    shape_tree = common_slide.find(f"{{{PML_NS}}}spTree")
    if shape_tree is None:
        raise RuntimeError(f"Flat structure part has no p:spTree: {part_path}")

    removed = 0
    retained: list[str] = []
    for child in list(shape_tree):
        if child.tag not in _TOP_LEVEL_SHAPE_TAGS:
            continue
        placeholder_type = _flat_placeholder_type(child)
        if placeholder_type in _FLAT_SYSTEM_PLACEHOLDER_TYPES:
            retained.append(placeholder_type)
            continue
        shape_tree.remove(child)
        removed += 1
    for parent in (common_slide, root):
        for extension_list in parent.findall(f"{{{PML_NS}}}extLst"):
            parent.remove(extension_list)

    common_slide.set("name", name)
    if is_layout:
        root.set("type", "blank")
        root.set("preserve", "1")
    _write_xml_tree(part_path, tree)
    return removed, tuple(retained)


def _name_flat_themes(extract_dir: Path, name: str) -> int:
    """Replace stock Office theme identities with the current deck identity."""
    theme_paths = sorted((extract_dir / "ppt" / "theme").glob("theme*.xml"))
    if not theme_paths:
        raise RuntimeError("Flat PPTX package has no theme part")
    for theme_path in theme_paths:
        try:
            tree = ET.parse(theme_path)
        except (OSError, ET.ParseError) as exc:
            raise RuntimeError(
                f"Cannot parse flat theme {theme_path}: {exc}"
            ) from exc
        root = tree.getroot()
        root.set("name", name)
        for tag in ("clrScheme", "fontScheme", "fmtScheme"):
            scheme = root.find(f".//{{{DML_NS}}}{tag}")
            if scheme is not None:
                scheme.set("name", name)
        _write_xml_tree(theme_path, tree)
    return len(theme_paths)


def _prepare_flat_structure(
    extract_dir: Path,
    structure: PptxStructureContext,
    slide_count: int,
    master_text_style_spec: MasterTextStyleSpec | None,
    structure_name: str | None,
    *,
    verbose: bool = False,
) -> None:
    """Materialize one clean current-deck Master and Blank Layout for flat export."""
    pruned = _prune_unused_slide_layouts(
        extract_dir,
        structure,
        slide_count,
        verbose=False,
    )
    live_structure = _read_slide_layout_targets(extract_dir, slide_count)
    layout_parts = {
        _resolve_package_target(
            f"ppt/slides/slide{slide_num}.xml",
            live_structure.slide_layout_target(slide_num),
        )
        for slide_num in range(1, slide_count + 1)
    }
    master_parts = set(live_structure.slide_master_parts.values())
    physical_layouts = {
        str(path.relative_to(extract_dir)).replace("\\", "/")
        for path in (extract_dir / "ppt" / "slideLayouts").glob("slideLayout*.xml")
    }
    physical_masters = {
        str(path.relative_to(extract_dir)).replace("\\", "/")
        for path in (extract_dir / "ppt" / "slideMasters").glob("slideMaster*.xml")
    }
    if len(layout_parts) != 1 or layout_parts != physical_layouts:
        raise RuntimeError(
            "Flat export must retain exactly one slide-referenced Blank Layout"
        )
    if len(master_parts) != 1 or master_parts != physical_masters:
        raise RuntimeError(
            "Flat export must retain exactly one slide-referenced Master"
        )

    identity = _flat_structure_name(structure_name)
    theme_name = identity
    master_name = f"{identity} — Master"
    layout_name = f"{identity} — Blank"
    master_path = extract_dir / next(iter(master_parts))
    layout_path = extract_dir / next(iter(layout_parts))
    removed_master_shapes, retained_master_placeholders = _clean_flat_structure_part(
        master_path,
        master_name,
        is_layout=False,
    )
    removed_layout_shapes, retained_layout_placeholders = _clean_flat_structure_part(
        layout_path,
        layout_name,
        is_layout=True,
    )
    theme_count = _name_flat_themes(extract_dir, theme_name)
    master_count = (
        apply_master_text_style_spec(extract_dir, master_text_style_spec)
        if master_text_style_spec is not None
        else 0
    )

    for part_path, expected_name, expected_layout in (
        (master_path, master_name, False),
        (layout_path, layout_name, True),
    ):
        root = ET.parse(part_path).getroot()
        common_slide = root.find(f"{{{PML_NS}}}cSld")
        if common_slide is None or common_slide.get("name") != expected_name:
            raise RuntimeError(
                f"Flat structure identity read-back failed: {part_path}"
            )
        shape_tree = common_slide.find(f"{{{PML_NS}}}spTree")
        if shape_tree is None:
            raise RuntimeError(f"Flat structure shell has no shape tree: {part_path}")
        actual_placeholder_types = tuple(
            sorted(
                _flat_placeholder_type(child) or ""
                for child in shape_tree
                if child.tag in _TOP_LEVEL_SHAPE_TAGS
            )
        )
        expected_placeholder_types = tuple(
            sorted(_FLAT_SYSTEM_PLACEHOLDER_TYPES)
        )
        if actual_placeholder_types != expected_placeholder_types:
            raise RuntimeError(
                "Flat structure shell must retain only one each of the date, "
                f"footer, and slide-number hooks: {part_path}"
            )
        if expected_layout and root.get("type") != "blank":
            raise RuntimeError(f"Flat layout is not typed as Blank: {part_path}")

    if verbose:
        print(
            "  Flat structure: project-owned Master + Blank Layout "
            f"({pruned} stock layout(s), "
            f"{removed_master_shapes + removed_layout_shapes} stock content "
            "shape(s) removed, "
            f"{len(retained_master_placeholders) + len(retained_layout_placeholders)} "
            "system footer hook(s) retained)"
        )
        text_style_status = (
            f"{master_count} master text style(s)"
            if master_text_style_spec is not None
            else "stock text defaults retained (no theme contract)"
        )
        print(f"  Flat theme: {theme_count} theme part(s), {text_style_status}")


def _append_relationship(
    rels_path: Path,
    rel_type: str,
    target: str,
    *,
    target_mode: str | None = None,
) -> str:
    """Append a relationship entry with the next available rId."""
    with open(rels_path, 'r', encoding='utf-8') as f:
        rels_content = f.read()

    rid_numbers = [int(match) for match in re.findall(r'Id="rId(\d+)"', rels_content)]
    next_rid = f'rId{max(rid_numbers, default=0) + 1}'
    mode_attr = (
        f" TargetMode={quoteattr(target_mode)}"
        if target_mode is not None
        else ""
    )
    rel_xml = (
        f"  <Relationship Id={quoteattr(next_rid)} "
        f"Type={quoteattr(rel_type)} Target={quoteattr(target)}{mode_attr}/>"
    )
    rels_content = rels_content.replace(
        '</Relationships>', rel_xml + '\n</Relationships>',
    )

    with open(rels_path, 'w', encoding='utf-8') as f:
        f.write(rels_content)

    return next_rid


def _add_default_content_type(content_types: str, extension: str, content_type: str) -> str:
    """Add a Default content type if it is not already present."""
    ext = extension.lstrip(".")
    if f'Extension="{ext}"' in content_types:
        return content_types
    entry = f'  <Default Extension="{ext}" ContentType="{content_type}"/>'
    override_pos = content_types.find('<Override ')
    if override_pos >= 0:
        return content_types[:override_pos] + entry + '\n' + content_types[override_pos:]
    return content_types.replace('</Types>', entry + '\n</Types>')


def _add_content_type_override(content_types: str, part_name: str, content_type: str) -> str:
    """Add an Override content type if it is not already present."""
    normalized = '/' + part_name.lstrip('/')
    if f'PartName="{normalized}"' in content_types:
        return content_types
    extension = PurePosixPath(part_name).suffix.lstrip('.').casefold()
    try:
        root = ET.fromstring(content_types)
    except ET.ParseError:
        root = None
    if root is not None:
        for child in root:
            if child.tag.rsplit('}', 1)[-1] != 'Default':
                continue
            if (
                child.get('Extension', '').casefold() == extension
                and child.get('ContentType') == content_type
            ):
                return content_types
    entry = f'  <Override PartName="{normalized}" ContentType="{content_type}"/>'
    return content_types.replace('</Types>', entry + '\n</Types>')


def _content_type_contract(payload: bytes) -> tuple[tuple[str, tuple[tuple[str, str], ...]], ...] | None:
    """Return the semantic Default/Override roster for exact-byte restoration."""
    try:
        root = ET.fromstring(payload)
    except ET.ParseError:
        return None
    entries = [
        (
            child.tag.rsplit('}', 1)[-1],
            tuple(sorted(child.attrib.items())),
        )
        for child in root
        if child.tag.rsplit('}', 1)[-1] in {'Default', 'Override'}
    ]
    return tuple(sorted(entries))


_IMAGE_CONTENT_TYPES = {
    'png': 'image/png',
    'jpg': 'image/jpeg',
    'jpeg': 'image/jpeg',
    'gif': 'image/gif',
    'webp': 'image/webp',
    'svg': 'image/svg+xml',
    'bmp': 'image/bmp',
    'emf': 'image/x-emf',
    'tif': 'image/tiff',
    'tiff': 'image/tiff',
    'wmf': 'image/x-wmf',
}


def _content_type_for_extension(ext: str) -> str:
    clean = ext.lower().lstrip('.')
    content_type = _IMAGE_CONTENT_TYPES.get(clean) or mimetypes.guess_type(f'x.{clean}')[0]
    if not content_type:
        raise ValueError(f"Unknown media content type for extension: {ext}")
    return content_type


def _as_dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _create_writable_work_dir(output_path: Path) -> Path:
    """Create a real writable work directory for PPTX assembly."""
    parents = [output_path.parent, Path.cwd(), Path(tempfile.gettempdir())]
    seen: set[str] = set()
    errors: list[str] = []

    for parent in parents:
        parent = parent if str(parent) else Path(".")
        try:
            key = str(parent.resolve())
        except OSError:
            key = str(parent.absolute())
        if key in seen:
            continue
        seen.add(key)

        try:
            parent.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            errors.append(f"{parent}: cannot create parent ({exc})")
            continue

        for _ in range(3):
            work_dir = parent / f".pptx-build-{os.getpid()}-{uuid.uuid4().hex}"
            try:
                work_dir.mkdir(mode=0o700)
                probe_path = work_dir / ".write-probe"
                probe_path.write_text("ok", encoding="utf-8")
                probe_path.unlink()
                return work_dir
            except OSError as exc:
                errors.append(f"{work_dir}: {exc}")
                shutil.rmtree(work_dir, ignore_errors=True)

    details = "\n  - ".join(errors) if errors else "no candidate directories available"
    raise PermissionError(
        "Unable to create a writable PPTX work directory. "
        "Set the output path to a writable project directory or adjust sandbox permissions. "
        f"Tried:\n  - {details}"
    )


def _relax_output_permissions(output_path: Path) -> list[str]:
    """Make exported files readable outside the sandbox owner where possible."""
    warnings: list[str] = []

    try:
        current_mode = output_path.stat().st_mode
        readable_mode = (
            current_mode
            | stat.S_IRUSR
            | stat.S_IWUSR
            | stat.S_IRGRP
            | stat.S_IROTH
        )
        os.chmod(output_path, readable_mode)
    except OSError as exc:
        warnings.append(f"chmod skipped for {output_path}: {exc}")

    if os.name != 'nt':
        return warnings

    # Windows ACLs can remain sandbox-only even when the file mode looks sane.
    # Grant the built-in Users SID read access; the SID avoids localization
    # issues on non-English Windows installations.
    try:
        result = subprocess.run(
            ['icacls', str(output_path), '/grant', '*S-1-5-32-545:R'],
            capture_output=True,
            text=True,
            check=False,
        )
    except OSError as exc:
        warnings.append(f"icacls skipped for {output_path}: {exc}")
    else:
        if result.returncode != 0:
            message = (result.stderr or result.stdout or '').strip()
            details = f": {message}" if message else ''
            warnings.append(f"icacls failed for {output_path}{details}")

    return warnings


_NOTES_MASTER_REL_TYPE = (
    'http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesMaster'
)
_NOTES_MASTER_CONTENT_TYPE = (
    'application/vnd.openxmlformats-officedocument.presentationml.notesMaster+xml'
)


def _content_type_overrides(extract_dir: Path) -> dict[str, str]:
    """Return package part content-type overrides keyed without a leading slash."""
    content_types_path = extract_dir / '[Content_Types].xml'
    try:
        root = ET.parse(content_types_path).getroot()
    except (OSError, ET.ParseError) as exc:
        raise RuntimeError(
            f'Cannot read package content types: {content_types_path}'
        ) from exc
    overrides: dict[str, str] = {}
    for child in root:
        if child.tag.rsplit('}', 1)[-1] != 'Override':
            continue
        part_name = child.get('PartName', '').lstrip('/')
        content_type = child.get('ContentType', '')
        if part_name:
            overrides[part_name] = content_type
    return overrides


def _notes_master_part_from_target(target: str) -> str | None:
    """Resolve one presentation relationship target to a notes-master part."""
    part_name = _resolve_package_target(
        'ppt/presentation.xml',
        target,
    ).lstrip('/')
    prefix = 'ppt/notesMasters/'
    if not part_name.startswith(prefix) or part_name == prefix:
        return None
    return part_name


def _validate_notes_master_part(
    extract_dir: Path,
    part_name: str,
    content_types: dict[str, str],
) -> str:
    """Validate one existing notes-master package part and return its name."""
    normalized = posixpath.normpath(part_name.lstrip('/'))
    if not normalized.startswith('ppt/notesMasters/'):
        raise RuntimeError(
            f'Notes master resolves outside ppt/notesMasters: {part_name}'
        )
    part_path = extract_dir / PurePosixPath(normalized)
    if not part_path.is_file():
        raise RuntimeError(f'Notes master part is missing: {normalized}')
    try:
        root = ET.parse(part_path).getroot()
    except (OSError, ET.ParseError) as exc:
        raise RuntimeError(f'Cannot parse notes master part: {normalized}') from exc
    if root.tag != f'{{{PML_NS}}}notesMaster':
        raise RuntimeError(f'Notes master part has an invalid root: {normalized}')
    declared_type = content_types.get(normalized)
    if declared_type is not None and declared_type != _NOTES_MASTER_CONTENT_TYPE:
        raise RuntimeError(
            'Notes master content type mismatch for '
            f'{normalized}: {declared_type}'
        )
    return normalized


def _register_notes_master_id(
    presentation_path: Path,
    relationship_id: str,
) -> None:
    """Register a notes-master relationship in presentation.xml by XML identity."""
    try:
        tree = ET.parse(presentation_path)
    except (OSError, ET.ParseError) as exc:
        raise RuntimeError(
            f'Cannot parse presentation package part: {presentation_path}'
        ) from exc
    root = tree.getroot()
    list_tag = f'{{{PML_NS}}}notesMasterIdLst'
    id_tag = f'{{{PML_NS}}}notesMasterId'
    relationship_attr = f'{{{REL_NS}}}id'
    notes_master_list = root.find(list_tag)
    if notes_master_list is not None:
        if any(
            child.tag == id_tag
            and child.get(relationship_attr) == relationship_id
            for child in notes_master_list
        ):
            return
    else:
        slide_master_list = root.find(f'{{{PML_NS}}}sldMasterIdLst')
        if slide_master_list is None:
            raise RuntimeError(
                'presentation.xml is missing p:sldMasterIdLst'
            )
        notes_master_list = ET.Element(list_tag)
        insert_at = list(root).index(slide_master_list) + 1
        root.insert(insert_at, notes_master_list)
    ET.SubElement(
        notes_master_list,
        id_tag,
        {relationship_attr: relationship_id},
    )
    _write_xml_tree(presentation_path, tree)


def _ensure_notes_master(
    extract_dir: Path,
    primary_language: str | None = None,
) -> _NotesMasterReference:
    """Reuse or create a notes master and wire it into the presentation package."""
    ppt_dir = extract_dir / 'ppt'
    notes_masters_dir = ppt_dir / 'notesMasters'
    notes_masters_dir.mkdir(exist_ok=True)
    presentation_rels_path = ppt_dir / '_rels' / 'presentation.xml.rels'
    presentation_path = ppt_dir / 'presentation.xml'
    content_types = _content_type_overrides(extract_dir)
    relationships = _read_relationships(presentation_rels_path)

    try:
        presentation_root = ET.parse(presentation_path).getroot()
    except (OSError, ET.ParseError) as exc:
        raise RuntimeError(
            f'Cannot parse presentation package part: {presentation_path}'
        ) from exc
    relationship_attr = f'{{{REL_NS}}}id'
    referenced_ids = [
        elem.get(relationship_attr, '')
        for elem in presentation_root.findall(
            f'./{{{PML_NS}}}notesMasterIdLst/{{{PML_NS}}}notesMasterId'
        )
    ]
    notes_relationship_ids = [
        rel_id
        for rel_id, attrs in relationships.items()
        if attrs.get('Type') == _NOTES_MASTER_REL_TYPE
        and attrs.get('TargetMode') != 'External'
    ]
    ordered_relationship_ids = list(
        dict.fromkeys(referenced_ids + notes_relationship_ids)
    )
    relationship_id_by_part: dict[str, str] = {}
    selected_part: str | None = None
    selected_rid: str | None = None
    for rel_id in ordered_relationship_ids:
        attrs = relationships.get(rel_id)
        if (
            not attrs
            or attrs.get('Type') != _NOTES_MASTER_REL_TYPE
            or attrs.get('TargetMode') == 'External'
        ):
            continue
        target = attrs.get('Target', '')
        part_name = _notes_master_part_from_target(target)
        if part_name is None:
            raise RuntimeError(
                f'Notes master relationship {rel_id} has an invalid target: {target}'
            )
        part_name = _validate_notes_master_part(
            extract_dir,
            part_name,
            content_types,
        )
        relationship_id_by_part.setdefault(part_name, rel_id)
        if selected_part is None:
            selected_part = part_name
            selected_rid = rel_id

    if selected_part is None:
        for part_name, content_type in sorted(content_types.items()):
            if content_type != _NOTES_MASTER_CONTENT_TYPE:
                continue
            selected_part = _validate_notes_master_part(
                extract_dir,
                part_name,
                content_types,
            )
            selected_rid = relationship_id_by_part.get(selected_part)
            break

    created_theme_part: str | None = None
    if selected_part is None:
        occupied_parts = set(content_types)
        occupied_parts.update(
            part_name
            for attrs in relationships.values()
            if attrs.get('Type') == _NOTES_MASTER_REL_TYPE
            and (part_name := _notes_master_part_from_target(
                attrs.get('Target', '')
            )) is not None
        )
        index = 1
        while True:
            selected_part = f'ppt/notesMasters/notesMaster{index}.xml'
            if (
                selected_part not in occupied_parts
                and not (extract_dir / selected_part).exists()
            ):
                break
            index += 1
        notes_master_path = extract_dir / selected_part
        notes_master_path.write_text(
            create_notes_master_xml(primary_language),
            encoding='utf-8',
        )

        theme_dir = ppt_dir / 'theme'
        theme_dir.mkdir(exist_ok=True)
        theme1_path = theme_dir / 'theme1.xml'
        theme2_path = theme_dir / 'theme2.xml'
        if not theme2_path.exists():
            if theme1_path.exists():
                shutil.copy2(theme1_path, theme2_path)
            else:
                raise RuntimeError(
                    'Cannot create notes theme: ppt/theme/theme1.xml is missing'
                )
        created_theme_part = 'ppt/theme/theme2.xml'

        notes_master_rels_path = _relationships_path_for_part(
            extract_dir,
            selected_part,
        )
        notes_master_rels_path.parent.mkdir(parents=True, exist_ok=True)
        notes_master_rels_path.write_text(
            create_notes_master_rels_xml(),
            encoding='utf-8',
        )

    if selected_rid is None:
        selected_rid = _append_relationship(
            presentation_rels_path,
            _NOTES_MASTER_REL_TYPE,
            posixpath.relpath(selected_part, 'ppt'),
        )
    _register_notes_master_id(presentation_path, selected_rid)
    return _NotesMasterReference(
        package_part=selected_part,
        created_theme_part=created_theme_part,
    )


def _slide_config(animation_config: dict[str, Any] | None, svg_stem: str) -> dict[str, Any]:
    if not animation_config:
        return {}
    slides_value = animation_config.get('slides', {})
    if not isinstance(slides_value, dict):
        raise ValueError('animations.json field "slides" must be an object')
    slide_value = slides_value.get(svg_stem, {})
    if not isinstance(slide_value, dict):
        raise ValueError(
            f'animations.json slide "{svg_stem}" must be an object'
        )
    return slide_value


def _slide_transition_settings(
    default_transition_cfg: dict[str, Any],
    slide_cfg: dict[str, Any],
    transition: str | None,
    transition_effect_options: dict[str, object] | None,
    duration: float,
    auto_advance: float | None,
    transition_sound: str | None,
    cli_overrides: dict[str, bool],
) -> tuple[str | None, dict[str, object], float, float | None, str | None]:
    trans_value = slide_cfg.get('transition', {})
    if not isinstance(trans_value, dict):
        raise ValueError('animations.json slide transition must be an object')
    trans_cfg = trans_value
    effect, effect_options = normalize_transition_effect_request(
        transition,
        transition_effect_options,
    )
    if not cli_overrides.get('transition'):
        if 'effect' in trans_cfg:
            raw_effect = trans_cfg['effect']
            raw_options = trans_cfg.get('effect_options')
            effect, effect_options = normalize_transition_effect_request(
                raw_effect,
                raw_options,
            )
        elif 'effect_options' in trans_cfg:
            raise ValueError(
                'animations.json transition effect_options requires '
                'an explicit effect'
            )
    if not cli_overrides.get('transition_duration'):
        if 'duration' in trans_cfg:
            duration = validate_seconds(
                trans_cfg.get('duration'),
                "transition duration",
                allow_zero=effect is None,
            )
    if not cli_overrides.get('auto_advance') and 'auto_advance' in trans_cfg:
        auto_advance = validate_seconds(
            trans_cfg.get('auto_advance'),
            "transition auto_advance",
            allow_zero=True,
        )
    raw_sound = transition_sound
    if raw_sound is None and not cli_overrides.get('transition_sound'):
        raw_sound = default_transition_cfg.get('sound')
    if 'sound' in trans_cfg:
        raw_sound = trans_cfg['sound']
    if raw_sound is not None and (
        not isinstance(raw_sound, str) or not raw_sound.strip()
    ):
        raise ValueError(
            'animations.json transition sound must be a non-empty '
            'project-relative .wav path or null'
        )
    return effect, effect_options, duration, auto_advance, raw_sound


def _slide_animation_settings(
    slide_cfg: dict[str, Any],
    default_animation_cfg: dict[str, Any],
    animation: str | None,
    duration: float,
    stagger: float,
    trigger: str,
    cli_overrides: dict[str, bool],
) -> tuple[str | None, float, float, str, dict[str, Any]]:
    anim_value = slide_cfg.get('animation', {})
    if not isinstance(anim_value, dict):
        raise ValueError('animations.json slide animation must be an object')
    anim_cfg = anim_value
    resolved_cfg = resolve_slide_animation_config(
        default_animation_cfg,
        anim_cfg,
    )
    if cli_overrides.get('animation'):
        effect, effect_options = normalize_animation_effect_request(
            animation,
            allow_none=True,
            allow_modes=True,
        )
        resolved_cfg['effect'] = effect or 'none'
        if effect_options:
            resolved_cfg['effect_options'] = effect_options
        else:
            resolved_cfg.pop('effect_options', None)
    else:
        raw_effect = resolved_cfg.get('effect', animation)
        effect, effect_options = normalize_animation_effect_request(
            raw_effect,
            resolved_cfg.get('effect_options'),
            allow_none=True,
            allow_modes=True,
        )
        resolved_cfg['effect'] = effect or 'none'
        if effect_options:
            resolved_cfg['effect_options'] = effect_options
        else:
            resolved_cfg.pop('effect_options', None)
    if not cli_overrides.get('animation_duration'):
        duration = validate_seconds(
            anim_cfg.get('duration', duration),
            'animation duration',
            allow_zero=False,
        )
    else:
        resolved_cfg['duration'] = duration
    if not cli_overrides.get('animation_stagger'):
        stagger = validate_seconds(
            anim_cfg.get('stagger', stagger),
            'animation stagger',
            allow_zero=True,
        )
    else:
        resolved_cfg['stagger'] = stagger
    if not cli_overrides.get('animation_trigger') and 'trigger' in anim_cfg:
        trigger = normalize_animation_trigger(anim_cfg.get('trigger'))
    else:
        trigger = normalize_animation_trigger(trigger)
        resolved_cfg['trigger'] = trigger
    animation_seconds_to_milliseconds(
        duration,
        'animation duration',
        allow_zero=False,
    )
    animation_seconds_to_milliseconds(
        stagger,
        'animation stagger',
        allow_zero=True,
    )
    resolved_cfg['effect'] = effect or 'none'
    resolved_cfg['duration'] = duration
    resolved_cfg['stagger'] = stagger
    resolved_cfg['trigger'] = trigger
    return effect, duration, stagger, trigger, resolved_cfg


def _build_sequence_targets(
    anim_targets: list[tuple[int, str]],
    slide_name: str,
    slide_cfg: dict[str, Any],
    animation: str | None,
    animation_cfg: dict[str, Any],
    duration: float,
    stagger: float,
    mixed_animation_offset: int,
    animation_rng: random.Random,
) -> tuple[list[dict[str, Any]], int]:
    groups_value = slide_cfg.get('groups', {})
    if not isinstance(groups_value, dict):
        raise ValueError('animations.json slide groups must be an object')
    groups_cfg = groups_value
    shape_ids_by_group = {
        svg_id: sid for sid, svg_id in anim_targets
    }
    ordered: list[tuple[int, int, int, str, str, dict[str, Any]]] = []
    # A group without a sidecar ``order`` follows the nearest listed group
    # before it in SVG order (0 before the first one). Numbering unlisted
    # groups by raw SVG index collided with explicit orders: a headline at
    # index 2 landed behind the body block the author had numbered 1 and 2.
    inherited_orders: list[int] = []
    current_order = 0
    for _sid, svg_id in anim_targets:
        group_value = groups_cfg.get(svg_id, {})
        explicit = [
            cfg.get('order')
            for _path, cfg in (
                animation_group_effect_entries(group_value, path='')
                if isinstance(group_value, dict) else []
            )
            if isinstance(cfg.get('order'), int)
            and not isinstance(cfg.get('order'), bool)
        ]
        if explicit:
            current_order = max(current_order, max(explicit))
        inherited_orders.append(current_order)
    for idx, (sid, svg_id) in enumerate(anim_targets):
        group_value = groups_cfg.get(svg_id, {})
        if not isinstance(group_value, dict):
            raise ValueError(
                f'animations.json group "{svg_id}" must be an object'
            )
        group_path = (
            f'slides[{json.dumps(slide_name, ensure_ascii=False)}]'
            f'.groups[{json.dumps(svg_id, ensure_ascii=False)}]'
        )
        effect_entries = animation_group_effect_entries(
            group_value,
            path=group_path,
        )
        for effect_idx, (effect_path, effect_cfg) in enumerate(effect_entries):
            raw_effect = effect_cfg.get('effect')
            if raw_effect is not None:
                normalized_effect = normalize_animation_effect(
                    raw_effect,
                    allow_none=True,
                    allow_modes=True,
                )
            else:
                normalized_effect = None
            if 'effect' in effect_cfg and normalized_effect is None:
                continue
            if animation is None and normalized_effect is None:
                continue
            order_value = effect_cfg.get('order')
            order = order_value if order_value is not None else inherited_orders[idx]
            if order_value is not None and (
                isinstance(order, bool)
                or not isinstance(order, int)
                or order <= 0
            ):
                raise ValueError(
                    f'animations.json {effect_path}.order must be '
                    'a positive integer'
                )
            effect_entry = dict(effect_cfg)
            effect_entry['_shape_id'] = sid
            effect_entry['_effect'] = normalized_effect
            effect_entry['_effect_raw'] = raw_effect
            ordered.append(
                (
                    order,
                    idx,
                    effect_idx,
                    svg_id,
                    effect_path,
                    effect_entry,
                )
            )

    ordered.sort(key=lambda item: (item[0], item[1], item[2]))

    seq_targets: list[dict[str, Any]] = []
    resolved_group_modes: list[str | None] = []
    main_sequence_count = 0
    for seq_idx, (
        _order,
        _original_idx,
        _effect_idx,
        _svg_id,
        effect_path,
        group_cfg,
    ) in enumerate(ordered):
        shape_id = int(group_cfg['_shape_id'])
        raw_effect = group_cfg.get('_effect')
        resolved_group_modes.append(
            raw_effect if raw_effect in ('auto', 'mixed', 'random') else None
        )
        if raw_effect in ('auto', 'mixed', 'random'):
            effect = pick_animation_effect(
                str(raw_effect), seq_idx, mixed_animation_offset, group_id=_svg_id,
                rng=animation_rng,
            )
            effect_options: dict[str, object] = {}
        else:
            effect = str(raw_effect or pick_animation_effect(
                animation, seq_idx, mixed_animation_offset, group_id=_svg_id,
                rng=animation_rng,
            ))
            request_effect = (
                group_cfg.get('_effect_raw')
                if group_cfg.get('_effect_raw') is not None
                else effect
            )
            option_value = (
                group_cfg.get('effect_options')
                if group_cfg.get('_effect_raw') is not None
                else animation_cfg.get('effect_options')
            )
            effect, effect_options = normalize_animation_effect_request(
                request_effect,
                option_value,
                allow_none=False,
                allow_modes=False,
            )
        item_duration = validate_seconds(
            group_cfg.get('duration', duration),
            f'animations.json {effect_path}.duration',
            allow_zero=False,
        )
        trigger_shape = group_cfg.get('trigger_shape')
        raw_trigger = group_cfg.get(
            'trigger',
            animation_cfg.get('trigger', 'after-previous'),
        )
        resolved_trigger = normalize_animation_trigger(raw_trigger)
        if trigger_shape is not None:
            if 'trigger' in group_cfg and resolved_trigger != 'on-click':
                raise ValueError(
                    f'animations.json {effect_path}.trigger_shape requires '
                    'trigger "on-click" when trigger is explicit'
                )
            resolved_trigger = 'on-click'
        default_delay = (
            stagger
            if (
                trigger_shape is None
                and resolved_trigger == 'after-previous'
                and main_sequence_count > 0
            )
            else 0
        )
        delay_seconds = validate_seconds(
            group_cfg.get('delay', default_delay),
            f'animations.json {effect_path}.delay',
            allow_zero=True,
        )
        delay_ms = animation_seconds_to_milliseconds(
            delay_seconds,
            f'animations.json {effect_path}.delay',
            allow_zero=True,
        )
        inherited_fields = {
            field: animation_cfg[field]
            for field in (
                *ANIMATION_TIMING_OPTION_FIELDS,
                'after_effect',
                'sound',
            )
            if field in animation_cfg
        }
        inherited_fields.update(
            {
                field: group_cfg[field]
                for field in (
                    *ANIMATION_TIMING_OPTION_FIELDS,
                    'after_effect',
                    'sound',
                )
                if field in group_cfg
            }
        )
        target_entry: dict[str, Any] = {
            'shape_id': shape_id,
            'delay_ms': delay_ms,
            'effect': effect,
            'effect_options': effect_options,
            'duration': item_duration,
            'trigger': resolved_trigger,
        }
        if trigger_shape is not None:
            if not isinstance(trigger_shape, str) or not trigger_shape.strip():
                raise ValueError(
                    f'animations.json {effect_path}.trigger_shape must '
                    'be a non-empty group id'
                )
            trigger_shape_id = shape_ids_by_group.get(trigger_shape)
            if trigger_shape_id is None:
                raise ValueError(
                    f'animations.json {effect_path}.trigger_shape '
                    f'references a missing or non-triggerable group: '
                    f'{trigger_shape}'
                )
            if trigger_shape_id == shape_id:
                raise ValueError(
                    f'animations.json {effect_path}.trigger_shape must '
                    'reference a different group'
                )
            target_entry['trigger_shape_id'] = trigger_shape_id
        else:
            main_sequence_count += 1
        target_entry.update(inherited_fields)
        if 'sound' in target_entry:
            target_entry['_sound_path'] = target_entry.pop('sound')
        seq_targets.append(target_entry)

    mixed_count = 0
    if animation == 'mixed':
        mixed_count = sum(1 for _target in seq_targets[1:])
    elif animation == 'auto':
        # 'auto' accumulates a cross-slide offset so the image pool and the
        # unmatched-id fallback rotate as the deck advances. Single-effect
        # semantic matches (title→entrance_fade, chart→entrance_wipe, etc.)
        # are unaffected
        # because they ignore the offset.
        mixed_count = len(seq_targets)
    else:
        mixed_count = sum(
            1
            for seq_idx, mode in enumerate(resolved_group_modes)
            if mode == 'auto' or (mode == 'mixed' and seq_idx > 0)
        )
    return seq_targets, mixed_count


def _next_relationship_id(rel_entries: list[dict[str, str]]) -> str:
    """Return the next slide relationship id, keeping rId1 for the layout."""
    used = {1}
    for rel in rel_entries:
        match = re.fullmatch(r'rId(\d+)', str(rel.get('id', '')))
        if match:
            used.add(int(match.group(1)))
    candidate = 2
    while candidate in used:
        candidate += 1
    return f'rId{candidate}'


def _materialize_slide_sound(
    project_path: Path,
    raw_sound: str,
    media_files: dict[str, bytes],
    rel_entries: list[dict[str, str]],
    audio_exts_used: set[str],
    packaged_by_source: dict[Path, tuple[str, str]],
    *,
    label: str,
    media_prefix: str,
    require_project_relative_wav: bool,
) -> dict[str, str]:
    """Package one slide sound and return its relationship descriptor."""
    if not isinstance(raw_sound, str) or not raw_sound.strip():
        raise ValueError(f'{label} sound must be a non-empty path string')
    sound_path = Path(raw_sound)
    if require_project_relative_wav:
        if sound_path.is_absolute() or PureWindowsPath(raw_sound).drive:
            raise ValueError(f'{label} sound must be project-relative: {raw_sound!r}')
        extension = sound_path.suffix.lower()
        if extension != '.wav':
            raise ValueError(f'{label} sound must use .wav')
        project_root = project_path.resolve()
        sound_path = (project_root / sound_path).resolve()
        try:
            sound_path.relative_to(project_root)
        except ValueError as exc:
            raise ValueError(
                f'{label} sound escapes the project root: {raw_sound!r}'
            ) from exc
    else:
        if not sound_path.is_absolute():
            sound_path = project_path / sound_path
        sound_path = sound_path.resolve()
        extension = sound_path.suffix.lower()

    if not sound_path.is_file():
        raise ValueError(f'{label} sound file not found: {sound_path}')
    if extension not in AUDIO_CONTENT_TYPES:
        valid = ', '.join(sorted(AUDIO_CONTENT_TYPES))
        raise ValueError(
            f'unsupported {label} sound format {extension or "(none)"}; '
            f'valid formats: {valid}'
        )

    packaged = packaged_by_source.get(sound_path)
    if packaged is None:
        payload = sound_path.read_bytes()
        if require_project_relative_wav and not (
            len(payload) >= 12
            and payload[:4] in {b'RIFF', b'RF64'}
            and payload[8:12] == b'WAVE'
        ):
            raise ValueError(f'{label} sound is not a valid WAV file: {sound_path}')
        digest = hashlib.sha256(payload).hexdigest()[:16]
        media_name = f'{media_prefix}_{digest}{extension}'
        relationship_id = _next_relationship_id(rel_entries)
        media_files.setdefault(media_name, payload)
        rel_entries.append(
            {
                'id': relationship_id,
                'type': AUDIO_REL_TYPE,
                'target': f'../media/{media_name}',
            }
        )
        packaged = (relationship_id, media_name)
        packaged_by_source[sound_path] = packaged
        audio_exts_used.add(extension)

    relationship_id, _media_name = packaged
    return {
        'relationship_id': relationship_id,
        'name': sound_path.name,
    }


def _materialize_transition_sound(
    project_path: Path,
    raw_sound: str | None,
    media_files: dict[str, bytes],
    rel_entries: list[dict[str, str]],
    audio_exts_used: set[str],
    packaged_by_source: dict[Path, tuple[str, str]],
) -> dict[str, str] | None:
    """Package one optional project-local WAV for a slide transition."""
    if raw_sound is None:
        return None
    return _materialize_slide_sound(
        project_path,
        raw_sound,
        media_files,
        rel_entries,
        audio_exts_used,
        packaged_by_source,
        label='transition',
        media_prefix='transition_sound',
        require_project_relative_wav=True,
    )


def _materialize_animation_sounds(
    project_path: Path,
    targets: list[dict[str, Any]],
    media_files: dict[str, bytes],
    rel_entries: list[dict[str, str]],
    audio_exts_used: set[str],
    packaged_by_source: dict[Path, tuple[str, str]] | None = None,
) -> list[dict[str, Any]]:
    """Package sidecar sound files and replace paths with OOXML relationships."""
    materialized: list[dict[str, Any]] = []
    packaged_by_source = packaged_by_source if packaged_by_source is not None else {}
    for index, raw_target in enumerate(targets, 1):
        target = dict(raw_target)
        raw_sound = target.pop('_sound_path', None)
        if raw_sound is None:
            materialized.append(target)
            continue
        target['sound'] = _materialize_slide_sound(
            project_path,
            raw_sound,
            media_files,
            rel_entries,
            audio_exts_used,
            packaged_by_source,
            label=f'animation target {index}',
            media_prefix='animation_sound',
            require_project_relative_wav=False,
        )
        materialized.append(target)
    return materialized


def _prerender_legacy_pngs(
    svg_files: list[Path],
    media_dir: Path,
    pixel_width: int,
    pixel_height: int,
    cache_dir: Path | None,
    workers: int,
    verbose: bool,
) -> dict[int, bool]:
    """Render every SVG→PNG into media_dir in parallel.

    Returns {1-based slide index: success}. Falls back to sequential when
    workers<=1 or len(svg_files)<=2.
    """
    results: dict[int, bool] = {}
    targets: list[tuple[int, Path, Path]] = [
        (i, svg, media_dir / f'image{i}.png')
        for i, svg in enumerate(svg_files, 1)
    ]

    if workers <= 1 or len(targets) <= 2:
        for i, svg, png in targets:
            ok = convert_svg_to_png_cached(svg, png, pixel_width, pixel_height, cache_dir)
            results[i] = ok
            if verbose:
                tag = 'cached/ok' if ok else 'failed'
                print(f"  [PNG {i}/{len(targets)}] {svg.name} - {tag}")
        return results

    with ProcessPoolExecutor(max_workers=workers) as pool:
        future_map = {
            pool.submit(
                convert_svg_to_png_cached,
                svg, png, pixel_width, pixel_height, cache_dir,
            ): (i, svg)
            for i, svg, png in targets
        }
        done = 0
        for future in as_completed(future_map):
            i, svg = future_map[future]
            try:
                ok = future.result()
            except Exception as exc:
                ok = False
                if verbose:
                    print(f"  [PNG] {svg.name} - worker error: {exc}")
            results[i] = ok
            done += 1
            if verbose:
                tag = 'cached/ok' if ok else 'failed'
                print(f"  [PNG {done}/{len(targets)}] {svg.name} - {tag}")

    return results


def _presentation_format(width: float, height: float) -> str:
    """Map the slide aspect ratio to PowerPoint's PresentationFormat label.
    Non-standard ratios (square, portrait, banner crops) report 'Custom'.
    """
    if width <= 0 or height <= 0:
        return 'Custom'
    ratio = width / height
    for target, label in (
        (4 / 3, 'On-screen Show (4:3)'),
        (16 / 9, 'On-screen Show (16:9)'),
        (16 / 10, 'On-screen Show (16:10)'),
    ):
        if abs(ratio - target) < 0.02:
            return label
    return 'Custom'


def _stamp_docprops(
    extract_dir: Path,
    slide_count: int,
    pres_format: str,
    meta: dict[str, Any] | None = None,
) -> None:
    """Overwrite the misleading python-pptx default metadata with accurate
    values. Factual fields (slide count, export timestamp, presentation format,
    application) are always machine-derived. Authored fields — including the
    title — come solely from an optional per-project ``metadata.json``
    (``meta``); whatever it omits stays blank. ``lastModifiedBy`` follows
    ``creator`` rather than ever carrying the base template's author or a tool
    name. No field is guessed from slide content: a blank title is preferable
    to an unreliable heuristic pick.
    """
    meta = meta or {}

    def field(key: str, default: str = '') -> str:
        value = meta.get(key)
        return value.strip() if isinstance(value, str) and value.strip() else default

    title = field('title')
    creator = field('creator')

    now = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')

    core_path = extract_dir / 'docProps' / 'core.xml'
    if core_path.exists():
        core_path.write_text(
            "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>\n"
            '<cp:coreProperties '
            'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" '
            'xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:dcmitype="http://purl.org/dc/dcmitype/" '
            'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f'<dc:title>{escape(title)}</dc:title>'
            f'<dc:subject>{escape(field("subject"))}</dc:subject>'
            f'<dc:creator>{escape(creator)}</dc:creator>'
            f'<cp:keywords>{escape(field("keywords"))}</cp:keywords>'
            f'<dc:description>{escape(field("description"))}</dc:description>'
            f'<dc:language>{escape(field("language"))}</dc:language>'
            f'<cp:lastModifiedBy>{escape(creator)}</cp:lastModifiedBy>'
            '<cp:revision>1</cp:revision>'
            f'<dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created>'
            f'<dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified>'
            f'<cp:category>{escape(field("category"))}</cp:category>'
            f'<cp:contentStatus>{escape(field("contentStatus"))}</cp:contentStatus>'
            '</cp:coreProperties>',
            encoding='utf-8',
        )

    app_path = extract_dir / 'docProps' / 'app.xml'
    if app_path.exists():
        app = app_path.read_text(encoding='utf-8')
        app = re.sub(r'<Slides>.*?</Slides>', f'<Slides>{slide_count}</Slides>', app)
        app = re.sub(
            r'<Company>.*?</Company>',
            f'<Company>{escape(field("company"))}</Company>',
            app,
        )
        app = re.sub(
            r'<Manager>.*?</Manager>',
            f'<Manager>{escape(field("manager"))}</Manager>',
            app,
        )
        app = re.sub(
            r'<Application>.*?</Application>',
            '<Application>Microsoft Office PowerPoint</Application>',
            app,
        )
        app = re.sub(
            r'<PresentationFormat>.*?</PresentationFormat>',
            f'<PresentationFormat>{escape(pres_format)}</PresentationFormat>',
            app,
        )
        app_path.write_text(app, encoding='utf-8')


def _create_preserved_base_pptx(
    contract: NativeStructureContract,
    specs: list[TemplateSlideSpec],
    output_path: Path,
    slide_size_emu: tuple[int, int],
    *,
    roundtrip_page_sources: tuple[int, ...] | None = None,
    package_overrides: dict[str, bytes] | None = None,
) -> bool:
    """Create the preserve base and report whether source Slides were retained."""
    if roundtrip_page_sources is not None:
        if len(roundtrip_page_sources) != len(specs):
            raise TemplateStructureError(
                "Round-trip page plan length does not match the SVG roster"
            )
        if contract.slide_size_emu != slide_size_emu:
            raise TemplateStructureError(
                "Generated SVG canvas does not match the preserved source template size"
            )
        for spec, source_index in zip(specs, roundtrip_page_sources):
            source_slide = contract.slide(source_index)
            if spec.layout_key != source_slide.layout_key:
                raise TemplateStructureError(
                    f"{spec.svg_path.name} declares layout {spec.layout_key!r}, "
                    f"but source slide {source_index} uses "
                    f"{source_slide.layout_key!r}"
                )
        try:
            clone_presentation_slides(
                contract.source_template,
                roundtrip_page_sources,
                output_path,
                package_overrides=package_overrides,
            )
        except (OSError, RuntimeError, zipfile.BadZipFile) as exc:
            raise TemplateStructureError(
                f"Round-trip page-plan cloning failed: {exc}"
            ) from exc
        return True

    source_roster_matches = (
        len(specs) == len(contract.slides)
        and all(
            spec.slide_num == source_slide.index
            and spec.layout_key == source_slide.layout_key
            for spec, source_slide in zip(specs, contract.slides)
        )
    )
    if source_roster_matches:
        shutil.copy2(contract.source_template, output_path)
        return True

    presentation = Presentation(str(contract.source_template))
    actual_size = (int(presentation.slide_width), int(presentation.slide_height))
    if actual_size != contract.slide_size_emu:
        raise TemplateStructureError(
            f"{contract.source_template.name} slide size does not match "
            f"{contract.contract_path.name}"
        )
    if actual_size != slide_size_emu:
        raise TemplateStructureError(
            "Generated SVG canvas does not match the preserved source template size"
        )

    layouts_by_part = {
        str(layout.part.partname).lstrip("/"): layout
        for master in presentation.slide_masters
        for layout in master.slide_layouts
    }
    slide_ids = presentation.slides._sldIdLst
    for slide_id in list(slide_ids):
        presentation.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)

    for spec in specs:
        layout_contract = contract.layout(spec.layout_key)
        layout = layouts_by_part.get(layout_contract.package_part)
        if layout is None:
            raise TemplateStructureError(
                f"Preserved source package did not load layout part "
                f"{layout_contract.package_part!r}"
            )
        presentation.slides.add_slide(layout)
    presentation.save(str(output_path))
    return False


def _clear_preserved_slide_collections(extract_dir: Path) -> None:
    """Remove source slide-order metadata that cannot apply to generated pages."""
    presentation_path = extract_dir / "ppt" / "presentation.xml"
    tree = ET.parse(presentation_path)
    root = tree.getroot()
    custom_shows = root.find(f"{{{PML_NS}}}custShowLst")
    if custom_shows is not None:
        root.remove(custom_shows)
    for extension_list in root.findall(f".//{{{PML_NS}}}extLst"):
        for extension in list(extension_list):
            if any(
                child.tag.rsplit("}", 1)[-1] == "sectionLst"
                for child in extension.iter()
                if isinstance(child.tag, str)
            ):
                extension_list.remove(extension)
    _write_xml_tree(presentation_path, tree)


def _validate_source_theme_xml(payload: bytes) -> None:
    """Validate one imported Theme part before package installation."""
    try:
        root = ET.fromstring(payload)
    except ET.ParseError as exc:
        raise ThemeColorError(f"Imported source theme is malformed: {exc}") from exc
    if root.tag != f"{{{DML_NS}}}theme":
        raise ThemeColorError("Imported source theme root must be a:theme")
    if any(
        isinstance(name, str) and name.startswith(f"{{{REL_NS}}}")
        for node in root.iter()
        for name in node.attrib
    ):
        raise ThemeColorError("Imported source theme cannot contain relationships")


def _install_source_theme_xml(extract_dir: Path, payload: bytes) -> None:
    """Install a validated imported theme into the generated flat package."""
    _validate_source_theme_xml(payload)

    theme_paths = sorted((extract_dir / "ppt" / "theme").glob("theme*.xml"))
    if not theme_paths:
        raise ThemeColorError("Generated PPTX package has no theme part")
    for theme_path in theme_paths:
        theme_path.write_bytes(payload)


def _install_source_themes_by_master(
    extract_dir: Path,
    master_parts_by_key: dict[str, str],
    payloads_by_key: dict[str, bytes],
) -> None:
    """Install one exact source Theme into each structured mirror Master."""
    expected = set(master_parts_by_key)
    actual = set(payloads_by_key)
    if actual != expected:
        raise ThemeColorError(
            "Structured source Theme roster differs from Master roster; "
            f"missing={sorted(expected - actual)}, extra={sorted(actual - expected)}"
        )
    for master_key, master_part in master_parts_by_key.items():
        payload = payloads_by_key[master_key]
        _validate_source_theme_xml(payload)
        master_rels = _relationships_path_for_part(extract_dir, master_part)
        targets = [
            attrs["Target"]
            for attrs in _read_relationships(master_rels).values()
            if attrs.get("Type") == THEME_REL_TYPE and attrs.get("Target")
        ]
        if len(targets) != 1:
            raise ThemeColorError(
                f"Structured Master {master_key!r} must own one Theme relationship"
            )
        theme_part = _resolve_package_target(master_part, targets[0])
        theme_path = extract_dir / theme_part
        if not theme_path.is_file():
            raise ThemeColorError(
                f"Structured Master {master_key!r} Theme part is missing: {theme_part}"
            )
        theme_path.write_bytes(payload)


def _install_source_embedded_fonts(
    extract_dir: Path,
    bundle: EmbeddedFontBundle,
) -> tuple[str, ...]:
    """Install validated source font parts into the generated PPTX package."""
    typefaces = embedded_font_typefaces(bundle)
    try:
        font_list = ET.fromstring(bundle.font_list_xml)
    except ET.ParseError as exc:
        raise EmbeddedFontError(
            f"Embedded font list XML is malformed: {exc}"
        ) from exc

    ppt_dir = extract_dir / "ppt"
    presentation_path = ppt_dir / "presentation.xml"
    presentation_rels_path = ppt_dir / "_rels" / "presentation.xml.rels"
    presentation_tree = ET.parse(presentation_path)
    presentation_root = presentation_tree.getroot()
    rels_tree = ET.parse(presentation_rels_path)
    rels_root = rels_tree.getroot()

    fonts_dir = ppt_dir / "fonts"
    fonts_dir.mkdir(exist_ok=True)
    for relationship in list(rels_root):
        if relationship.attrib.get("Type") != FONT_REL_TYPE:
            continue
        target = relationship.attrib.get("Target", "")
        target_path = PurePosixPath(target)
        if (
            not target_path.is_absolute()
            and target_path.parts[:1] == ("fonts",)
            and target_path.suffix.lower() == ".fntdata"
        ):
            candidate = ppt_dir.joinpath(*target_path.parts)
            if candidate.is_file():
                candidate.unlink()
        rels_root.remove(relationship)

    rid_numbers = [
        int(match.group(1))
        for relationship in rels_root
        if (match := re.fullmatch(
            r"rId(\d+)",
            relationship.attrib.get("Id", ""),
        ))
    ]
    next_rid = max(rid_numbers, default=0) + 1
    relationship_mapping: dict[str, str] = {}
    for index, part in enumerate(bundle.parts, start=1):
        filename = f"font{index}.fntdata"
        (fonts_dir / filename).write_bytes(part.payload)
        relationship_id = f"rId{next_rid}"
        next_rid += 1
        ET.SubElement(
            rels_root,
            f"{{{PACKAGE_REL_NS}}}Relationship",
            {
                "Id": relationship_id,
                "Type": FONT_REL_TYPE,
                "Target": f"fonts/{filename}",
            },
        )
        relationship_mapping[part.relationship_id] = relationship_id

    for node in font_list.iter():
        relationship_attr = f"{{{EMBEDDED_FONT_REL_NS}}}id"
        original_id = node.attrib.get(relationship_attr)
        if original_id is None:
            continue
        replacement_id = relationship_mapping.get(original_id)
        if replacement_id is None:
            raise EmbeddedFontError(
                f"Embedded font list references an unknown part: {original_id}"
            )
        node.set(relationship_attr, replacement_id)

    existing = presentation_root.find(
        f"{{{EMBEDDED_FONT_PML_NS}}}embeddedFontLst"
    )
    if existing is not None:
        insertion_index = list(presentation_root).index(existing)
        presentation_root.remove(existing)
    else:
        default_text_style = presentation_root.find(
            f"{{{EMBEDDED_FONT_PML_NS}}}defaultTextStyle"
        )
        insertion_index = (
            list(presentation_root).index(default_text_style)
            if default_text_style is not None
            else len(presentation_root)
        )
    presentation_root.insert(insertion_index, font_list)
    _write_xml_tree(presentation_path, presentation_tree)
    _write_xml_tree(presentation_rels_path, rels_tree)

    content_types_path = extract_dir / "[Content_Types].xml"
    content_types_path.write_text(
        _add_default_content_type(
            content_types_path.read_text(encoding="utf-8"),
            "fntdata",
            FONT_CONTENT_TYPE,
        ),
        encoding="utf-8",
    )
    return typefaces


def create_pptx_with_native_svg(
    svg_files: list[Path],
    output_path: Path,
    *,
    resource_root: Path,
    canvas_format: str | None = None,
    verbose: bool = True,
    transition: str | None = 'fade',
    transition_duration: float = 0.5,
    auto_advance: float | None = None,
    use_compat_mode: bool = True,
    notes: dict[str, str] | None = None,
    enable_notes: bool = True,
    use_native_shapes: bool = True,
    animation: str | None = None,
    animation_duration: float = 0.4,
    animation_stagger: float = 0.5,
    animation_trigger: str = 'after-previous',
    animation_config: dict[str, Any] | None = None,
    animation_cli_overrides: dict[str, bool] | None = None,
    narration_audio: dict[str, Path] | None = None,
    use_narration_timings: bool = False,
    narration_padding: float = 0.5,
    cache_dir: Path | None = None,
    workers: int | None = None,
    merge_paragraphs: bool | None = None,
    image_optimize: bool = True,
    image_max_dimension: int | None = 2560,
    image_sizing: str = 'cap',
    image_scale: float = 2.0,
    image_quality: int = 85,
    native_objects: bool = False,
    conversion_trace_path: Path | None = None,
    dangerous_nonconforming_export: bool = False,
    doc_metadata: dict[str, Any] | None = None,
    structure_name: str | None = None,
    pptx_structure: str = "structured",
    use_layout_placeholder_frames: bool = False,
    native_structure_contract: NativeStructureContract | None = None,
    roundtrip_passthrough_slides: set[int] | None = None,
    roundtrip_slide_patches: dict[int, RoundtripSlidePatch] | None = None,
    roundtrip_resources: tuple[WorkspaceResourceSpec, ...] = (),
    roundtrip_page_sources: tuple[int, ...] | None = None,
    theme_font_spec: ThemeFontSpec | None = None,
    master_text_style_spec: MasterTextStyleSpec | None = None,
    theme_color_spec: ThemeColorSpec | None = None,
    source_theme_xml: bytes | None = None,
    source_theme_xml_by_master: dict[str, bytes] | None = None,
    source_embedded_fonts: EmbeddedFontBundle | None = None,
    structured_baseline: bool = False,
    baseline_layout_specs: list[TemplateSlideSpec] | None = None,
    layout_definition_files: list[Path] | None = None,
    expected_viewbox: str | None = None,
    animation_resource_root: Path | None = None,
    transition_effect_options: dict[str, object] | None = None,
    transition_sound: str | None = None,
    text_flow: str | None = None,
    primary_language: str | None = None,
    narration_start_floor: float = DEFAULT_NARRATION_START_FLOOR,
) -> bool:
    """Create a PPTX file with native DrawingML shapes.

    Args:
        svg_files: List of SVG files.
        output_path: Output PPTX path.
        layout_definition_files: Optional complete Slide SVG prototypes for
            Layouts that no generated page uses. They are converted on internal
            carrier slides, registered, and removed before publication.
        canvas_format: Canvas format key.
        expected_viewbox: Optional project/template-lock canvas contract. Every
            public page and internal Layout carrier prototype must match it.
        animation_resource_root: Project root for sidecar sound paths. Object
            animation sounds retain existing absolute-path compatibility;
            transition sounds must remain project-relative WAV files.
        verbose: Whether to output detailed information.
        transition: Transition effect name.
        transition_effect_options: PowerPoint Effect Options for the selected
            native page transition.
        transition_sound: Optional project-relative WAV path used by the
            generated page transition.
        transition_duration: Transition duration in seconds.
        auto_advance: Auto-advance interval in seconds.
        use_compat_mode: Retained for API compatibility; ignored in native mode.
        notes: Notes dict, key is SVG stem, value is notes content.
        enable_notes: Whether to enable notes embedding.
        use_native_shapes: Must remain true; SVG-image PPTX export is unsupported.
        animation: Per-element object-animation mode (compatibility alias,
            PowerPoint-native ``entrance_*``/``emphasis_*``/``path_*``/
            ``exit_*`` effect, ``'mixed'``, ``'random'``, or None to disable).
            Native shapes mode only.
        animation_duration: Per-element animation duration in seconds.
            Instantaneous native presets retain their PowerPoint-authored
            duration.
        animation_stagger: Delay between elements in ``after-previous``
            trigger mode (seconds). Ignored otherwise.
        animation_trigger: PowerPoint Start mode — ``'after-previous'`` (default),
            ``'on-click'``, or ``'with-previous'``.
        animation_config: Optional sidecar overrides loaded from animations.json.
        animation_cli_overrides: Flags indicating explicit CLI overrides.
        narration_audio: Optional dict mapping SVG stem to narration audio file.
        use_narration_timings: Whether to set slide auto-advance from audio duration.
        narration_padding: Extra seconds added after each narration before advancing.
        narration_start_floor: Minimum seconds from transition start to narration
            start. Any remainder after the transition becomes silent lead-in.
        merge_paragraphs: Legacy compatibility option. True selects reflow;
            False selects split. Do not combine with ``text_flow``.
        text_flow: Positional-tspan policy: preserve authored line breaks in
            one frame, reflow text, or split visual lines into separate frames.
        image_optimize: Whether native export optimizes raster images when needed.
        image_max_dimension: Preferred optimized image dimension cap in pixels.
        image_sizing: ``cap`` preserves unchanged source bytes and limits
            oversized sources; ``display`` sizes from rendered SVG boxes.
        image_scale: Target image pixels per SVG display pixel.
        image_quality: JPEG quality used when opaque rasters are re-encoded.
        native_objects: Replace opt-in ``data-pptx-replace-with`` chart/table
            fallback groups with native PowerPoint Chart/Table objects. Formula
            markers remain intrinsically native; Chart/Table markers stay off otherwise.
        conversion_trace_path: Optional JSON path for native conversion diagnostics.
        dangerous_nonconforming_export: Apply narrowly defined compatibility
            normalizations before strict SVG conversion. Actual contract,
            conversion, and package failures remain blocking.
        resource_root: Explicit project boundary for SVG-local resource paths.
        structure_name: Current deck identity used to name a flat Master, Layout,
            and theme.
        pptx_structure: PPTX structure strategy. ``baseline`` promotes safe
            shared native backgrounds and leading chrome to slide masters,
            then extracts semantic page-role layout families and exact
            family-wide structurally marked leading chrome; marker-free legacy
            SVGs retain filename/id fallback;
            ``structured`` consumes explicit SVG master/layout/placeholder
            metadata; ``preserve`` reuses an imported source PPTX package;
            ``flat`` keeps generated content Slide-local and builds one clean
            project-owned Master/Blank-Layout shell.
        use_layout_placeholder_frames: In structured template-review decks, size
            each Slide placeholder carrier to its reusable Layout bounds instead
            of the tight SVG content frame. Default off for generated decks.
        native_structure_contract: Validated source package contract for
            ``preserve`` mode.
        roundtrip_passthrough_slides: Source slide indices whose SVG, notes,
            and motion sidecars are unchanged and may retain their original
            slide XML and relationships byte-for-byte.
        roundtrip_slide_patches: Authoring metadata for rebuilt slides that
            overlay edited objects onto their preserved source slide parts.
        roundtrip_resources: Semantic workspace payloads copied back to their
            exact source package parts after slide conversion.
        roundtrip_page_sources: Optional output-order tuple of one-based source
            slide indices. Presence activates deck-level source slide cloning.
        theme_font_spec: Locked project major/minor fonts for flat/structured
            release-theme inheritance. Direct diagnostic flat callers may omit it.
        master_text_style_spec: Required declared or inferred title/body anchors
            for structured and release flat slide-master text styles. Direct
            diagnostic flat callers may omit it; other routes ignore this value.
        theme_color_spec: Locked project color scheme for context-aware
            flat/structured theme inheritance. Preserve mode ignores this value.
        source_theme_xml: Complete validated source theme used only by an
            explicit PPTX-import diagnostic round-trip.
        source_theme_xml_by_master: Exact validated source themes keyed by
            structured mirror Master identity.
        source_embedded_fonts: Validated source font-list metadata and font
            parts used only by an explicit PPTX-import diagnostic round-trip.
        primary_language: Canonical BCP-47 deck content language. ``None``
            preserves legacy per-run language detection.
        structured_baseline: Obsolete compatibility argument; must remain false.
        baseline_layout_specs: Obsolete compatibility argument; must remain None.

    Returns:
        Whether all slides were successfully created.
    """
    if source_theme_xml is not None and source_theme_xml_by_master is not None:
        raise ThemeColorError(
            "Use either one diagnostic source Theme or per-Master mirror Themes"
        )
    if source_theme_xml_by_master is not None and pptx_structure != "structured":
        raise ThemeColorError(
            "Per-Master source Themes are allowed only in structured export"
        )
    text_flow = resolve_text_flow(text_flow, merge_paragraphs)
    if primary_language is not None:
        primary_language = normalize_language_tag(primary_language)
    public_svg_files = list(svg_files)
    passthrough_slides = set(roundtrip_passthrough_slides or set())
    slide_patches = dict(roundtrip_slide_patches or {})
    overlay_slides = set(slide_patches)
    page_plan_export = roundtrip_page_sources is not None
    roundtrip_export = bool(
        passthrough_slides
        or overlay_slides
        or roundtrip_resources
        or page_plan_export
    )
    source_motion_slides = passthrough_slides | {
        index
        for index, patch in slide_patches.items()
        if not patch.animation_changed
    }
    invalid_passthrough = sorted(
        index
        for index in passthrough_slides
        if index < 1 or index > len(public_svg_files)
    )
    if invalid_passthrough:
        raise ValueError(
            "Round-trip passthrough slide indices are outside the SVG roster: "
            + ", ".join(str(index) for index in invalid_passthrough)
        )
    invalid_overlays = sorted(
        index
        for index in overlay_slides
        if index < 1 or index > len(public_svg_files)
    )
    if invalid_overlays:
        raise ValueError(
            "Round-trip overlay slide indices are outside the SVG roster: "
            + ", ".join(str(index) for index in invalid_overlays)
        )
    overlap = sorted(passthrough_slides & overlay_slides)
    if overlap:
        raise ValueError(
            "Round-trip slides cannot be both passthrough and overlay: "
            + ", ".join(str(index) for index in overlap)
        )
    if passthrough_slides and pptx_structure != "preserve":
        raise ValueError(
            "Round-trip slide passthrough is available only in preserve mode"
        )
    if (overlay_slides or roundtrip_resources) and pptx_structure != "preserve":
        raise ValueError(
            "Round-trip overlays and resource reinjection require preserve mode"
        )
    if page_plan_export:
        if pptx_structure != "preserve":
            raise ValueError(
                "Round-trip page plans are available only in preserve mode"
            )
        if len(roundtrip_page_sources) != len(public_svg_files):
            raise ValueError(
                "Round-trip page-plan source roster differs from the SVG roster"
            )
    definition_svg_files = list(layout_definition_files or [])
    public_slide_names = [path.stem for path in public_svg_files]
    morph_pairs = resolve_morph_pairs(
        public_slide_names,
        animation_config,
    )
    public_slide_numbers = {
        slide_name: slide_number
        for slide_number, slide_name in enumerate(public_slide_names, 1)
    }
    morph_expectations = tuple(
        MorphPairExpectation(
            source_slide_number=public_slide_numbers[pair.source_slide],
            destination_slide_number=public_slide_numbers[
                pair.destination_slide
            ],
            key=pair.key,
        )
        for pair in morph_pairs
    )
    morph_pairs_by_destination: dict[str, list[MorphPair]] = {}
    morph_group_overrides_by_slide: dict[str, set[str]] = {}
    for pair in morph_pairs:
        morph_pairs_by_destination.setdefault(
            pair.destination_slide,
            [],
        ).append(pair)
        morph_group_overrides_by_slide.setdefault(
            pair.source_slide,
            set(),
        ).add(pair.source_group_id)
        morph_group_overrides_by_slide.setdefault(
            pair.destination_slide,
            set(),
        ).add(pair.destination_group_id)
    morph_shape_ids: dict[tuple[str, str], int] = {}
    if definition_svg_files and pptx_structure != "structured":
        raise ValueError(
            "unselected Layout prototypes require pptx_structure='structured'"
        )
    public_paths = {path.resolve() for path in public_svg_files}
    seen_definition_paths: set[Path] = set()
    for path in definition_svg_files:
        resolved = path.resolve()
        if not path.is_file():
            raise ValueError(f"Layout prototype SVG does not exist: {path}")
        if resolved in public_paths:
            raise ValueError(
                f"Layout prototype SVG is already a generated page: {path}"
            )
        if resolved in seen_definition_paths:
            raise ValueError(f"Layout prototype SVG is repeated: {path}")
        seen_definition_paths.add(resolved)
    public_slide_count = len(public_svg_files)
    svg_files = public_svg_files + definition_svg_files
    total_slide_count = len(svg_files)

    if not use_native_shapes:
        raise ValueError(
            "SVG-image PPTX export is no longer supported; use svg_final/ "
            "directly for preview and native DrawingML PPTX for delivery"
        )
    if not public_svg_files:
        print("Error: No SVG files found")
        return False

    use_compat_mode = False
    if pptx_structure not in {"baseline", "structured", "preserve", "flat"}:
        raise ValueError(f"Unsupported pptx_structure: {pptx_structure}")
    requested_canvas_format = canvas_format
    canvas, detected_canvas_format = resolve_svg_canvas(
        svg_files,
        canvas_format=canvas_format,
        expected_viewbox=expected_viewbox,
    )
    if canvas_format is None:
        canvas_format = detected_canvas_format
    if pptx_structure == "flat":
        flat_errors = flat_structure_metadata_errors(public_svg_files)
        if flat_errors:
            details = "\n".join(f"  - {error}" for error in flat_errors)
            raise TemplateStructureError(
                "Flat PPTX structure validation failed:\n" + details
            )
    if use_layout_placeholder_frames and pptx_structure != "structured":
        raise ValueError(
            "use_layout_placeholder_frames requires pptx_structure='structured'"
        )
    if structured_baseline:
        raise ValueError(
            "structured_baseline is obsolete; use pptx_structure='structured'"
        )
    if baseline_layout_specs is not None:
        raise ValueError(
            "baseline_layout_specs is obsolete; structured export parses SVG metadata"
        )
    if pptx_structure == "structured" and master_text_style_spec is None:
        raise ValueError(
            "Structured export requires declared or inferred typography "
            "title/body anchors in master_text_style_spec"
        )
    if use_native_shapes and pptx_structure == "structured":
        template_specs = parse_template_slides(svg_files)
        public_template_specs = template_specs[:public_slide_count]
    elif use_native_shapes and pptx_structure == "preserve":
        if native_structure_contract is None:
            raise TemplateStructureError(
                "Preserve export requires a validated native structure contract"
            )
        template_specs = parse_preserve_slides(svg_files)
        for spec in template_specs:
            native_structure_contract.layout(spec.layout_key)
        public_template_specs = template_specs
    else:
        template_specs = None
        public_template_specs = None
    template_background_expectations: dict[str, str | None] | None = None
    template_shape_roster_expectations: (
        dict[str, tuple[str, ...]] | None
    ) = None
    template_layout_parts_by_key: dict[str, str] | None = None
    template_master_parts_by_key: dict[str, str] | None = None
    if template_specs is not None and not native_objects:
        native_placeholders = sorted({
            item.placeholder
            for spec in template_specs
            for item in spec.placeholders
            if item.placeholder in {"chart", "table"}
        })
        if native_placeholders:
            kinds = ", ".join(str(kind) for kind in native_placeholders)
            context = (
                pptx_structure.capitalize()
            )
            raise TemplateStructureError(
                f"{context} {kinds} placeholder(s) require "
                "--native-charts-and-tables so each marker becomes one native "
                "PowerPoint Chart/Table object"
            )

    # Check compatibility mode dependencies
    renderer_name, renderer_status, renderer_hint = get_png_renderer_info()
    if not use_native_shapes and use_compat_mode and PNG_RENDERER is None:
        print("Warning: No PNG rendering library installed, cannot use compatibility mode")
        print(f"  {renderer_hint}")
        print("  Will use pure SVG mode (may not display in Office LTSC 2021 and similar versions)")
        use_compat_mode = False

    width_emu, height_emu = canvas.emu_dimensions
    pixel_width, pixel_height = canvas.pixel_dimensions
    pixel_width_label, pixel_height_label = canvas.canonical.split()[2:]
    if verbose and requested_canvas_format is None:
        if canvas_format:
            format_name = CANVAS_FORMATS.get(canvas_format, {}).get('name', canvas_format)
            print(f"  Detected canvas format: {format_name}")
        else:
            print(
                "  Using SVG viewBox dimensions: "
                f"{canvas.canonical.removeprefix('0 0 ')} px"
            )

    if verbose:
        print(
            f"  Slide dimensions: {pixel_width_label} x "
            f"{pixel_height_label} px"
        )
        print(f"  SVG file count: {public_slide_count}")
        if definition_svg_files:
            print(
                "  Unselected Layout prototype carriers: "
                f"{len(definition_svg_files)}"
            )
        if use_native_shapes:
            print(f"  Mode: Native DrawingML shapes (directly editable)")
            native_object_mode = (
                "Chart/Table replacement enabled"
                if native_objects
                else "Chart/Table replacement disabled"
            )
            print(
                "  Native table/chart objects: "
                f"{native_object_mode}"
            )
            print(f"  PPTX structure: {pptx_structure}")
            if image_optimize:
                if image_sizing == 'display':
                    image_mode = (
                        f"display scale {image_scale:g}, "
                        f"preferred max {image_max_dimension or 'unlimited'} px"
                    )
                else:
                    image_mode = (
                        f"preferred cap {image_max_dimension or 'unlimited'} px, "
                        "unchanged bytes preserved"
                    )
                print(
                    "  Image optimization: Enabled "
                    f"({image_mode}, JPEG q{image_quality} when re-encoded)"
                )
            else:
                print("  Image optimization: Disabled (original bytes)")
        elif use_compat_mode:
            print(f"  Compatibility mode: Enabled (PNG + SVG dual format)")
            print(f"  PNG renderer: {renderer_name} {renderer_status}")
        else:
            print(f"  Compatibility mode: Disabled (pure SVG)")
        if transition:
            canonical_transition, _transition_options = (
                normalize_transition_effect_request(
                    transition,
                    transition_effect_options,
                )
            )
            trans_name = (
                NATIVE_TRANSITIONS.get(canonical_transition, {}).get(
                    'name',
                    canonical_transition,
                )
                if canonical_transition
                else transition
            )
            print(f"  Transition effect: {trans_name}")
        if enable_notes and notes:
            print(f"  Speaker notes: {len(notes)} page(s)")
        elif enable_notes:
            print(f"  Speaker notes: Enabled (no notes files found)")
        else:
            print(f"  Speaker notes: Disabled")
        print()

    animation_cli_overrides = animation_cli_overrides or {}

    temp_dir = _create_writable_work_dir(output_path)

    try:
        base_pptx = temp_dir / 'base.pptx'
        preserved_source_slides = False
        page_plan_package_overrides = (
            _roundtrip_resource_payloads(resource_root, roundtrip_resources)
            if page_plan_export
            else None
        )
        late_roundtrip_resources = (
            () if page_plan_export else roundtrip_resources
        )
        if (
            use_native_shapes
            and pptx_structure == "preserve"
            and native_structure_contract is not None
            and template_specs is not None
        ):
            preserved_source_slides = _create_preserved_base_pptx(
                native_structure_contract,
                template_specs,
                base_pptx,
                (width_emu, height_emu),
                roundtrip_page_sources=roundtrip_page_sources,
                package_overrides=page_plan_package_overrides,
            )
        else:
            # Create the standard base PPTX with python-pptx.
            prs = Presentation()
            prs.slide_width = width_emu
            prs.slide_height = height_emu

            blank_layout = prs.slide_layouts[6]
            for _ in svg_files:
                prs.slides.add_slide(blank_layout)
            prs.save(str(base_pptx))

        # Extract PPTX
        extract_dir = temp_dir / 'pptx_content'
        with zipfile.ZipFile(base_pptx, 'r') as zf:
            zf.extractall(extract_dir)
        source_content_types_bytes = (
            (extract_dir / '[Content_Types].xml').read_bytes()
            if roundtrip_export
            else None
        )
        roundtrip_source_parts = (
            frozenset(
                path.relative_to(extract_dir).as_posix()
                for path in extract_dir.rglob('*')
                if path.is_file()
            )
            if roundtrip_export
            else frozenset()
        )
        if (
            use_native_shapes
            and pptx_structure == "preserve"
            and not preserved_source_slides
        ):
            _clear_preserved_slide_collections(extract_dir)
        if (
            passthrough_slides
            or overlay_slides
            or roundtrip_resources
            or page_plan_export
        ) and not preserved_source_slides:
            raise TemplateStructureError(
                "Round-trip export requires the original source package roster"
            )
        active_theme_font_spec = (
            theme_font_spec
            if use_native_shapes
            and pptx_structure in {"baseline", "flat", "structured"}
            else None
        )
        if active_theme_font_spec is not None:
            apply_theme_font_spec(extract_dir, active_theme_font_spec)
        active_theme_color_spec = (
            theme_color_spec
            if use_native_shapes
            and pptx_structure in {"baseline", "flat", "structured"}
            else None
        )
        if active_theme_color_spec is not None:
            apply_theme_color_spec(extract_dir, active_theme_color_spec)
        if (
            source_theme_xml is not None
            and native_structure_contract is None
        ):
            _install_source_theme_xml(extract_dir, source_theme_xml)
        if source_embedded_fonts is not None and not roundtrip_export:
            installed_typefaces = _install_source_embedded_fonts(
                extract_dir,
                source_embedded_fonts,
            )
            if verbose:
                print(
                    "  Embedded fonts: preserved "
                    f"{len(source_embedded_fonts.parts)} part(s) for "
                    + ", ".join(installed_typefaces)
                )
        structure = _read_slide_layout_targets(extract_dir, len(svg_files))

        media_dir = extract_dir / 'ppt' / 'media'
        media_dir.mkdir(exist_ok=True)

        prerender_results: dict[int, bool] | None = None
        if not use_native_shapes and use_compat_mode and PNG_RENDERER is not None:
            if workers is None:
                resolved_workers = min(os.cpu_count() or 2, len(svg_files), 8)
            else:
                resolved_workers = max(0, workers)
            if verbose:
                cache_label = str(cache_dir) if cache_dir else 'disabled'
                mode = f'parallel x{resolved_workers}' if resolved_workers > 1 else 'sequential'
                print(f"  Pre-rendering PNGs ({mode}, cache: {cache_label})")
            prerender_results = _prerender_legacy_pngs(
                svg_files, media_dir, pixel_width, pixel_height,
                cache_dir, resolved_workers, verbose,
            )
            if verbose:
                print()

        success_count = 0
        has_any_image = False
        media_cache: dict[tuple[str, str], str] = {}
        image_exts_used: set[str] = set()
        package_exts_used: set[str] = set()
        package_content_overrides: dict[str, str] = {}
        notes_slides_created: set[int] = set()
        notes_master_parts_used: set[str] = set()
        notes_master_theme_parts_created: set[str] = set()
        narration_slides_created: set[int] = set()
        audio_exts_used: set[str] = set()
        package_uses_timings = False
        mixed_animation_offset = 0
        config_defaults = _as_dict(_as_dict(animation_config).get('defaults'))
        transition_defaults_value = config_defaults.get('transition', {})
        if not isinstance(transition_defaults_value, dict):
            raise ValueError(
                'animations.json defaults transition must be an object'
            )
        default_transition_cfg = transition_defaults_value
        animation_defaults_value = config_defaults.get('animation', {})
        if not isinstance(animation_defaults_value, dict):
            raise ValueError(
                'animations.json defaults animation must be an object'
            )
        default_animation_cfg = animation_defaults_value
        animation_seed = json.dumps(
            {
                'animation': animation,
                'config': animation_config,
                'slides': [path.name for path in svg_files],
            },
            ensure_ascii=False,
            sort_keys=True,
            separators=(',', ':'),
        )
        animation_rng = random.Random(animation_seed)
        conversion_trace: list[dict[str, Any]] | None = [] if conversion_trace_path else None
        structure_trace: list[dict[str, Any]] | None = (
            []
            if use_native_shapes and pptx_structure in {"baseline", "structured", "preserve"}
            else None
        )

        for i, svg_path in enumerate(svg_files, 1):
            slide_num = i
            is_layout_definition = slide_num > public_slide_count
            progress_label = (
                f"[Layout carrier {slide_num - public_slide_count}/"
                f"{len(definition_svg_files)}]"
                if is_layout_definition
                else f"[Slide {slide_num}/{public_slide_count}]"
            )
            expected_animation_targets: list[dict[str, Any]] = []
            expected_animation_duration = animation_duration
            expected_animation_trigger = normalize_animation_trigger(animation_trigger)
            expected_transition_sound: dict[str, str] | None = None

            try:
                slide_xml_path = (
                    extract_dir / 'ppt' / 'slides' / f'slide{slide_num}.xml'
                )
                rels_path = (
                    extract_dir / 'ppt' / 'slides' / '_rels'
                    / f'slide{slide_num}.xml.rels'
                )
                slide_patch = slide_patches.get(slide_num)
                source_slide_bytes: bytes | None = None
                source_rels_bytes: bytes | None = None
                if slide_patch is not None:
                    if is_layout_definition:
                        raise TemplateStructureError(
                            "Internal Layout definitions cannot use round-trip overlays"
                        )
                    try:
                        source_slide_bytes = slide_xml_path.read_bytes()
                        source_rels_bytes = rels_path.read_bytes()
                    except OSError as exc:
                        raise TemplateStructureError(
                            f"Cannot read source slide {slide_num} for overlay: {exc}"
                        ) from exc
                if slide_num in passthrough_slides:
                    if is_layout_definition:
                        raise TemplateStructureError(
                            "Internal Layout definitions cannot use source-slide passthrough"
                        )
                    if verbose:
                        print(
                            f"  {progress_label} {svg_path.name} "
                            "(Source slide passthrough)"
                        )
                    success_count += 1
                    continue
                if (
                    slide_patch is not None
                    and not slide_patch.visual_changed
                    and slide_patch.motion_changed
                ):
                    overlay_slide_cfg = _slide_config(
                        animation_config,
                        svg_path.stem,
                    )
                    overlay_groups = overlay_slide_cfg.get("groups", {})
                    if not isinstance(overlay_groups, dict):
                        raise ValueError(
                            "animations.json slide groups must be an object"
                        )
                    (
                        overlay_transition,
                        overlay_transition_options,
                        overlay_transition_duration,
                        overlay_auto_advance,
                        overlay_transition_sound,
                    ) = _slide_transition_settings(
                        default_transition_cfg,
                        overlay_slide_cfg,
                        transition,
                        transition_effect_options,
                        transition_duration,
                        auto_advance,
                        transition_sound,
                        animation_cli_overrides,
                    )
                    (
                        overlay_animation,
                        _overlay_animation_duration,
                        _overlay_animation_stagger,
                        _overlay_animation_trigger,
                        _overlay_animation_cfg,
                    ) = _slide_animation_settings(
                        overlay_slide_cfg,
                        default_animation_cfg,
                        animation,
                        animation_duration,
                        animation_stagger,
                        animation_trigger,
                        animation_cli_overrides,
                    )
                    animation_override_requested = any(
                        animation_cli_overrides.get(key, False)
                        for key in (
                            "animation",
                            "animation_duration",
                            "animation_stagger",
                            "animation_trigger",
                        )
                    )
                    direct_transition_overlay = (
                        overlay_animation is None
                        and not overlay_groups
                        and not animation_override_requested
                        and overlay_transition_sound is None
                        and not morph_group_overrides_by_slide.get(
                            svg_path.stem
                        )
                        and not (
                            narration_audio
                            and narration_audio.get(svg_path.stem) is not None
                        )
                        and not use_narration_timings
                    )
                    if direct_transition_overlay:
                        if _apply_roundtrip_transition_overlay(
                            slide_xml_path,
                            effect=overlay_transition,
                            effect_options=overlay_transition_options,
                            duration=overlay_transition_duration,
                            auto_advance=overlay_auto_advance,
                            replace_transition=slide_patch.transition_replaced,
                        ):
                            package_uses_timings = True
                        if slide_patch.notes_changed:
                            notes_content = (
                                notes.get(svg_path.stem, '') if notes else ''
                            )
                            notes_master = _apply_slide_notes(
                                extract_dir,
                                rels_path,
                                slide_num,
                                notes_content,
                                primary_language,
                                enable_notes=enable_notes,
                            )
                            if notes_master is not None:
                                notes_slides_created.add(slide_num)
                                notes_master_parts_used.add(
                                    notes_master.package_part
                                )
                                if notes_master.created_theme_part is not None:
                                    notes_master_theme_parts_created.add(
                                        notes_master.created_theme_part
                                    )
                        if verbose:
                            print(
                                f"  {progress_label} {svg_path.name} "
                                "(Source slide motion overlay)"
                            )
                        success_count += 1
                        continue
                if (
                    slide_patch is not None
                    and slide_patch.notes_changed
                    and not slide_patch.visual_changed
                    and not slide_patch.motion_changed
                ):
                    notes_content = (
                        notes.get(svg_path.stem, '') if notes else ''
                    )
                    notes_master = _apply_slide_notes(
                        extract_dir,
                        rels_path,
                        slide_num,
                        notes_content,
                        primary_language,
                        enable_notes=enable_notes,
                    )
                    if notes_master is not None:
                        notes_slides_created.add(slide_num)
                        notes_master_parts_used.add(notes_master.package_part)
                        if notes_master.created_theme_part is not None:
                            notes_master_theme_parts_created.add(
                                notes_master.created_theme_part
                            )
                    if verbose:
                        print(
                            f"  {progress_label} {svg_path.name} "
                            "(Source slide notes overlay)"
                        )
                    success_count += 1
                    continue

                # ---- Native shapes mode ----
                if use_native_shapes:
                    slide_cfg = (
                        {}
                        if is_layout_definition
                        else _slide_config(animation_config, svg_path.stem)
                    )
                    if is_layout_definition:
                        slide_transition = None
                        slide_transition_effect_options = {}
                        slide_transition_duration = transition_duration
                        slide_auto_advance = None
                        slide_transition_sound_path = None
                        slide_animation = None
                        slide_animation_duration = animation_duration
                        slide_animation_stagger = animation_stagger
                        slide_animation_trigger = animation_trigger
                        slide_animation_cfg = {}
                    else:
                        (
                            slide_transition,
                            slide_transition_effect_options,
                            slide_transition_duration,
                            slide_auto_advance,
                            slide_transition_sound_path,
                        ) = _slide_transition_settings(
                            default_transition_cfg,
                            slide_cfg,
                            transition,
                            transition_effect_options,
                            transition_duration,
                            auto_advance,
                            transition_sound,
                            animation_cli_overrides,
                        )
                        (
                            slide_animation,
                            slide_animation_duration,
                            slide_animation_stagger,
                            slide_animation_trigger,
                            slide_animation_cfg,
                        ) = _slide_animation_settings(
                            slide_cfg,
                            default_animation_cfg,
                            animation,
                            animation_duration,
                            animation_stagger,
                            animation_trigger,
                            animation_cli_overrides,
                        )
                        if morph_pairs_by_destination.get(svg_path.stem):
                            if (
                                slide_transition != "morph"
                                or slide_transition_effect_options.get(
                                    "morph_by",
                                    "object",
                                )
                                != "object"
                            ):
                                raise ValueError(
                                    f'animations.json slide "{svg_path.stem}" '
                                    'declares deterministic Morph pairs, but '
                                    'the resolved transition is not Morph by object'
                                )
                    groups_value = slide_cfg.get('groups', {})
                    if not isinstance(groups_value, dict):
                        raise ValueError(
                            'animations.json slide groups must be an object'
                        )
                    animation_hard_disabled = (
                        animation_cli_overrides.get('animation', False)
                        and animation is None
                    )
                    explicit_group_ids: set[str] = set()
                    trigger_group_ids: set[str] = set()
                    if not animation_hard_disabled:
                        for group_id, group_cfg in groups_value.items():
                            if not isinstance(group_cfg, dict):
                                continue
                            group_path = (
                                f'slides['
                                f'{json.dumps(svg_path.stem, ensure_ascii=False)}'
                                f'].groups['
                                f'{json.dumps(str(group_id), ensure_ascii=False)}'
                                f']'
                            )
                            effect_entries = animation_group_effect_entries(
                                group_cfg,
                                path=group_path,
                            )
                            if any(
                                effect_cfg.get('effect') != 'none'
                                and (
                                    slide_animation is not None
                                    or 'effect' in effect_cfg
                                )
                                for _effect_path, effect_cfg in effect_entries
                            ):
                                explicit_group_ids.add(str(group_id))
                            for _effect_path, effect_cfg in effect_entries:
                                trigger_shape = effect_cfg.get('trigger_shape')
                                if (
                                    isinstance(trigger_shape, str)
                                    and trigger_shape.strip()
                                ):
                                    trigger_group_ids.add(trigger_shape)
                    explicit_animation_groups = frozenset(
                        explicit_group_ids | trigger_group_ids
                    )
                    if trigger_group_ids:
                        hyperlink_trigger_errors = trigger_shape_hyperlink_errors(
                            ET.parse(svg_path).getroot(),
                            trigger_group_ids,
                        )
                        if hyperlink_trigger_errors:
                            raise ValueError('; '.join(hyperlink_trigger_errors))
                    converter_group_overrides = (
                        explicit_animation_groups
                        | frozenset(
                            morph_group_overrides_by_slide.get(
                                svg_path.stem,
                                set(),
                            )
                        )
                    )
                    (
                        slide_xml,
                        media_files_dict,
                        rel_entries,
                        anim_targets,
                        package_files_dict,
                        content_type_overrides,
                    ) = (
                        convert_svg_to_slide_shapes(
                            svg_path, slide_num=slide_num,
                            slide_count=public_slide_count,
                            verbose=verbose,
                            text_flow=text_flow,
                            image_optimize=image_optimize,
                            image_max_dimension=image_max_dimension,
                            image_sizing=image_sizing,
                            image_scale=image_scale,
                            image_quality=image_quality,
                            native_objects=native_objects,
                            animation_group_overrides=converter_group_overrides,
                            theme_font_spec=active_theme_font_spec,
                            theme_color_spec=active_theme_color_spec,
                            primary_language=primary_language,
                            promote_background=pptx_structure != "structured",
                            dangerous_nonconforming_export=(
                                dangerous_nonconforming_export
                            ),
                            resource_root=resource_root,
                            trace_out=conversion_trace
                            if conversion_trace is not None
                            else structure_trace,
                        )
                    )
                    morph_group_ids = morph_group_overrides_by_slide.get(
                        svg_path.stem,
                        set(),
                    )
                    if morph_group_ids:
                        target_ids_by_group: dict[str, list[int]] = {}
                        for shape_id, group_id in anim_targets:
                            target_ids_by_group.setdefault(
                                str(group_id),
                                [],
                            ).append(int(shape_id))
                        for group_id in sorted(morph_group_ids):
                            resolved_shape_ids = target_ids_by_group.get(
                                group_id,
                                [],
                            )
                            if len(resolved_shape_ids) != 1:
                                raise ValueError(
                                    f'Morph target "{svg_path.stem}/{group_id}" '
                                    'must resolve to exactly one Slide-local '
                                    'PowerPoint shape'
                                )
                            morph_shape_ids[
                                (svg_path.stem, group_id)
                            ] = resolved_shape_ids[0]
                    # Order matters: OOXML schema requires <p:transition>
                    # to precede <p:timing> inside <p:sld>. Both use the same
                    # </p:sld> string-replace anchor, so transition must be
                    # injected first and timing second.
                    packaged_sounds_by_source: dict[Path, tuple[str, str]] = {}
                    expected_transition_sound = _materialize_transition_sound(
                        (
                            animation_resource_root
                            if animation_resource_root is not None
                            else svg_files[0].parent.parent
                        ),
                        slide_transition_sound_path,
                        media_files_dict,
                        rel_entries,
                        audio_exts_used,
                        packaged_sounds_by_source,
                    )
                    if (
                        slide_transition is not None
                        or slide_auto_advance is not None
                        or expected_transition_sound is not None
                    ):
                        transition_fragment = create_transition_xml(
                            effect=slide_transition,
                            duration=slide_transition_duration,
                            advance_after=slide_auto_advance,
                            effect_options=slide_transition_effect_options,
                            sound=expected_transition_sound,
                        )
                        if transition_fragment:
                            slide_xml = slide_xml.replace(
                                '</p:sld>',
                                '\n' + transition_fragment + '\n</p:sld>',
                            )
                        if slide_auto_advance is not None:
                            package_uses_timings = True

                    expected_animation_duration = slide_animation_duration
                    expected_animation_trigger = slide_animation_trigger
                    if (
                        not animation_hard_disabled
                        and (slide_animation or explicit_animation_groups)
                        and anim_targets
                    ):
                        seq_targets, mixed_count = _build_sequence_targets(
                            anim_targets,
                            svg_path.stem,
                            slide_cfg,
                            slide_animation,
                            slide_animation_cfg,
                            slide_animation_duration,
                            slide_animation_stagger,
                            mixed_animation_offset,
                            animation_rng,
                        )
                        seq_targets = _materialize_animation_sounds(
                            (
                                animation_resource_root
                                if animation_resource_root is not None
                                else svg_files[0].parent.parent
                            ),
                            seq_targets,
                            media_files_dict,
                            rel_entries,
                            audio_exts_used,
                            packaged_sounds_by_source,
                        )
                        expected_animation_targets = seq_targets
                        if mixed_count:
                            mixed_animation_offset += mixed_count
                        timing_xml = '\n' + create_sequence_timing_xml(
                            seq_targets, duration=slide_animation_duration,
                            trigger=slide_animation_trigger,
                        )
                        slide_xml = slide_xml.replace(
                            '</p:sld>',
                            timing_xml + '\n</p:sld>',
                        )

                    # Write slide XML
                    slide_xml_path = extract_dir / 'ppt' / 'slides' / f'slide{slide_num}.xml'
                    with open(slide_xml_path, 'w', encoding='utf-8') as f:
                        f.write(slide_xml)

                    # Write media files
                    media_name_map: dict[str, str] = {}
                    for media_name, media_data in media_files_dict.items():
                        ext = media_name.rsplit('.', 1)[-1].lower()
                        media_hash = hashlib.sha256(media_data).hexdigest()
                        cache_key = (ext, media_hash)
                        cached_name = media_cache.get(cache_key)

                        if cached_name is None:
                            prefix = (
                                'audio'
                                if f'.{ext}' in AUDIO_CONTENT_TYPES
                                else 'image'
                            )
                            cached_name = f'{prefix}_{media_hash[:16]}.{ext}'
                            media_cache[cache_key] = cached_name
                            with open(media_dir / cached_name, 'wb') as f:
                                f.write(media_data)

                        media_name_map[media_name] = cached_name

                    for rel in rel_entries:
                        target = rel.get('target', '')
                        if not target.startswith('../media/'):
                            continue
                        media_name = target.split('../media/', 1)[1]
                        mapped_name = media_name_map.get(media_name)
                        if mapped_name:
                            rel['target'] = f'../media/{mapped_name}'

                    # Write non-media OOXML package parts produced by native
                    # object converters, e.g. chart XML, chart rels, and
                    # embedded workbooks.
                    for part_name, part_data in package_files_dict.items():
                        if (
                            part_name.startswith('ppt/charts/')
                            and part_name.endswith('.xml')
                        ):
                            part_data = rewrite_chart_accent_colors(
                                part_data,
                                active_theme_color_spec,
                            )
                        package_path = extract_dir / part_name
                        package_path.parent.mkdir(parents=True, exist_ok=True)
                        with open(package_path, 'wb') as f:
                            f.write(part_data)
                        suffix = package_path.suffix.lstrip('.').lower()
                        if suffix:
                            package_exts_used.add(suffix)
                    package_content_overrides.update(content_type_overrides)

                    # Build relationships XML
                    rels_dir = extract_dir / 'ppt' / 'slides' / '_rels'
                    rels_dir.mkdir(exist_ok=True)
                    rels_path = rels_dir / f'slide{slide_num}.xml.rels'

                    extra_rels = ''
                    for rel in rel_entries:
                        target_mode = rel.get('target_mode')
                        mode_attr = (
                            f" TargetMode={quoteattr(target_mode)}"
                            if target_mode is not None
                            else ''
                        )
                        extra_rels += (
                            f"\n  <Relationship Id={quoteattr(rel['id'])} "
                            f"Type={quoteattr(rel['type'])} "
                            f"Target={quoteattr(rel['target'])}{mode_attr}/>"
                        )

                    rels_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1"
                Type="{SLIDE_LAYOUT_REL_TYPE}"
                Target="{structure.slide_layout_target(slide_num)}"/>{extra_rels}
</Relationships>'''
                    with open(rels_path, 'w', encoding='utf-8') as f:
                        f.write(rels_xml)

                    # Track image formats for Content_Types
                    for media_name in media_name_map.values():
                        ext = media_name.rsplit('.', 1)[-1].lower()
                        dotted_ext = f'.{ext}'
                        if dotted_ext in AUDIO_CONTENT_TYPES:
                            audio_exts_used.add(dotted_ext)
                        else:
                            _content_type_for_extension(ext)
                            image_exts_used.add(ext)
                            has_any_image = True

                # ---- Legacy SVG embedding mode ----
                else:
                    slide_cfg = _slide_config(animation_config, svg_path.stem)
                    (
                        slide_transition,
                        slide_transition_effect_options,
                        slide_transition_duration,
                        slide_auto_advance,
                        slide_transition_sound_path,
                    ) = (
                        _slide_transition_settings(
                            default_transition_cfg,
                            slide_cfg,
                            transition,
                            transition_effect_options,
                            transition_duration,
                            auto_advance,
                            transition_sound,
                            animation_cli_overrides,
                        )
                    )
                    svg_filename = f'image{i}.svg'
                    png_filename = f'image{i}.png'
                    png_rid = 'rId2'
                    svg_rid = 'rId3' if use_compat_mode else 'rId2'

                    shutil.copy(svg_path, media_dir / svg_filename)

                    slide_has_png = False
                    if use_compat_mode:
                        if prerender_results is not None:
                            png_success = prerender_results.get(i, False)
                        else:
                            png_path = media_dir / png_filename
                            png_success = convert_svg_to_png(
                                svg_path, png_path,
                                width=pixel_width, height=pixel_height,
                            )
                        if png_success:
                            slide_has_png = True
                            has_any_image = True
                            image_exts_used.add('png')
                        else:
                            if verbose:
                                print(
                                    f"  [{i}/{len(svg_files)}] {svg_path.name} - "
                                    "PNG generation failed, using pure SVG"
                                )
                            svg_rid = 'rId2'

                    slide_xml_path = extract_dir / 'ppt' / 'slides' / f'slide{slide_num}.xml'
                    slide_xml = create_slide_xml_with_svg(
                        slide_num,
                        png_rid=png_rid, svg_rid=svg_rid,
                        width_emu=width_emu, height_emu=height_emu,
                        transition=slide_transition,
                        transition_effect_options=slide_transition_effect_options,
                        transition_duration=slide_transition_duration,
                        auto_advance=slide_auto_advance,
                        use_compat_mode=(use_compat_mode and slide_has_png),
                    )
                    with open(slide_xml_path, 'w', encoding='utf-8') as f:
                        f.write(slide_xml)

                    rels_dir = extract_dir / 'ppt' / 'slides' / '_rels'
                    rels_dir.mkdir(exist_ok=True)
                    rels_path = rels_dir / f'slide{slide_num}.xml.rels'
                    rels_xml = create_slide_rels_xml(
                        png_rid=png_rid, png_filename=png_filename,
                        svg_rid=svg_rid, svg_filename=svg_filename,
                        use_compat_mode=(use_compat_mode and slide_has_png),
                        slide_layout_target=structure.slide_layout_target(slide_num),
                    )
                    with open(rels_path, 'w', encoding='utf-8') as f:
                        f.write(rels_xml)

                if slide_patch is not None:
                    if source_slide_bytes is None or source_rels_bytes is None:
                        raise TemplateStructureError(
                            f"Round-trip slide {slide_num} source snapshot is missing"
                        )
                    trace_sink = (
                        conversion_trace
                        if conversion_trace is not None
                        else structure_trace
                    )
                    slide_conversion_trace = next(
                        (
                            entry
                            for entry in reversed(trace_sink or [])
                            if entry.get("slide_num") == slide_num
                        ),
                        None,
                    )
                    overlay_shape_id_map = _apply_roundtrip_slide_overlay(
                        source_slide_bytes,
                        source_rels_bytes,
                        slide_xml_path,
                        rels_path,
                        slide_patch,
                        slide_conversion_trace,
                    )
                    for target in expected_animation_targets:
                        for field in ("shape_id", "trigger_shape_id"):
                            shape_id = target.get(field)
                            if (
                                isinstance(shape_id, int)
                                and not isinstance(shape_id, bool)
                                and shape_id in overlay_shape_id_map
                            ):
                                target[field] = overlay_shape_id_map[shape_id]
                    for morph_key, shape_id in tuple(
                        morph_shape_ids.items()
                    ):
                        if (
                            morph_key[0] == svg_path.stem
                            and shape_id in overlay_shape_id_map
                        ):
                            morph_shape_ids[morph_key] = (
                                overlay_shape_id_map[shape_id]
                            )

                resolved_advance_after = slide_auto_advance
                resolved_advance_on_click = True

                # --- Process notes (shared between native and legacy mode) ---
                notes_changed = slide_patch is None or slide_patch.notes_changed
                if not is_layout_definition and notes_changed:
                    notes_content = (
                        notes.get(svg_path.stem, '') if notes else ''
                    )
                    notes_master = _apply_slide_notes(
                        extract_dir,
                        rels_path,
                        slide_num,
                        notes_content,
                        primary_language,
                        enable_notes=enable_notes,
                    )
                    if notes_master is not None:
                        notes_slides_created.add(slide_num)
                        notes_master_parts_used.add(notes_master.package_part)
                        if notes_master.created_theme_part is not None:
                            notes_master_theme_parts_created.add(
                                notes_master.created_theme_part
                            )

                # --- Process narration audio (shared between native and legacy mode) ---
                svg_stem = svg_path.stem
                audio_path = (
                    narration_audio.get(svg_stem)
                    if narration_audio and not is_layout_definition
                    else None
                )
                if audio_path:
                    slide_xml_path = extract_dir / 'ppt' / 'slides' / f'slide{slide_num}.xml'
                    rels_path = extract_dir / 'ppt' / 'slides' / '_rels' / f'slide{slide_num}.xml.rels'

                    ext = audio_path.suffix.lower()
                    media_name = f'narration{slide_num}{ext}'
                    shutil.copy2(audio_path, media_dir / media_name)
                    audio_exts_used.add(ext)

                    poster_name = 'narration_poster.png'
                    poster_path = media_dir / poster_name
                    if not poster_path.exists():
                        poster_path.write_bytes(AUDIO_MARKER_PNG_BYTES)
                    has_any_image = True
                    image_exts_used.add('png')

                    media_rid = _append_relationship(
                        rels_path,
                        MEDIA_REL_TYPE,
                        f'../media/{media_name}',
                    )
                    audio_rid = _append_relationship(
                        rels_path,
                        AUDIO_REL_TYPE,
                        f'../media/{media_name}',
                    )
                    poster_rid = _append_relationship(
                        rels_path,
                        IMAGE_REL_TYPE,
                        f'../media/{poster_name}',
                    )

                    slide_xml = slide_xml_path.read_text(encoding='utf-8')
                    narration_shape_id = next_shape_id(slide_xml)
                    narration_transition_duration = (
                        slide_transition_duration
                        if slide_transition is not None
                        else 0.0
                    )
                    narration_lead_in = narration_lead_in_seconds(
                        narration_transition_duration,
                        start_floor=narration_start_floor,
                    )
                    slide_xml = inject_narration(
                        slide_xml,
                        shape_id=narration_shape_id,
                        shape_name=media_name,
                        audio_rid=audio_rid,
                        media_rid=media_rid,
                        poster_rid=poster_rid,
                        start_delay=narration_lead_in,
                    )

                    if use_narration_timings:
                        duration = probe_audio_duration(audio_path)
                        if duration is None:
                            raise RuntimeError(
                                f"Unable to read narration duration with ffprobe: {audio_path}"
                            )
                        narration_advance_after = (
                            narration_lead_in + duration + narration_padding
                        )
                        slide_xml = apply_recorded_timing(
                            slide_xml,
                            advance_after=narration_advance_after,
                            transition_duration=slide_transition_duration,
                            transition_effect=slide_transition,
                        )
                        resolved_advance_after = narration_advance_after
                        resolved_advance_on_click = False
                        package_uses_timings = True
                    slide_xml_path.write_text(slide_xml, encoding='utf-8')
                    narration_slides_created.add(slide_num)

                final_slide_xml = slide_xml_path.read_text(encoding='utf-8')
                preserved_source_transition = (
                    slide_patch is not None
                    and not slide_patch.transition_replaced
                )
                preserved_source_animation = (
                    slide_patch is not None and not slide_patch.animation_changed
                )
                resolved_motion = None
                resolved_animation = None
                if not preserved_source_transition:
                    try:
                        resolved_motion = validate_generated_transition_xml(
                            final_slide_xml,
                            effect=slide_transition,
                            effect_options=slide_transition_effect_options,
                            duration=slide_transition_duration,
                            advance_on_click=resolved_advance_on_click,
                            advance_after=resolved_advance_after,
                            sound=expected_transition_sound,
                        )
                    except ValueError as exc:
                        raise RuntimeError(
                            f'Slide {slide_num} transition validation failed: {exc}'
                        ) from exc
                if not preserved_source_animation:
                    try:
                        resolved_animation = validate_generated_animation_xml(
                            final_slide_xml,
                            expected_animation_targets,
                            duration=expected_animation_duration,
                            trigger=expected_animation_trigger,
                        )
                    except ValueError as exc:
                        raise RuntimeError(
                            f'Slide {slide_num} animation validation failed: {exc}'
                        ) from exc

                if conversion_trace is not None:
                    motion_summary = (
                        {"source_preserved": True}
                        if resolved_motion is None
                        else asdict(resolved_motion)
                    )
                    for trace_entry in reversed(conversion_trace):
                        if trace_entry.get('slide_num') == slide_num:
                            trace_entry['motion'] = motion_summary
                            trace_entry['animation'] = (
                                {"source_preserved": True}
                                if resolved_animation is None
                                else asdict(resolved_animation)
                            )
                            break

                if verbose:
                    if use_native_shapes:
                        mode_str = " (Native)"
                    elif use_compat_mode and not use_native_shapes:
                        mode_str = " (PNG+SVG)" if has_any_image else " (SVG)"
                    else:
                        mode_str = " (SVG)"
                    has_notes = slide_num in notes_slides_created
                    notes_str = " +notes" if has_notes else ""
                    narration_str = " +narration" if slide_num in narration_slides_created else ""
                    print(
                        f"  {progress_label} {svg_path.name}{mode_str}"
                        f"{notes_str}{narration_str}"
                    )

                success_count += 1

            except Exception as e:
                if verbose:
                    print(
                        f"  {progress_label} {svg_path.name} - Error: {e}"
                    )
                if use_native_shapes:
                    raise

        if (
            use_native_shapes
            and pptx_structure == "baseline"
            and success_count == len(svg_files)
        ):
            _convert_page_number_texts_to_fields(
                extract_dir,
                len(svg_files),
                conversion_trace if conversion_trace is not None else structure_trace,
                context="Baseline",
                verbose=verbose,
            )
            _promote_common_slide_backgrounds_to_masters(
                extract_dir,
                structure,
                len(svg_files),
                verbose=verbose,
            )
            _promote_common_chrome_shapes_to_masters(
                extract_dir,
                structure,
                len(svg_files),
                conversion_trace if conversion_trace is not None else structure_trace,
                verbose=verbose,
            )
            _extract_baseline_layout_families(
                extract_dir,
                structure,
                svg_files,
                verbose=verbose,
            )
            _promote_common_chrome_shapes_to_layouts(
                extract_dir,
                len(svg_files),
                conversion_trace if conversion_trace is not None else structure_trace,
                verbose=verbose,
            )
            _prune_unused_slide_layouts(
                extract_dir,
                structure,
                len(svg_files),
                verbose=verbose,
            )

        if (
            use_native_shapes
            and pptx_structure == "flat"
            and success_count == len(svg_files)
        ):
            _prepare_flat_structure(
                extract_dir,
                structure,
                len(svg_files),
                master_text_style_spec,
                structure_name,
                verbose=verbose,
            )

        if (
            use_native_shapes
            and pptx_structure == "structured"
            and success_count == len(svg_files)
        ):
            _convert_page_number_texts_to_fields(
                extract_dir,
                len(svg_files),
                conversion_trace if conversion_trace is not None else structure_trace,
                context="Structured",
                verbose=verbose,
            )
            if template_specs is None:
                raise TemplateStructureError(
                    "Structured metadata was not parsed before export"
                )
            (
                template_background_expectations,
                template_shape_roster_expectations,
                template_layout_parts_by_key,
                template_master_parts_by_key,
            ) = _apply_explicit_layout_structure(
                extract_dir,
                structure,
                template_specs,
                conversion_trace if conversion_trace is not None else structure_trace,
                active_theme_font_spec,
                use_layout_placeholder_frames=use_layout_placeholder_frames,
                verbose=verbose,
            )
            if source_theme_xml_by_master is not None:
                _install_source_themes_by_master(
                    extract_dir,
                    template_master_parts_by_key,
                    source_theme_xml_by_master,
                )
            master_count = apply_master_text_style_spec(
                extract_dir,
                master_text_style_spec,
            )
            if verbose:
                print(
                    "  Structured master text styles: "
                    f"{master_count} master(s), "
                    f"title {master_text_style_spec.title_hpt / 100:g}pt, "
                    "body levels "
                    f"{master_text_style_spec.body_levels_hpt[0] / 100:g}–"
                    f"{master_text_style_spec.body_levels_hpt[-1] / 100:g}pt"
                )
            _prune_unused_slide_layouts(
                extract_dir,
                structure,
                len(svg_files),
                verbose=verbose,
            )

        if (
            use_native_shapes
            and pptx_structure == "preserve"
            and success_count == len(svg_files)
        ):
            _convert_page_number_texts_to_fields(
                extract_dir,
                len(svg_files),
                conversion_trace if conversion_trace is not None else structure_trace,
                context="Preserve",
                verbose=verbose,
            )
            if template_specs is None or native_structure_contract is None:
                raise TemplateStructureError(
                    "Preserved structure metadata was not parsed before export"
                )
            rebuilt_specs = [
                spec
                for spec in template_specs
                if spec.slide_num not in passthrough_slides | overlay_slides
            ]
            if rebuilt_specs:
                _apply_preserved_structure(
                    extract_dir,
                    rebuilt_specs,
                    native_structure_contract,
                    conversion_trace if conversion_trace is not None else structure_trace,
                    verbose=verbose,
                )

        if roundtrip_export:
            pruned_roundtrip_payloads = (
                _prune_unreferenced_definition_payload_parts(
                    extract_dir,
                    preserved_parts=roundtrip_source_parts,
                )
            )
            if verbose and pruned_roundtrip_payloads:
                print(
                    "  Round-trip overlay: pruned "
                    f"{pruned_roundtrip_payloads} unused generated payload part(s)"
                )

        if (
            use_native_shapes
            and pptx_structure == "structured"
            and definition_svg_files
            and success_count == total_slide_count
        ):
            removed = _remove_trailing_layout_definition_slides(
                extract_dir,
                public_slide_count,
                total_slide_count,
            )
            pruned_payload_parts = _prune_unreferenced_definition_payload_parts(
                extract_dir
            )
            for slide_num in range(public_slide_count + 1, total_slide_count + 1):
                slide_part = f"ppt/slides/slide{slide_num}.xml"
                if template_background_expectations is not None:
                    template_background_expectations.pop(slide_part, None)
                if template_shape_roster_expectations is not None:
                    template_shape_roster_expectations.pop(slide_part, None)
            if verbose:
                print(
                    "  Layout definition carriers: "
                    f"removed {removed} internal slide(s), pruned "
                    f"{pruned_payload_parts} orphan payload part(s)"
                )

        morph_trace_names = _apply_morph_shape_names(
            extract_dir,
            morph_pairs,
            public_slide_numbers,
            morph_shape_ids,
        )
        if template_shape_roster_expectations is not None:
            for slide_number in morph_trace_names:
                slide_part = f"ppt/slides/slide{slide_number}.xml"
                template_shape_roster_expectations[
                    slide_part
                ] = _top_level_shape_name_roster(
                    ET.parse(extract_dir / slide_part).getroot()
                )
        if conversion_trace is not None:
            for trace_entry in conversion_trace:
                slide_number = int(trace_entry.get("slide_num", 0))
                names = morph_trace_names.get(slide_number)
                if names:
                    trace_entry["morph_names"] = dict(sorted(names.items()))

        # Update [Content_Types].xml
        content_types_path = extract_dir / '[Content_Types].xml'
        with open(content_types_path, 'r', encoding='utf-8') as f:
            content_types = f.read()

        if not use_native_shapes:
            content_types = _add_default_content_type(content_types, 'svg', 'image/svg+xml')
        for ext in sorted(image_exts_used):
            content_types = _add_default_content_type(
                content_types,
                ext,
                _content_type_for_extension(ext),
            )
        if 'xlsx' in package_exts_used:
            content_types = _add_default_content_type(
                content_types,
                'xlsx',
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        for part_name, content_type in sorted(package_content_overrides.items()):
            content_types = _add_content_type_override(content_types, part_name, content_type)
        with open(content_types_path, 'w', encoding='utf-8') as f:
            f.write(content_types)

        if audio_exts_used:
            for ext in sorted(audio_exts_used):
                content_type = AUDIO_CONTENT_TYPES.get(ext)
                if content_type:
                    content_types = _add_default_content_type(
                        content_types,
                        ext.removeprefix('.'),
                        content_type,
                    )
            if 'Extension="png"' not in content_types:
                content_types = _add_default_content_type(content_types, 'png', 'image/png')
            with open(content_types_path, 'w', encoding='utf-8') as f:
                f.write(content_types)

        # Add notes master / slides content types
        if enable_notes and notes_slides_created:
            for part_name in sorted(notes_master_theme_parts_created):
                content_types = _add_content_type_override(
                    content_types,
                    part_name,
                    THEME_CONTENT_TYPE,
                )
            for part_name in sorted(notes_master_parts_used):
                content_types = _add_content_type_override(
                    content_types,
                    part_name,
                    _NOTES_MASTER_CONTENT_TYPE,
                )
            for i in sorted(notes_slides_created):
                content_types = _add_content_type_override(
                    content_types,
                    f'/ppt/notesSlides/notesSlide{i}.xml',
                    'application/vnd.openxmlformats-officedocument.presentationml.'
                    'notesSlide+xml',
                )
            with open(content_types_path, 'w', encoding='utf-8') as f:
                f.write(content_types)

        if source_content_types_bytes is not None:
            current_content_types = content_types_path.read_bytes()
            if (
                _content_type_contract(current_content_types)
                == _content_type_contract(source_content_types_bytes)
            ):
                content_types_path.write_bytes(source_content_types_bytes)

        if page_plan_export:
            pruned_page_plan_parts = prune_unreferenced_directory_parts(
                extract_dir
            )
            if verbose and pruned_page_plan_parts:
                print(
                    "  Round-trip page plan: pruned "
                    f"{pruned_page_plan_parts} unreachable package part(s)"
                )

        if package_uses_timings:
            set_directory_use_timings(extract_dir)

        reinjected_resources = _reinject_roundtrip_resources(
            extract_dir,
            resource_root,
            late_roundtrip_resources,
        )
        if verbose and reinjected_resources:
            print(
                "  Round-trip resources: reinjected "
                f"{reinjected_resources} source package part(s)"
            )

        rels_problems = verify_internal_relationships(extract_dir)
        if rels_problems:
            details = '\n'.join(f'  - {p}' for p in rels_problems)
            raise RuntimeError(
                'PPTX package contains dangling internal relationship targets; '
                'PowerPoint will report the file as corrupt:\n' + details
            )

        # Replace the python-pptx base-template metadata (stale "Steve Canny"
        # author, 2013 dates, "generated using python-pptx", Slides=0) with
        # accurate, tool-neutral document properties.
        pres_format = _presentation_format(width_emu, height_emu)
        effective_doc_metadata = dict(doc_metadata or {})
        if primary_language is not None:
            effective_doc_metadata['language'] = primary_language
        if not roundtrip_export or doc_metadata:
            _stamp_docprops(
                extract_dir,
                public_slide_count,
                pres_format,
                effective_doc_metadata,
            )

        # Repackage PPTX to a temporary file first. The public output path is
        # replaced only after every slide and relationship has succeeded.
        temp_output_path = temp_dir / 'result.pptx'
        with zipfile.ZipFile(temp_output_path, 'w', zipfile.ZIP_DEFLATED) as zf:
            for file_path in extract_dir.rglob('*'):
                if file_path.is_file():
                    arcname = file_path.relative_to(extract_dir)
                    zf.write(file_path, arcname)
        if (
            use_native_shapes
            and pptx_structure == "structured"
            and success_count == len(svg_files)
        ):
            if template_specs is None or public_template_specs is None:
                raise TemplateStructureError(
                    "Explicit Layout metadata was not parsed before validation"
                )
            try:
                validate_pptx_template_package(
                    temp_output_path,
                    public_template_specs,
                    layout_specs=template_specs,
                    expected_layout_parts=template_layout_parts_by_key,
                    expected_master_parts=template_master_parts_by_key,
                    expected_backgrounds=template_background_expectations,
                    expected_shape_rosters=template_shape_roster_expectations,
                )
            except ValueError as exc:
                raise TemplateStructureError(
                    f"PPTX structured package validation failed: {exc}"
                ) from exc
        try:
            validate_pptx_transition_package(
                temp_output_path,
                require_use_timings=package_uses_timings,
            )
        except ValueError as exc:
            raise RuntimeError(
                f'PPTX transition package validation failed: {exc}'
            ) from exc
        try:
            validate_pptx_morph_pairs(
                temp_output_path,
                morph_expectations,
            )
        except ValueError as exc:
            raise RuntimeError(
                f'PPTX Morph package validation failed: {exc}'
            ) from exc
        try:
            validate_pptx_animation_package(
                temp_output_path,
                require_supported_effects=True,
                skip_slide_numbers=source_motion_slides,
            )
        except ValueError as exc:
            raise RuntimeError(
                f'PPTX animation package validation failed: {exc}'
            ) from exc
        shutil.move(str(temp_output_path), str(output_path))
        permission_warnings = _relax_output_permissions(output_path)

        if conversion_trace_path and conversion_trace is not None:
            conversion_trace_path.parent.mkdir(parents=True, exist_ok=True)
            payload = {
                'output': str(output_path),
                'slide_count': public_slide_count,
                'project_contract': {
                    'mode': (
                        'strict-after-dangerous-normalization'
                        if dangerous_nonconforming_export
                        else 'strict'
                    ),
                },
                'slides': [
                    entry
                    for entry in conversion_trace
                    if int(entry.get('slide_num', 0)) <= public_slide_count
                ],
            }
            conversion_trace_path.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2),
                encoding='utf-8',
            )

        if verbose:
            print()
            print(f"[Done] Saved: {output_path}")
            for warning in permission_warnings:
                print(f"  [warn] {warning}")
            if conversion_trace_path and conversion_trace is not None:
                print(f"  Trace: {conversion_trace_path}")
            print(
                f"  Slides: {public_slide_count}; "
                "unselected Layout prototype carriers: "
                f"{len(definition_svg_files)}"
            )
            if use_compat_mode and has_any_image:
                print(f"  Mode: Office compatibility mode (supports all Office versions)")
                if PNG_RENDERER == 'svglib' and renderer_hint:
                    print(f"  [Tip] {renderer_hint}")

        return success_count == len(svg_files)

    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)
